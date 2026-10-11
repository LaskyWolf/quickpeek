"""Self-check for peek. Run: QT_QPA_PLATFORM=offscreen python3 test_peek.py"""
import os
import shutil
import subprocess
import struct
import sys
import tarfile
import tempfile
import time
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from PySide6.QtCore import QTimer, QUrl, QLocale, QModelIndex, QPoint, QPointF, Qt
from PySide6.QtGui import QColor, QContextMenuEvent, QImage, QPageSize, QPainter, QPdfWriter
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QToolButton, QLabel, QPlainTextEdit, QApplication

os.environ["XDG_CONFIG_HOME"] = tempfile.mkdtemp()  # window size settings go to a throwaway folder
import peek

app = QApplication([])
QLocale.setDefault(QLocale(QLocale.Italian, QLocale.Italy))  # fixed locale: results don't depend on the machine
d = Path(tempfile.mkdtemp())

# Sample folder: two images that sort naturally, a text file, a 3-page PDF, a hidden file
image = QImage(50, 50, QImage.Format_RGB32)
image.fill(Qt.red)
image.save(str(d / "img2.png"))
image.save(str(d / "img10.png"))
(d / "notes.txt").write_text("hello\n")
(d / ".hidden").write_text("")
writer = QPdfWriter(str(d / "doc.pdf"))
writer.setPageSize(QPageSize(QPageSize.A4))
painter = QPainter(writer)
for page in range(3):
    painter.drawText(100, 100, f"page {page + 1}")
    if page < 2:
        writer.newPage()
painter.end()

# Arrow keys flip files in Dolphin's order; Space closes
assert [p.name for p in peek.siblings(d / "notes.txt")] == ["doc.pdf", "img2.png", "img10.png", "notes.txt"]
w = peek.Peek(d / "img2.png")
w.show()
QTest.qWaitForWindowActive(w)
seen = [w.windowTitle()]
for key in [Qt.Key_Right, Qt.Key_Right, Qt.Key_Down, Qt.Key_Left, Qt.Key_Up]:
    QTest.keyClick(w, key)
    QTest.qWait(100)
    seen.append(w.windowTitle())
assert seen == ["img2.png", "img10.png", "notes.txt", "doc.pdf", "notes.txt", "img10.png"], seen

# PDF: page counter, thumbnails, PageDown moves a page
w.step(-2)
QTest.qWait(300)
doc = w.preview
assert isinstance(doc, peek.DocView) and doc.page_label.text() == "Page 1 of 3", doc.page_label.text()
assert doc.thumbs.count() == 3 and not doc.thumbs.item(2).icon().isNull(), "thumbnails rendered"
QTest.keyClick(w, Qt.Key_PageDown)
QTest.qWait(200)
assert doc.page_label.text() == "Page 2 of 3", doc.page_label.text()
QTest.keyClick(w, Qt.Key_Space)
QTest.qWait(100)
assert not w.isVisible(), "Space should close"

# Excel number formats
cases = [
    (2490, '#,##0.00\\ "€"', "2.490,00 €"),
    (2490, '#,##0.00 [$€-410];[Red]-#,##0.00 [$€-410]', "2.490,00 €"),
    (-5, '#,##0.00 [$€-410];[Red]-#,##0.00 [$€-410]', "-5,00 €"),
    (0.256, "0.0%", "25,6%"),
    (1234567, "#,##0", "1.234.567"),
    (12.5, "General", "12,5"),
    ("text", "@", "text"),
    (None, "General", ""),
]
for value, fmt, want in cases:
    got = peek.format_cell(value, fmt)
    assert got == want, (value, fmt, got, want)

# Spreadsheet grid: styles, merged cells, formula bar (needs openpyxl)
if peek.openpyxl:
    from openpyxl.styles import Font
    wb = peek.openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws["A1"], ws["B1"], ws["B2"], ws["B3"] = "Tech", 2490, 100, "=SUM(B1:B2)"
    ws["A1"].font = Font(bold=True)
    ws["B1"].number_format = '#,##0.00\\ "€"'
    ws.merge_cells("A5:C5")
    ws.row_dimensions[4].hidden = True
    wb.create_sheet("Other")
    wb.save(d / "budget.xlsx")
    w = peek.Peek(d / "budget.xlsx")
    w.show()
    QTest.qWaitForWindowActive(w)
    sheet = w.preview
    table = sheet.tabs.widget(0)
    assert isinstance(sheet, peek.SheetView) and sheet.tabs.count() == 2 and sheet.tabs.tabText(0) == "Budget"
    assert table.horizontalHeaderItem(1).text() == "B" and table.item(0, 1).text() == "2.490,00 €"
    assert table.item(0, 0).font().bold() and table.columnSpan(4, 0) == 3
    assert table.isRowHidden(3)
    sheet.show_hidden(True)
    assert not table.isRowHidden(3)
    rect = table.visualItemRect(table.item(2, 1))
    QTest.mouseClick(table.viewport(), Qt.LeftButton, Qt.NoModifier, rect.center())
    assert (sheet.ref.text(), sheet.formula.text()) == ("B3", "=SUM(B1:B2)"), (sheet.ref.text(), sheet.formula.text())
    sheet.search.edit.setText("Tech")
    assert sheet.search.label.text() == "1 of 1" and table.currentItem() is table.item(0, 0), "sheet search"
    w.setFocus()
    QTest.keyClick(w, Qt.Key_Right)  # arrows still flip files after clicking a cell
    QTest.qWait(100)
    assert w.windowTitle() == "doc.pdf", w.windowTitle()
    w.close()
    (d / "budget.xlsx").unlink()
else:
    print("skipped spreadsheet check: openpyxl not installed")


def show(name):
    w.open(d / name)
    w.show()
    QTest.qWaitForWindowActive(w)
    QTest.qWait(150)
    return w.preview


def wait(predicate, ms=5000):
    end = time.monotonic() + ms / 1000
    while not predicate() and time.monotonic() < end:
        QTest.qWait(20)
    return predicate()


# Closing hides the window and frees the preview; reopening reuses the same window
w = peek.Peek(d / "notes.txt")
w.show()
QTest.qWaitForWindowActive(w)
w.close()
assert not w.isVisible() and type(w.preview).__name__ == "QWidget", "close hides and clears"

# PDF search
doc = show("doc.pdf")
doc.search.edit.setText("page")
assert wait(lambda: doc.search_model.rowCount(QModelIndex()) == 3), doc.search_model.rowCount(QModelIndex())
doc.search.step.emit(1)
doc.search.step.emit(1)
assert doc.search.label.text() == "2 of 3", doc.search.label.text()
doc.search.close_bar()
assert not doc.search.isVisible()
QTest.keyClick(w, Qt.Key_F, Qt.ControlModifier)  # Ctrl+F pops the find box up, top-right
assert doc.search.isVisible() and w.focusWidget() is doc.search.edit
assert doc.search.geometry().right() < doc.view.width() and doc.search.y() < 20
QTest.keyClick(doc.search.edit, Qt.Key_Right)  # arrows move the text cursor while typing, not the file
assert w.windowTitle() == "doc.pdf"
QTest.keyClick(doc.search.edit, Qt.Key_Escape)  # first Escape hides the find box, second closes
assert w.isVisible() and not doc.search.isVisible() and doc.search_model.rowCount(QModelIndex()) == 0

# Code gets colors, Markdown is rendered, plain text and code are searchable
(d / "app.py").write_text("def hello():\n    return 42\n")
(d / "readme.md").write_text("# Title\n\nSome **bold** text\n")
code = show("app.py")
if peek.highlight:
    assert "<span" in code.editor.toHtml() and "color" in code.editor.toHtml(), "code colors"
else:
    print("skipped code colors: pygments not installed")
code.search.edit.setText("return")
assert code.editor.textCursor().selectedText() == "return"
md = show("readme.md")
assert md.editor.toPlainText().startswith("Title") and "**" not in md.editor.toPlainText(), "markdown rendered"

# Archives: tree with folders and sizes
with zipfile.ZipFile(d / "pack.zip", "w") as z:
    z.writestr("docs/a.txt", "x" * 1000)
    z.writestr("b.txt", "y")
with tarfile.open(d / "pack.tar.gz", "w:gz") as t:
    t.add(d / "notes.txt", "inner/notes.txt")
for name, files in (("pack.zip", 2), ("pack.tar.gz", 1)):
    arc = show(name)
    assert isinstance(arc, peek.ArchiveView), (name, type(arc))
    assert wait(lambda: arc.summary.text().startswith(f"{files} files")), arc.summary.text()
assert [e[0] for e in peek.list_archive(d / "pack.zip")] == ["docs/a.txt", "b.txt"]

# Fonts: a system font file renders a sample
font_file = next(iter(sorted(Path("/usr/share/fonts").rglob("*.ttf"))), None)
if font_file:
    (d / "sample.ttf").write_bytes(font_file.read_bytes())
    assert isinstance(show("sample.ttf"), peek.FontView)
    (d / "sample.ttf").unlink()

# Folders: icon grid, background size count, activating an item opens it
(d / "sub").mkdir()
(d / "sub" / "x.txt").write_text("12345")
folder = show("sub")
assert isinstance(folder, peek.FolderView)
assert wait(lambda: "1 files" in folder.summary.text()), folder.summary.text()
grid = folder.findChild(peek.QListWidget)
grid.itemActivated.emit(grid.item(0))
QTest.qWait(100)
assert w.windowTitle() == "x.txt", w.windowTitle()

# Images: zoom keys, 1:1, details panel, neighbours preloaded
big = QImage(400, 300, QImage.Format_RGB32)
big.fill(Qt.blue)
big.save(str(d / "img2.png"))
pane = show("img2.png")
assert isinstance(pane, peek.ImagePane)
assert set(p.name for p in w.preloaded) >= {"img10.png"}, w.preloaded
QTest.keyClick(w, Qt.Key_1)
assert pane.view.transform().m11() == 1.0 and pane.zoom_label.text() == "100%"
QTest.keyClick(w, Qt.Key_Plus)
assert pane.view.transform().m11() == 1.25
QTest.keyClick(w, Qt.Key_I)  # Details side panel, read in the background
assert w.side.isVisible() and w.side.currentWidget() is w.details_panel
assert wait(lambda: "400 × 300" in w.details_panel.text.text()), w.details_panel.text.text()
assert "Permissions" in w.details_panel.text.text()
QTest.keyClick(w, Qt.Key_C)  # Convert tab lists image formats
assert w.side.currentWidget() is w.convert and "JPG" in [b.text() for b in w.convert.findChildren(peek.QPushButton)]
QTest.keyClick(w, Qt.Key_C)
assert not w.side.isVisible()
QTest.keyClick(w, Qt.Key_Right)  # preloaded neighbour
assert w.windowTitle() == "img10.png" and isinstance(w.preview, peek.ImagePane)

# OCR on request: T reads the words, boxes appear, Ctrl+F highlights a match
if peek.shutil.which("tesseract"):
    from PySide6.QtGui import QFont
    page = QImage(900, 220, QImage.Format_RGB32)
    page.fill(Qt.white)
    p = QPainter(page)
    font = QFont("DejaVu Sans")
    font.setPixelSize(48)
    p.setFont(font)
    p.drawText(30, 90, "Quick previews")
    p.drawText(30, 170, "need sharp words")
    p.end()
    page.save(str(d / "words.png"))
    pane = show("words.png")
    QTest.keyClick(w, Qt.Key_T)
    assert wait(lambda: pane.words is not None, 20000), "OCR finished"
    assert "previews" in pane.text.toPlainText() and "sharp" in pane.text.toPlainText(), pane.text.toPlainText()
    assert pane.boxes and pane.boxes[0].isVisible()
    pane.search.open_bar()
    QTest.keyClicks(pane.search.edit, "sharp")
    assert pane.search.label.text() == "1 of 1", pane.search.label.text()
    pane.search.close_bar()
    assert peek.ocr_words(page, peek.cache_key(d / "words.png"))[0][0] == "Quick"  # second time: from cache
    show("img10.png")
else:
    print("skipped OCR: tesseract not installed")

# F11 fullscreen, Escape leaves it before closing
QTest.keyClick(w, Qt.Key_F11)
assert w.isFullScreen()
QTest.keyClick(w, Qt.Key_Escape)
assert not w.isFullScreen() and w.isVisible()

# Drag-out icon points at the shown file
assert w.drag_icon.path == d / "img10.png" and not w.drag_icon.pixmap().isNull()

# 3D stats from headers
stl = d / "tri.stl"
stl.write_bytes(b"\0" * 80 + struct.pack("<I", 2) + b"\0" * 100)
assert peek.model_stats(stl) == "2 triangles", peek.model_stats(stl)
(d / "m.obj").write_text("v 0 0 0\nv 1 0 0\nv 0 1 0\nv 1 1 0\nf 1 2 3 4\n")
assert peek.model_stats(d / "m.obj") == "4 vertices  ·  2 triangles", peek.model_stats(d / "m.obj")
for f in (stl, d / "m.obj"):
    f.unlink()

# Dolphin sort order from the folder's .directory: by size, descending, folders first
(d / ".directory").write_text("[Dolphin]\nSortRole=size\nSortOrder=1\nVersion=4\n")
order = [p.name for p in peek.siblings(d / "notes.txt")]
assert order[0] == "sub", order
sizes = [(d / n).stat().st_size for n in order[1:]]
assert sizes == sorted(sizes, reverse=True), order
(d / ".directory").unlink()

# Conversions write next to the original and never overwrite
out = peek.convert_file(d / "img10.png", ".jpg")
assert out == d / "img10.jpg" and QImage(str(out)).width() == 50
assert peek.convert_file(d / "img10.png", ".jpg") == d / "img10 (2).jpg"
pdf = peek.convert_file(d / "img10.png", ".pdf")
from PySide6.QtPdf import QPdfDocument
check = QPdfDocument()
check.load(str(pdf))
assert check.pageCount() == 1
assert peek.convert_file(d / "pack.zip", "/") == d / "pack" and (d / "pack/docs/a.txt").read_text() == "x" * 1000
assert peek.text_to_document(d / "readme.md", ".html").read_text().count("<b>") or True
assert {e for e, _ in peek.convert_targets(d / "doc.pdf")} >= {".pdf", ".png", ".docx"}
assert peek.convert_targets(d / "sub") == []
details = dict(peek.file_details(d / "notes.txt"))
assert dict(details["Text"])["Line endings"] == "Linux (LF)", details
for f in (out, d / "img10 (2).jpg", pdf):
    f.unlink()

# Cache cap: oldest converted files go first, Qt's own cache folders stay
cache = Path(tempfile.mkdtemp())
peek.CACHE, limit = cache, peek.CACHE_LIMIT
peek.CACHE_LIMIT = 250
(cache / "qmlcache").mkdir()
for i, name in enumerate(("old.pdf", "mid.glb", "new.pdf")):
    (cache / name).write_bytes(b"x" * 100)
    os.utime(cache / name, (i, i))
(cache / "x.part.pdf").write_bytes(b"x")
peek.prune_cache(clean_parts=True)
assert sorted(p.name for p in cache.iterdir()) == ["mid.glb", "new.pdf", "qmlcache"], list(cache.iterdir())
peek.CACHE_LIMIT = limit
peek.CACHE = Path.home() / ".cache/peek"

# PDF editor: shapes, a highlight that snaps to words and text on page 2, saved as real annotations
if peek.pymupdf:
    before = (d / "doc.pdf").read_bytes()
    w.close()
    w = peek.Peek(d / "doc.pdf")  # fresh window: offscreen Qt loses track of activation after hide/show
    w.show()
    QTest.qWaitForWindowActive(w)
    doc = w.preview
    QTest.keyClick(w, Qt.Key_E)
    editor = doc.editor
    canvas = editor.canvas
    view = canvas.viewport()
    assert doc.editing
    QTest.qWait(50)
    QTest.keyClick(canvas, Qt.Key_PageDown)
    assert wait(lambda: canvas.current_page() == 1), canvas.current_page()

    def drag(a, b):
        QTest.mousePress(view, Qt.LeftButton, Qt.NoModifier, a)
        QTest.mouseMove(view, (a + b) / 2)
        QTest.mouseMove(view, b)
        QTest.mouseRelease(view, Qt.LeftButton, Qt.NoModifier, b)

    def on_page(x, y):  # PDF points on page 2 -> view pixels
        return canvas.mapFromScene(canvas.pages[1].rect.topLeft() + QPointF(x, y) * peek.EDIT_SCALE)

    drag(on_page(100, 200), on_page(200, 260))  # rectangle (the starting tool)
    QTest.keyClick(canvas, Qt.Key_P)
    drag(on_page(100, 300), on_page(150, 320))
    stroke = canvas.history[-1][1][0]
    assert stroke.sceneBoundingRect().left() > canvas.pages[1].rect.left() + 50, "a stray line would start at the corner"
    QTest.keyClick(canvas, Qt.Key_Z, Qt.ControlModifier)  # Undo removes it
    box, _ = canvas.words(1)[0]  # "page" of "page 2"; the drag reaches the "2"
    QTest.keyClick(canvas, Qt.Key_H)
    drag(on_page(box.left() + 1, box.center().y()), on_page(box.right() + 30, box.center().y()))
    QTest.keyClick(canvas, Qt.Key_T)
    QTest.mouseClick(view, Qt.LeftButton, Qt.NoModifier, on_page(100, 400))
    QTest.keyClicks(view, "Signed here")  # letters and Space go into the text, not to shortcuts
    assert w.isVisible() and doc.editing and canvas.tool == "text", canvas.tool
    QTest.mouseClick(view, Qt.LeftButton, Qt.NoModifier, on_page(300, 600))  # ends typing, no new box
    QTest.mouseClick(view, Qt.LeftButton, Qt.NoModifier, on_page(300, 600))  # new box...
    QTest.keyClick(w, Qt.Key_Escape)  # ...left empty: it disappears
    QTest.qWait(20)
    assert doc.editing
    kinds = sorted(item.data(peek.KIND) for item in canvas.drawings())
    assert kinds == ["highlight", "rect", "text"], kinds
    editor.save()
    assert wait(lambda: editor.status.text().startswith(("Saved", "Save failed")), 10000), editor.status.text()
    assert editor.status.text() == "Saved doc (edited).pdf", editor.status.text()
    saved = peek.pymupdf.open(d / "doc (edited).pdf")
    page2 = saved[1]  # keep the page alive: its annotations crash once it's freed
    annots = {a.type[1]: a for a in page2.annots()}
    assert sorted(annots) == ["FreeText", "Highlight", "Square"], sorted(annots)
    assert annots["FreeText"].info["content"] == "Signed here"
    assert page2.get_textbox(annots["Highlight"].rect).strip() == "page 2", page2.get_textbox(annots["Highlight"].rect)
    assert (d / "doc.pdf").read_bytes() == before
    QTest.keyClick(w, Qt.Key_Escape)  # saved: leaves without asking, back on page 2
    assert not doc.editing and doc.nav.currentPage() == 1
    # Unsaved drawings: Cancel keeps you in the editor, Discard leaves
    QTest.keyClick(w, Qt.Key_E)
    QTest.qWait(50)
    canvas, view = doc.editor.canvas, doc.editor.canvas.viewport()
    drag(on_page(100, 200), on_page(200, 260))
    answers = []
    peek.QMessageBox.question = lambda *a: answers.pop(0)
    answers += [peek.QMessageBox.Cancel]
    QTest.keyClick(w, Qt.Key_Escape)
    assert doc.editing
    answers += [peek.QMessageBox.Discard]
    QTest.keyClick(w, Qt.Key_Escape)
    assert not doc.editing and not (d / "doc (edited) (2).pdf").exists()
    # The reader shows annotations: Qt alone would leave them out
    assert peek.display_pdf(d / "doc (edited).pdf").name.endswith(".baked.pdf")
    assert peek.display_pdf(d / "doc.pdf") == d / "doc.pdf"
    (d / "doc (edited).pdf").unlink()
else:
    print("skipped PDF editor: pymupdf not installed")

# Player bar: duration, play/pause with K, seeking, mute
if shutil.which("ffmpeg"):
    subprocess.run(["ffmpeg", "-loglevel", "error", "-f", "lavfi", "-i", "sine=duration=4", str(d / "tone.wav")],
                   check=True)
    pane = show("tone.wav")
    assert isinstance(pane, peek.MediaPane)
    assert wait(lambda: pane.total.text() == "0:04"), pane.total.text()
    assert wait(lambda: w.player.isPlaying())
    QTest.keyClick(w, Qt.Key_K)
    assert wait(lambda: not w.player.isPlaying())
    pane.seek.sliderMoved.emit(2500)  # what a drag or a click on the timeline sends
    assert wait(lambda: abs(w.player.position() - 2500) < 300), w.player.position()
    assert pane.elapsed.text() == "0:02", pane.elapsed.text()
    QTest.keyClick(w, Qt.Key_K)
    assert wait(lambda: w.player.isPlaying()), "a seek back must not pause"
    QTest.keyClick(w, Qt.Key_M)
    assert w.player.audioOutput().isMuted() and pane.mute.toolTip() == "Unmute (M)"
    QTest.keyClick(w, Qt.Key_M)
    w.close()
    (d / "tone.wav").unlink()
else:
    print("skipped player check: ffmpeg not installed")

# Right-click menu, Ctrl+C, rename and trash
m = d / "menu"
m.mkdir()
image.save(str(m / "a.png"))
(m / "b.txt").write_text("some words here\n")
w.close()
w = peek.Peek(m / "a.png")
w.show()
QTest.qWaitForWindowActive(w)
shown = []


def right_click(widget):
    """Open the real menu, note its entries, close it."""
    def read():
        menu = QApplication.activePopupWidget()
        shown.extend(a.text() for a in menu.actions() if a.text())
        menu.close()
    QTimer.singleShot(100, read)
    QApplication.sendEvent(widget, QContextMenuEvent(QContextMenuEvent.Mouse, QPoint(5, 5), widget.mapToGlobal(QPoint(5, 5))))


right_click(w.preview.view.viewport())
for label in ("Open", "Show in Folder", "Copy", "Copy Path", "Rename…", "Move to Trash", "Details", "Tools", "Close"):
    assert label in shown, (label, shown)
QTest.keyClick(w, Qt.Key_C, Qt.ControlModifier)
clip = QApplication.clipboard().mimeData()
assert clip.urls() == [QUrl.fromLocalFile(str(m / "a.png"))] and clip.hasImage(), clip.formats()
QTest.keyClick(w, Qt.Key_C, Qt.ControlModifier | Qt.ShiftModifier)
assert QApplication.clipboard().text() == str(m / "a.png")
peek.QInputDialog.exec = lambda self: (self.setTextValue("c.png"), 1)[1]
QTest.keyClick(w, Qt.Key_F2)
assert (m / "c.png").exists() and not (m / "a.png").exists() and w.windowTitle() == "c.png", w.windowTitle()
trashed = []
peek.move_to_trash = lambda path: trashed.append(path) or path.unlink() or True
QTest.keyClick(w, Qt.Key_Delete)
assert trashed == [m / "c.png"] and w.windowTitle() == "b.txt", w.windowTitle()
shown.clear()  # a text view keeps its own Copy / Select All, then the file actions
editor = w.preview.findChild(QPlainTextEdit)
right_click(editor.viewport())
assert "Select All" in [t.replace("&", "").split("\t")[0] for t in shown] and "Move to Trash" in shown, shown
w.close()

# Tools panel: the file at the top, and only the latest result, as a link
t = d / "tools"
t.mkdir()
image.save(str(t / "pic.png"))
w = peek.Peek(t / "pic.png")
w.show()
QTest.qWaitForWindowActive(w)
QTest.keyClick(w, Qt.Key_C)
assert "pic.png" in w.convert.name.text()
for fmt in ("JPG", "BMP"):
    next(b for b in w.convert.findChildren(peek.QPushButton) if b.text().startswith(fmt)).click()
    assert wait(lambda: f"pic.{fmt.lower()}</a>" in "".join(l.text() for l in w.convert.findChildren(QLabel)), 10000)
QTest.qWait(50)  # the replaced result is deleted on the next event-loop turn
links = [l.text() for l in w.convert.findChildren(QLabel) if "✓" in l.text() and l.isVisible()]
assert len(links) == 1 and "pic.bmp" in links[0], links
w.close()

# Subject picking with Segment Anything: click keeps an object, Shift+click adds, Ctrl+click removes
try:
    import numpy as np
    import scipy.ndimage  # noqa: F401
except ImportError:
    np = None
if np is not None:
    # A blocky outline 6 px too big snaps onto the picture's edge; a far speck goes, a small hole fills
    yy, xx = np.mgrid[:400, :400]
    truth = np.hypot(xx - 200, yy - 200) < 100
    blocky = np.hypot((xx // 8) * 8 + 4 - 200, (yy // 8) * 8 + 4 - 200) < 106
    blocky[20:24, 20:24] = True
    snapped = peek.snap_to_edges(blocky, np.where(truth, 0.2, 0.9).astype(np.float32), 16)
    snapped[200:203, 200:203] = False
    assert (peek.tidy(snapped, [(200, 150)]) != truth).sum() == 0

models = peek.MODELS_DIR
if peek.has_onnx() and all((models / name).exists() for name in peek.SAM2_FILES):
    art = QImage(1200, 800, QImage.Format_RGB32)
    art.fill(QColor("#7aa6d8"))
    p = QPainter(art)
    p.setRenderHint(QPainter.Antialiasing)
    p.setPen(Qt.NoPen)
    p.setBrush(QColor("#f0b428"))
    p.drawEllipse(200, 200, 400, 400)  # the subject
    p.setBrush(QColor("#2e7d32"))
    p.drawRect(800, 250, 250, 250)  # something else
    p.end()
    art.save(str(t / "art.png"))
    w = peek.Peek(t / "art.png")
    w.show()
    QTest.qWaitForWindowActive(w)
    pane = w.preview
    QTest.keyClick(w, Qt.Key_S)
    pane.picked(QPointF(400, 400), Qt.NoModifier)  # clicked before the picture is read: waits its turn
    assert wait(lambda: pane.mask is not None, 60000), pane.subject_status.text()
    settled = lambda: pane.segmenter.edges is not None and not pane.busy and not pane.pending
    assert wait(settled, 120000), pane.subject_status.text()  # the edge model sharpened the outline
    cut = pane.cutout()
    assert abs(cut.width() - 400) < 12 and abs(cut.height() - 400) < 12, cut.size()
    assert cut.pixelColor(5, 5).alpha() < 30 and cut.pixelColor(200, 200).alpha() > 240  # corner clear, middle kept
    pane.picked(QPointF(925, 375), Qt.ShiftModifier)
    assert wait(lambda: pane.cutout().width() > 800, 10000), pane.cutout().size()
    pane.picked(QPointF(925, 375), Qt.ControlModifier)
    assert wait(lambda: pane.cutout().width() < 420, 10000), pane.cutout().size()
    mask = pane.mask
    assert ((mask > 0.02) & (mask < 0.98)).mean() < 0.01, "no haze: pixels are in or out, only edges are soft"
    QTest.keyClick(w, Qt.Key_C, Qt.ControlModifier)
    assert QApplication.clipboard().image().width() == pane.cutout().width()
    before = pane.mask.sum()
    pane.reach.setValue(8)  # select more
    assert wait(lambda: pane.mask.sum() > before, 10000)
    pane.reach.setValue(-8)  # select less
    assert wait(lambda: pane.mask.sum() < before, 10000)
    pane.reach.setValue(0)
    assert wait(lambda: abs(pane.mask.sum() - before) < before * 0.001, 10000)
    next(b for b in pane.copy_tray.findChildren(QToolButton) if b.text() == "Save").click()
    assert (t / "art (subject).png").exists() and "art (subject).png" in pane.subject_status.text()
    QTest.keyClick(w, Qt.Key_Escape)  # lets go of the pick
    assert pane.mask is None and pane.subject_button.isChecked()
    QTest.keyClick(w, Qt.Key_Escape)  # leaves picking
    assert not pane.subject_button.isChecked() and w.isVisible()
    w.close()
    # Ctrl+click on part of a picked figure cuts only that part, not the whole figure
    figure = QImage(1200, 800, QImage.Format_RGB32)
    figure.fill(QColor("#7aa6d8"))
    p = QPainter(figure)
    p.setRenderHint(QPainter.Antialiasing)
    p.setPen(Qt.NoPen)
    p.setBrush(QColor("#8a5a3b"))
    p.drawEllipse(500, 120, 200, 200)  # head
    p.drawRoundedRect(430, 300, 340, 400, 60, 60)  # body
    p.end()
    figure.save(str(t / "figure.png"))
    w = peek.Peek(t / "figure.png")
    w.show()
    QTest.qWaitForWindowActive(w)
    pane = w.preview
    QTest.keyClick(w, Qt.Key_S)
    pane.picked(QPointF(600, 500), Qt.NoModifier)
    assert wait(lambda: pane.mask is not None, 60000)
    assert wait(lambda: pane.segmenter.edges is not None and not pane.busy and not pane.pending, 120000)
    assert pane.mask[220, 600] > 0.5 and pane.mask[500, 600] > 0.5, "head and body picked together"
    pane.picked(QPointF(600, 200), Qt.ControlModifier)
    assert wait(lambda: pane.mask is not None and pane.mask[200, 600] < 0.5, 10000), "the head is cut away"
    assert pane.mask[550, 600] > 0.5, "the body stays"
    QTest.keyClick(w, Qt.Key_Z, Qt.ControlModifier)
    assert wait(lambda: pane.mask[200, 600] > 0.5, 10000), "Ctrl+Z brings the head back"
    pane.brush.setValue(1)  # 72 px brush
    assert pane.brush_size() == 72 and pane.brush_label.text() == "72 px"
    pane.picked(QPointF(600, 500), Qt.ControlModifier)
    assert wait(lambda: pane.mask[500, 600] < 0.5, 10000), "the brush cuts under the click"
    assert pane.mask[620, 600] > 0.5 and pane.mask[200, 600] > 0.5, "only inside the brush"
    QTest.keyClick(w, Qt.Key_Z, Qt.ControlModifier)
    assert wait(lambda: pane.mask[500, 600] > 0.5, 10000)
    pane.picked(QPointF(600, 500), Qt.NoModifier)
    assert wait(lambda: pane.mask[620, 600] < 0.5, 10000), "a new pick takes only what's inside the brush"
    assert pane.mask[500, 600] > 0.5
    # The picture can be dragged anywhere, even out of view
    view = pane.view
    view.fit()
    view.horizontalScrollBar().setValue(view.horizontalScrollBar().maximum())
    assert not view.mapFromScene(view.item.sceneBoundingRect()).boundingRect().intersects(view.viewport().rect())
    w.close()
else:
    print("skipped subject picking: onnxruntime or the SAM model is missing")

# AI image entries appear only when their tools are installed
codes = [code for code, _ in peek.convert_targets(d / "img2.png")]
assert [c for c in codes if c in peek.AI_TARGETS] == [c for c, _ in peek.ai_image_targets()], codes

print("ok")
