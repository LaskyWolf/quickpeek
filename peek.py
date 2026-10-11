#!/usr/bin/env python3
"""peek: quick file previewer (GNOME Sushi clone) for KDE. Usage: peek FILE | peek --quit"""
import configparser
import datetime
import grp
import gzip
import importlib.util
import hashlib
import json
import mmap
import os
import pwd
import re
import shutil
import sqlite3
import stat
import struct
import subprocess
import sys
import tarfile
import threading
import time
import types
import urllib.request
import zipfile
from compression import zstd
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from xml.etree import ElementTree
from xml.sax.saxutils import escape

import shiboken6
from PySide6.QtCore import (QAbstractAnimation, QEasingCurve, QVariantAnimation, QFile, Q_ARG, Qt, QCollator, QDate, QEvent, QLibraryInfo, QLocale, QLineF, QMarginsF, QMimeData, QMetaObject, QMimeDatabase, QModelIndex, QObject,
                            QPoint, QPointF, QProcess, QRectF, QSettings, QSize, QSizeF, QStandardPaths, QTimer, QUrl, Signal)
from PySide6.QtGui import (QTextCursor, QColor, QDesktopServices, QDrag, QFont, QFontDatabase, QIcon, QImage, QImageReader, QImageWriter, QKeySequence,
                           QMovie, QPageSize, QPainter, QPainterPath, QPalette, QPen, QPdfWriter, QPixmap, QShortcut, QCursor, QFontMetrics, QTextDocument, QTransform)
from PySide6.QtMultimedia import QAudioOutput, QMediaMetaData, QMediaPlayer, QtAudio
from PySide6.QtMultimediaWidgets import QVideoWidget
from PySide6.QtNetwork import QLocalServer, QLocalSocket
from PySide6.QtPdf import QPdfDocumentRenderOptions, QPdfDocument, QPdfSearchModel
from PySide6.QtPdfWidgets import QPdfView
from PySide6.QtQuickWidgets import QQuickWidget
from PySide6.QtWidgets import (QGraphicsBlurEffect, QInputDialog, QMenu, QMessageBox, QSlider, QStyleOptionSlider, QTextEdit, QAbstractItemView, QAbstractScrollArea, QApplication, QButtonGroup, QComboBox,
                               QFrame, QGraphicsDropShadowEffect, QGraphicsEllipseItem, QGraphicsItem,
                               QGraphicsPathItem, QGraphicsPixmapItem, QGraphicsRectItem, QGraphicsTextItem, QGraphicsScene, QGraphicsView,
                               QHBoxLayout, QLabel, QLineEdit, QListWidget, QListWidgetItem, QPlainTextEdit,
                               QProgressBar, QPushButton, QScrollArea, QStackedWidget, QStyle, QTableWidget, QTableWidgetItem, QTabWidget, QTextBrowser,
                               QToolButton, QTreeWidget, QTreeWidgetItem, QVBoxLayout, QWidget)

# Optional libraries: each feature falls back to a simpler preview without its library
try:
    import openpyxl
    from openpyxl.utils import get_column_letter
except ImportError:
    openpyxl = None
try:
    from pygments import highlight
    from pygments.formatters import HtmlFormatter
    from pygments.lexers import TextLexer, get_lexer_for_filename
    from pygments.styles import get_all_styles, get_style_by_name
    from pygments.util import ClassNotFound
except ImportError:
    highlight = None
try:
    from PIL import ExifTags, Image
except ImportError:
    Image = None
try:
    import pymupdf
except ImportError:
    pymupdf = None

# GStreamer plays short sounds fine; Qt's default FFmpeg backend freezes on them (Qt 6.12)
if (Path(QLibraryInfo.path(QLibraryInfo.PluginsPath)) / "multimedia/libgstreamermediaplugin.so").exists():
    os.environ.setdefault("QT_MEDIA_BACKEND", "gstreamer")
SOCKET = f"peek-{os.getuid()}"  # one peek window per user; later launches hand their file to it

TEXT_LIMIT = 1_000_000  # bytes shown for text files, so a 2 GB log doesn't freeze the window
CODE_LIMIT = 200_000  # characters highlighted; bigger files show as plain text (pygments is pure Python)
BLEND_HEAD = 4_000_000  # the .blend thumbnail sits near the start; never decompress a whole scene
THUMB_CACHE = Path.home() / ".cache/thumbnails"
CACHE = Path.home() / ".cache/peek"  # converted .glb/.pdf files, keyed by path + mtime + size
CACHE_LIMIT = 500 * 1024 ** 2  # oldest converted files are deleted beyond this
VIEW_PROPS = Path.home() / ".local/share/dolphin/view_properties"
MODELS = {".glb", ".gltf", ".obj", ".fbx", ".stl", ".ply", ".dae", ".3ds"}  # Qt's assimp importer
FONTS = {".ttf", ".otf", ".ttc", ".otc", ".woff", ".woff2"}
X2T = Path("/opt/onlyoffice/desktopeditors/converter/x2t")  # OnlyOffice's document converter
X2T_FONTS = Path.home() / ".local/share/onlyoffice/desktopeditors/data/fonts"  # made on first launch
X2T_PDF = 513  # x2t's format code for PDF
SHEETS = {".xlsx", ".xlsm"}  # opened directly by openpyxl
OFFICE = {".docx", ".xlsx", ".pptx", ".odt", ".ods", ".odp", ".docm", ".xlsm", ".doc", ".xls", ".ppt", ".rtf"}
EBOOKS = {".epub", ".fb2", ".djvu", ".xps", ".oxps"}  # x2t turns these into PDF too
MAX_ROWS, MAX_COLS = 5000, 200  # ponytail: hard cap keeps huge sheets responsive; lazy model if it bites
MAX_ENTRIES = 20_000  # archive entries listed
FOLDER_ITEMS = 200  # folder preview icons
ARCHIVES = {"application/zip", "application/x-tar", "application/x-compressed-tar", "application/x-bzip2-compressed-tar",
            "application/x-bzip-compressed-tar", "application/x-xz-compressed-tar", "application/x-zstd-compressed-tar",
            "application/x-lzma-compressed-tar", "application/x-7z-compressed", "application/vnd.rar",
            "application/x-rar", "application/x-iso9660-image", "application/x-cd-image", "application/java-archive",
            "application/vnd.android.package-archive", "application/x-cpio", "application/gzip", "application/x-xz",
            "application/x-bzip2", "application/zstd"}

VIEWER_QML = """
import QtQuick
import QtQuick3D
import QtQuick3D.Helpers
import QtQuick3D.AssetUtils

Item {
    id: root
    property url source
    property color background
    property color textColor
    property bool flying: false
    property bool wireframe: false
    property bool spin: false
    property bool grid: true
    property real gridStep: 1
    property url scene  // balsam's QML version of the model: same look, plus its animations
    property var timelines: []
    signal animationsReady(var names)

    onFlyingChanged: {
        if (flying) {  // start flying from wherever the orbit camera is looking
            flyCam.position = camera.scenePosition
            flyCam.rotation = camera.sceneRotation
            wasd.forceActiveFocus()
        } else {  // land: orbit around the point the fly camera was looking at, from the same spot
            const distance = camera.z
            orbit.position = flyCam.position.plus(flyCam.forward.times(distance))
            orbit.rotation = flyCam.rotation
            camera.position = Qt.vector3d(0, 0, distance)
        }
    }

    View3D {
        id: view
        anchors.fill: parent
        environment: SceneEnvironment {
            backgroundMode: SceneEnvironment.Color
            clearColor: root.background
            antialiasingMode: SceneEnvironment.MSAA
            antialiasingQuality: SceneEnvironment.High
            debugSettings: DebugSettings { wireframeEnabled: root.wireframe }
            InfiniteGrid { visible: root.grid; gridInterval: root.gridStep }
        }
        camera: root.flying ? flyCam : camera
        Node {
            id: orbit; objectName: "orbit"  // the controller rotates this node; the camera rides on it
            PerspectiveCamera {
                id: camera; objectName: "camera"
                DirectionalLight {  // headlight
                    visible: !root.flying
                    brightness: 1.0
                    ambientColor: Qt.rgba(0.35, 0.35, 0.35, 1)
                }
            }
        }
        PerspectiveCamera {
            id: flyCam; objectName: "flyCam"
            clipNear: camera.clipNear
            clipFar: camera.clipFar
            DirectionalLight { visible: root.flying; brightness: 1.0; ambientColor: Qt.rgba(0.35, 0.35, 0.35, 1) }
        }
        DirectionalLight { eulerRotation: Qt.vector3d(-50, 40, 0); brightness: 0.5 }
        RuntimeLoader {
            id: model
            source: root.source
            visible: scene.status !== Loader3D.Ready
            onBoundsChanged: root.fit()
        }
        Loader3D {
            id: scene
            source: root.scene
            onLoaded: root.collectAnimations()
        }
    }
    FrameAnimation {
        running: root.spin && !root.flying
        onTriggered: orbit.eulerRotation.y += frameTime * 30  // degrees per second
    }
    OrbitCameraController { anchors.fill: view; origin: orbit; camera: camera; enabled: !root.flying }
    WasdController { id: wasd; anchors.fill: view; controlledObject: flyCam; enabled: root.flying }
    Text {
        anchors { bottom: parent.bottom; horizontalCenter: parent.horizontalCenter; bottomMargin: 8 }
        color: root.textColor
        opacity: 0.7
        text: root.flying
              ? "W A S D / arrows: move  ·  R / F: up / down  ·  Shift: faster  ·  drag: look  ·  Esc: stop flying"
              : "Drag: rotate  ·  Ctrl+drag: pan  ·  Wheel: zoom"
    }
    Text {
        anchors.centerIn: parent
        visible: model.status === RuntimeLoader.Error
        text: model.errorString
        color: root.textColor
    }

    // balsam makes one Timeline per animation and starts them all; keep only the first running
    function collectAnimations() {
        const found = []
        const walk = node => {
            for (const child of node.data) {
                if (child.endFrame !== undefined && child.animations !== undefined)
                    found.push(child)
                else if (child.data !== undefined)
                    walk(child)
            }
        }
        walk(scene.item)
        timelines = found
        playAnimation(0, true)
        animationsReady(found.map((t, i) => t.objectName || "Animation " + (i + 1)))
    }

    function playAnimation(index, playing) {
        timelines.forEach((t, i) => {
            t.enabled = i === index
            t.animations[0].running = i === index
            t.animations[0].paused = i === index && !playing
        })
    }

    // Center the model and back the camera off far enough to see all of it
    function fit() {
        const min = model.bounds.minimum, max = model.bounds.maximum
        const radius = max.minus(min).length() / 2 || 1
        orbit.position = min.plus(max).times(0.5)
        orbit.eulerRotation = Qt.vector3d(-20, 30, 0)
        camera.position = Qt.vector3d(0, 0, radius * 2.6)
        camera.clipNear = radius / 100
        camera.clipFar = radius * 100
        wasd.speed = radius / 150  // fly speed relative to model size
        root.gridStep = Math.pow(10, Math.floor(Math.log10(radius)))  // 1 m lines for a 1-5 m model
    }
}
"""


# ---------------------------------------------------------------- background work

class _Bridge(QObject):
    done = Signal(object)


# Separate lanes, so slow work never queues in front of quick work:
_pool = ThreadPoolExecutor(max_workers=4)  # quick reads: details, stats, archive lists
_preload_pool = ThreadPoolExecutor(max_workers=1)  # decoding the neighbouring images
_convert_pool = ThreadPoolExecutor(max_workers=2)  # conversions, which can take minutes
_bridge = None


def in_background(fn, owner, on_done, pool=_pool):
    """Run fn() on a worker thread, then on_done(result) on the GUI thread if `owner` still exists."""
    global _bridge
    if _bridge is None:  # created on first use, on the GUI thread: its signal delivers there
        _bridge = _Bridge()
        _bridge.done.connect(lambda call: call())

    def finished(future):
        if future.exception() is None:
            _bridge.done.emit(lambda: shiboken6.isValid(owner) and on_done(future.result()))

    pool.submit(fn).add_done_callback(finished)


def decode_image(path):
    reader = QImageReader(str(path))
    reader.setAutoTransform(True)  # honour EXIF rotation
    return reader.read()


# ---------------------------------------------------------------- cache and conversions

def prune_cache(clean_parts=False):
    """Keep converted files under CACHE_LIMIT, newest first. Leftover partial exports go at startup."""
    if not CACHE.exists():
        return
    if clean_parts:
        for part in CACHE.glob("*.part.*"):
            shutil.rmtree(part) if part.is_dir() else part.unlink(missing_ok=True)
        shutil.rmtree(CACHE / "tmp", ignore_errors=True)
    files = sorted((p for p in CACHE.iterdir() if p.suffix in {".glb", ".pdf", ".scene", ".json"} and ".part." not in p.name),
                   key=lambda p: p.stat().st_mtime, reverse=True)
    total = 0
    for p in files:
        total += sum(f.stat().st_size for f in p.rglob("*") if f.is_file()) if p.is_dir() else p.stat().st_size
        if total > CACHE_LIMIT:
            shutil.rmtree(p) if p.is_dir() else p.unlink(missing_ok=True)


def blender_command():
    if shutil.which("blender"):
        return [shutil.which("blender")]
    for root in (Path("/var/lib/flatpak"), Path.home() / ".local/share/flatpak"):
        if (root / "app/org.blender.Blender").exists():
            return ["flatpak", "run", "org.blender.Blender"]
    return None


def cache_key(path):
    """Changes whenever the file does, so an edited file is never shown from a stale cache."""
    st = path.stat()
    return hashlib.md5(f"{path}:{st.st_mtime_ns}:{st.st_size}".encode()).hexdigest()


def animated_scene(path):
    """balsam's QML version of a model, which keeps its animations, or None if it has none.
    Runs on a worker thread; cached in CACHE/<key>.scene/."""
    suffix = path.suffix.lower()
    if suffix in {".glb", ".gltf"}:
        stats = model_stats(path) or ""
        if "animation" not in stats:
            return None
    elif suffix not in {".fbx", ".dae"}:  # formats that can carry animations; others never do
        return None
    balsam = Path(QLibraryInfo.path(QLibraryInfo.BinariesPath)) / "balsam"
    if not balsam.exists():
        return None
    out = CACHE / f"{cache_key(path)}.scene"
    if not out.exists():
        part = CACHE / f"{cache_key(path)}.part.scene"
        shutil.rmtree(part, ignore_errors=True)
        run([str(balsam), "-o", str(part), str(path)], timeout=300)
        part.rename(out)
    out.touch()  # recently used: the cache cleanup keeps it longer
    qml = next(out.glob("*.qml"), None)
    return qml if qml and "Timeline {" in qml.read_text(errors="replace") else None


def conversion(path):
    """(command, temporary output, final output, x2t task) that turns `path` into a .glb or .pdf, or None."""
    suffix = path.suffix.lower()
    key = cache_key(path)
    if suffix == ".blend" and (blender := blender_command()):
        out = CACHE / f"{key}.glb"
        part = CACHE / f"{key}.part.glb"  # renamed when done, so a killed export never looks finished
        script = f"import bpy; bpy.ops.export_scene.gltf(filepath={str(part)!r}, export_format='GLB', export_apply=True)"
        return blender + ["-b", str(path), "--python-expr", script], part, out, None
    if (suffix in OFFICE | EBOOKS and not (openpyxl and suffix in SHEETS)
            and X2T.exists() and X2T_FONTS.exists()):
        out = CACHE / f"{key}.pdf"
        part = CACHE / f"{key}.part.pdf"
        task = f"""<?xml version="1.0" encoding="utf-8"?>
<TaskQueueDataConvert>
  <m_sFileFrom>{escape(str(path))}</m_sFileFrom>
  <m_sFileTo>{escape(str(part))}</m_sFileTo>
  <m_nFormatTo>{X2T_PDF}</m_nFormatTo>
  <m_sTempDir>{escape(str(CACHE / "tmp"))}</m_sTempDir>
  <m_sFontDir>{escape(str(X2T_FONTS))}</m_sFontDir>
  <m_sAllFontsPath>{escape(str(X2T_FONTS / "AllFonts.js"))}</m_sAllFontsPath>
  <m_bIsNoBase64>true</m_bIsNoBase64>
</TaskQueueDataConvert>"""
        return [str(X2T), str(out.with_suffix(".xml"))], part, out, task
    return None


# ---------------------------------------------------------------- format readers

def blend_thumbnail(path):
    """Blender stores a thumbnail in a 'TEST' block: width, height, then bottom-up RGBA pixels."""
    with open(path, "rb") as f:
        head = f.read(4)
        f.seek(0)
        if head == b"\x28\xb5\x2f\xfd":  # zstd, Blender 3.0+ "Compress" option
            data = zstd.ZstdDecompressor().decompress(f.read(), max_length=BLEND_HEAD)
        elif head[:2] == b"\x1f\x8b":  # gzip, older Blender
            data = gzip.GzipFile(fileobj=f).read(BLEND_HEAD)
        else:
            data = f.read(BLEND_HEAD)
    if not data.startswith(b"BLENDER"):
        return None
    if data[7:9] == b"17":  # Blender 5.0+ header: "BLENDER17-01v0500", 64-bit block headers
        pos, bhead, size_of = 17, struct.Struct("<4siQqq"), lambda h: h[3]
    else:  # "BLENDER-v405": '-' = 8-byte pointers, 'v' = little endian
        ptr = "Q" if data[7:8] == b"-" else "I"
        bhead = struct.Struct(("<" if data[8:9] == b"v" else ">") + "4si" + ptr + "ii")
        pos, size_of = 12, lambda h: h[1]
    while pos + bhead.size <= len(data):
        h = bhead.unpack_from(data, pos)
        pos += bhead.size
        if h[0] == b"TEST":
            w, h_ = struct.unpack_from(bhead.format[0] + "ii", data, pos)
            pixels = data[pos + 8:pos + 8 + w * h_ * 4]
            return QImage(pixels, w, h_, QImage.Format_RGBA8888).flipped(Qt.Vertical)
        if h[0] == b"ENDB":
            return None
        pos += size_of(h)
    return None


def clip_preview(path):
    """Clip Studio files embed an SQLite database whose CanvasPreview table holds a PNG."""
    data = Path(path).read_bytes()
    chunk = data.find(b"CHNKSQLi")
    if chunk < 0:
        return None
    size = struct.unpack_from(">Q", data, chunk + 8)[0]
    db = sqlite3.connect(":memory:")
    db.deserialize(data[chunk + 16:chunk + 16 + size])
    row = db.execute("SELECT ImageData FROM CanvasPreview LIMIT 1").fetchone()
    return QImage.fromData(row[0]) if row else None


def office_text(path):
    """Plain text of Office/OpenDocument files: paragraphs on lines, table cells split by tabs."""
    z = zipfile.ZipFile(path)
    names = z.namelist()
    if "content.xml" in names:  # OpenDocument: odt, ods, odp
        parts, para, cell, row = ["content.xml"], {"p", "h"}, "table-cell", "table-row"
    elif "word/document.xml" in names:
        parts, para, cell, row = ["word/document.xml"], {"p"}, "tc", "tr"
    elif "ppt/presentation.xml" in names:
        slides = [n for n in names if n.startswith("ppt/slides/slide") and n.endswith(".xml")]
        parts = sorted(slides, key=lambda n: int(n[16:-4]))
        para, cell, row = {"p"}, "tc", "tr"
    elif "xl/workbook.xml" in names:
        return xlsx_text(z, names)
    else:
        return None
    out = []

    def walk(el):
        tag = el.tag.rsplit("}", 1)[-1]  # drop the XML namespace: "{urn:...}p" -> "p"
        if tag == row:
            cells = ["".join(c.itertext()) for c in el if c.tag.rsplit("}", 1)[-1] == cell]
            out.append("\t".join(cells).rstrip("\t"))
        elif tag in para:
            out.append("".join(el.itertext()))
        else:
            for child in el:
                walk(child)

    for part in parts:
        walk(ElementTree.fromstring(z.read(part)))
        out.append("")  # blank line between slides
    return "\n".join(out)


def xlsx_text(z, names):
    """First worksheet as tab-separated rows. Shared strings are stored once and referenced by index."""
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    shared = []
    if "xl/sharedStrings.xml" in names:
        shared = ["".join(si.itertext()) for si in ElementTree.fromstring(z.read("xl/sharedStrings.xml"))]
    sheet = ElementTree.fromstring(z.read("xl/worksheets/sheet1.xml"))
    lines = []
    for r in sheet.iter(ns + "row"):
        cells = []
        for c in r.iter(ns + "c"):
            v = c.find(ns + "v")
            text = v.text if v is not None else "".join(c.itertext())
            cells.append(shared[int(text)] if c.get("t") == "s" else text or "")
        lines.append("\t".join(cells))
    return "\n".join(lines)


def cached_thumbnail(path, sizes=("xx-large", "x-large", "large", "normal")):
    """Thumbnail Dolphin already generated (freedesktop spec: md5 of the file URI)."""
    key = hashlib.md5(QUrl.fromLocalFile(str(path)).toEncoded().data()).hexdigest() + ".png"
    for size in sizes:
        thumb = THUMB_CACHE / size / key
        if thumb.exists():
            return QImage(str(thumb))
    return None


def model_stats(path):
    """'12.340 vertices · 20.000 triangles · ...' for formats whose headers say so, else None."""
    suffix = path.suffix.lower()
    n = QLocale().toString
    parts = []
    if suffix in {".glb", ".gltf"}:
        if suffix == ".glb":  # 12-byte header, then the JSON chunk: length, type, data
            with open(path, "rb") as f:
                f.seek(12)
                length, _ = struct.unpack("<II", f.read(8))
                gltf = json.loads(f.read(length))
        else:
            gltf = json.loads(path.read_text())
        accessors = gltf.get("accessors", [])
        vertices = triangles = 0
        for mesh in gltf.get("meshes", []):
            for prim in mesh.get("primitives", []):
                count = accessors[prim["attributes"]["POSITION"]]["count"]
                vertices += count
                if prim.get("mode", 4) == 4:  # 4 = triangles
                    triangles += (accessors[prim["indices"]]["count"] if "indices" in prim else count) // 3
        parts = [f"{n(vertices)} vertices", f"{n(triangles)} triangles",
                 f"{len(gltf.get('materials', []))} materials", f"{len(gltf.get('images', []))} textures"]
        if animations := len(gltf.get("animations", [])):
            parts.append(f"{animations} animation{'s' if animations > 1 else ''}")
    elif suffix == ".obj" and path.stat().st_size < 200_000_000:
        vertices = triangles = 0
        with open(path, "rb") as f:
            for line in f:
                if line.startswith(b"v "):
                    vertices += 1
                elif line.startswith(b"f "):
                    triangles += len(line.split()) - 3  # an n-sided face is n - 2 triangles
        parts = [f"{n(vertices)} vertices", f"{n(triangles)} triangles"]
    elif suffix == ".stl":
        with open(path, "rb") as f:
            head = f.read(84)
        count = struct.unpack_from("<I", head, 80)[0] if len(head) == 84 else 0
        if path.stat().st_size == 84 + 50 * count:  # binary STL: 50 bytes per triangle
            parts = [f"{n(count)} triangles"]
        else:
            parts = [f"{n(path.read_bytes().count(b'facet normal'))} triangles"]
    elif suffix == ".ply":
        with open(path, "rb") as f:
            for line in f:
                if line.startswith(b"element vertex"):
                    parts.append(f"{n(int(line.split()[2]))} vertices")
                elif line.startswith(b"element face"):
                    parts.append(f"{n(int(line.split()[2]))} faces")
                elif line.strip() == b"end_header":
                    break
    return "  ·  ".join(parts) or None


def list_archive(path):
    """[(name, size or None, modified, is_dir)] from an archive. Python reads zip/tar; 7z the rest."""
    entries = []
    try:
        if zipfile.is_zipfile(path):
            for info in zipfile.ZipFile(path).infolist()[:MAX_ENTRIES]:
                entries.append((info.filename.rstrip("/"), info.file_size,
                                datetime.datetime(*info.date_time), info.is_dir()))
            return entries
        with tarfile.open(path) as tar:  # reads gz/bz2/xz/zst tars
            for member in tar:
                entries.append((member.name, member.size, datetime.datetime.fromtimestamp(member.mtime),
                                member.isdir()))
                if len(entries) >= MAX_ENTRIES:
                    break
            return entries
    except (tarfile.TarError, zipfile.BadZipFile, OSError, ValueError, EOFError):
        entries = []
    if not shutil.which("7z"):
        return None
    # -slt prints "Key = value" blocks; -p with no password stops 7z asking for one
    out = subprocess.run(["7z", "l", "-slt", "-p", "--", str(path)], capture_output=True, text=True,
                         stdin=subprocess.DEVNULL, timeout=30).stdout
    if "----------" not in out:
        return None
    for block in out.split("----------", 1)[1].split("\n\n"):
        fields = dict(line.split(" = ", 1) for line in block.strip().splitlines() if " = " in line)
        if "Path" in fields:
            modified = fields.get("Modified", "")[:19]
            entries.append((fields["Path"], int(fields["Size"]) if fields.get("Size") else None,
                            datetime.datetime.fromisoformat(modified) if modified else None,
                            fields.get("Folder") == "+" or "D" in fields.get("Attributes", "")[:1]))
        if len(entries) >= MAX_ENTRIES:
            break
    return entries


def folder_summary(folder, deadline=3.0):
    """(files, folders, bytes, complete) for everything under `folder`, giving up after `deadline` s."""
    files = folders = size = 0
    stop = time.monotonic() + deadline
    for root, dirs, names in os.walk(folder):
        folders += len(dirs)
        for name in names:
            files += 1
            try:
                size += os.lstat(os.path.join(root, name)).st_size
            except OSError:
                pass
        if time.monotonic() > stop:
            return files, folders, size, False
    return files, folders, size, True


def code_html(text, filename):
    """Syntax-highlighted HTML and its background color, or None when pygments doesn't know the language."""
    if not highlight or len(text) > CODE_LIMIT:
        return None
    try:
        lexer = get_lexer_for_filename(filename, stripnl=False)
    except ClassNotFound:
        return None
    if isinstance(lexer, TextLexer):
        return None
    dark = QApplication.palette().window().color().lightness() < 128
    style = ("github-dark" if "github-dark" in get_all_styles() else "monokai") if dark else "default"
    html = highlight(text, lexer, HtmlFormatter(noclasses=True, style=style, nobackground=True))
    return html, get_style_by_name(style).background_color


# ---------------------------------------------------------------- Dolphin's sort order

def dolphin_view_settings(folder):
    """Sorting Dolphin uses for `folder`: per-folder settings first, then the global default."""
    settings = {"role": "text", "descending": False, "folders_first": True, "hidden_last": False,
                "show_hidden": False}
    dolphinrc = configparser.ConfigParser(interpolation=None, strict=False)
    dolphinrc.read(Path.home() / ".config/dolphinrc")
    texts = []
    if dolphinrc.get("General", "GlobalViewProps", fallback="false") != "true":
        try:  # Dolphin 24+ stores per-folder view settings as an extended attribute on the folder
            texts.append(os.getxattr(folder, "user.kde.fm.viewproperties#1").decode())
        except OSError:
            pass
        for file in (Path(folder) / ".directory", VIEW_PROPS / str(folder).lstrip("/") / ".directory"):
            if file.is_file():
                texts.append(file.read_text(errors="replace"))
    if (VIEW_PROPS / "global/.directory").is_file():
        texts.append((VIEW_PROPS / "global/.directory").read_text(errors="replace"))
    for text in texts:
        parser = configparser.ConfigParser(interpolation=None, strict=False)
        try:
            parser.read_string(text)
        except configparser.Error:
            continue
        if parser.has_section("Dolphin"):
            section = parser["Dolphin"]  # keys are case-insensitive here
            settings["role"] = section.get("SortRole", settings["role"])
            settings["descending"] = section.get("SortOrder", "0") == "1"
            settings["folders_first"] = section.get("SortFoldersFirst", "true") == "true"
            settings["hidden_last"] = section.get("SortHiddenLast", "false") == "true"
            settings["show_hidden"] = section.get("HiddenFilesShown", "false") == "true"
            break
    return settings


def sorted_like_dolphin(folder, include=None):
    """Entries of `folder` in the order Dolphin shows them. `include` stays even if hidden."""
    s = dolphin_view_settings(folder)
    items = [p for p in Path(folder).iterdir() if s["show_hidden"] or not p.name.startswith(".") or p == include]
    collator = QCollator()
    collator.setNumericMode(True)  # "img2" before "img10"
    collator.setCaseSensitivity(Qt.CaseInsensitive)
    mimes = QMimeDatabase()

    def stat(p, field):
        try:
            st = p.stat()
        except OSError:
            return 0
        return getattr(st, field, st.st_mtime)  # no birth time on some filesystems: use modified

    role = {
        "size": lambda p: 0 if p.is_dir() else stat(p, "st_size"),
        "modificationtime": lambda p: stat(p, "st_mtime"),
        "accesstime": lambda p: stat(p, "st_atime"),
        "creationtime": lambda p: stat(p, "st_birthtime"),
        "type": lambda p: collator.sortKey(mimes.mimeTypeForFile(str(p), QMimeDatabase.MatchExtension).comment()),
        "extension": lambda p: p.suffix.lower(),
    }.get(s["role"].lower(), lambda p: 0)
    items.sort(key=lambda p: (role(p), collator.sortKey(p.name)), reverse=s["descending"])
    # Grouping never flips with the sort direction; Python's sort is stable, so the order above survives
    items.sort(key=lambda p: (s["hidden_last"] and p.name.startswith("."), s["folders_first"] and not p.is_dir()))
    return items


def siblings(path):
    return sorted_like_dolphin(Path(path).parent, include=Path(path))


# ---------------------------------------------------------------- small shared widgets

class Tray(QWidget):
    """Related buttons on one rounded background, so a toolbar reads as a few groups, not a row of icons."""

    def __init__(self, *widgets):
        super().__init__()
        layout = QHBoxLayout(self)
        layout.setContentsMargins(3, 3, 3, 3)
        layout.setSpacing(1)
        for w in widgets:
            layout.addWidget(w)

    def paintEvent(self, event):
        painter = QPainter(self)
        painter.setRenderHint(QPainter.Antialiasing)
        fill = QColor(self.palette().button().color())
        fill.setAlpha(150)
        edge = QColor(self.palette().mid().color())
        edge.setAlpha(90)
        painter.setPen(QPen(edge, 1))
        painter.setBrush(fill)
        painter.drawRoundedRect(QRectF(self.rect()).adjusted(0.5, 0.5, -0.5, -0.5), 7, 7)


PANEL_BUTTON_ROOM = 44  # width kept free at the top right for the Details and Tools button


def toolbar(first, *items):
    """`first` takes the free space on the left; each tuple after it becomes a Tray."""
    row = QHBoxLayout()
    row.setSpacing(8)
    row.addWidget(first, 1)
    for item in items:
        row.addWidget(Tray(*item) if isinstance(item, tuple) else item)
    row.addSpacing(PANEL_BUTTON_ROOM)  # the window's panel button floats over this corner
    return row


def tool_button(icon, tip, slot, text=None, checkable=False, checked=False):
    button = QToolButton(icon=QIcon.fromTheme(icon), toolTip=tip, autoRaise=True, checkable=checkable)
    if text:
        button.setText(text)
        button.setToolButtonStyle(Qt.ToolButtonTextBesideIcon)
    button.setChecked(checked)
    (button.toggled if checkable else button.clicked).connect(slot)
    return button


def vbox(*items, margins=0):
    box = QWidget()
    layout = QVBoxLayout(box)
    layout.setContentsMargins(margins, margins, margins, margins)
    for item in items:
        if isinstance(item, QWidget):
            layout.addWidget(item, 1 if isinstance(item, (QAbstractItemView, QGraphicsView, QPlainTextEdit,
                                                          QVideoWidget, QQuickWidget, QPdfView)) else 0)
        else:
            layout.addLayout(item)
    return box


class SearchBar(QFrame):
    """Find box floating over the top-right corner of `host`. Hidden until Ctrl+F; Escape hides it.
    Typing searches, Enter / Shift+Enter jump to the next / previous match."""
    changed = Signal(str)
    step = Signal(int)

    def __init__(self, host, owner):
        super().__init__(host)
        self.host = host
        self.setFrameShape(QFrame.StyledPanel)
        self.setAutoFillBackground(True)
        shadow = QGraphicsDropShadowEffect(self, blurRadius=18, offset=QPointF(0, 3))
        shadow.setColor(QColor(0, 0, 0, 90))
        self.setGraphicsEffect(shadow)
        self.edit = QLineEdit(placeholderText="Find", clearButtonEnabled=True)
        self.edit.setProperty("peekFocus", True)  # the one widget allowed to take keyboard focus
        self.edit.setMinimumWidth(220)
        self.edit.addAction(QIcon.fromTheme("edit-find"), QLineEdit.LeadingPosition)
        self.edit.textChanged.connect(self.changed)
        self.edit.returnPressed.connect(lambda: self.step.emit(1))
        QShortcut("Shift+Return", self.edit, lambda: self.step.emit(-1), context=Qt.WidgetShortcut)
        self.label = QLabel()
        self.label.setMinimumWidth(70)
        self.label.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
        layout = QHBoxLayout(self)
        layout.setContentsMargins(6, 4, 4, 4)
        layout.setSpacing(2)
        layout.addWidget(self.edit)
        layout.addWidget(self.label)
        layout.addWidget(tool_button("go-up", "Previous match (Shift+Enter)", lambda: self.step.emit(-1)))
        layout.addWidget(tool_button("go-down", "Next match (Enter)", lambda: self.step.emit(1)))
        layout.addWidget(tool_button("window-close", "Close (Esc)", self.close_bar))
        self.hide()
        host.installEventFilter(self)  # follow the host's corner when the window resizes
        QShortcut(QKeySequence.Find, owner, self.open_bar)

    def button(self):
        """Magnifier for a toolbar, so the find box can be discovered without knowing Ctrl+F."""
        return tool_button("edit-find", "Find (Ctrl+F)", self.open_bar)

    def open_bar(self):
        self.show()
        self.raise_()
        self.place()
        self.edit.setFocus()
        self.edit.selectAll()

    def close_bar(self):
        self.edit.clear()  # clears the highlights too
        self.hide()
        self.window().setFocus()

    def place(self):
        self.adjustSize()
        right = self.host.width() - 12
        if isinstance(self.host, QAbstractScrollArea) and self.host.verticalScrollBar().isVisible():
            right -= self.host.verticalScrollBar().width()
        self.move(right - self.width(), 10)

    def eventFilter(self, obj, event):
        if event.type() == QEvent.Resize and self.isVisible():
            self.place()
        return False


class DragIcon(QLabel):
    """The file's icon in the bottom bar: drag it into another app to hand over the file."""

    def __init__(self):
        super().__init__(cursor=Qt.OpenHandCursor, toolTip="Drag the file into another app")
        self.path = None
        self.start = QPoint()

    def set_file(self, path, icon):
        self.path = path
        if icon.isNull():  # a type with no themed icon still needs something to grab
            icon = self.style().standardIcon(QStyle.StandardPixmap.SP_FileIcon)
        self.setPixmap(icon.pixmap(32))

    def mousePressEvent(self, event):
        self.start = event.position().toPoint()

    def mouseMoveEvent(self, event):
        if not (event.buttons() & Qt.LeftButton) or not self.path:
            return
        if (event.position().toPoint() - self.start).manhattanLength() < QApplication.startDragDistance():
            return
        data = QMimeData()
        data.setUrls([QUrl.fromLocalFile(str(self.path))])
        drag = QDrag(self)
        drag.setMimeData(data)
        drag.setPixmap(self.pixmap())
        drag.exec(Qt.CopyAction | Qt.LinkAction)  # never Move: dropping must not take the file away


def icon_label(mime):
    label = QLabel(alignment=Qt.AlignCenter)
    icon = QIcon.fromTheme(mime.iconName(), QIcon.fromTheme(mime.genericIconName()))
    label.setPixmap(icon.pixmap(128))
    return label


# ---------------------------------------------------------------- images

class ZoomView(QGraphicsView):
    """Image canvas: fits the window until you zoom; wheel zooms at the cursor, drag pans."""
    zoomed = Signal(float)

    def __init__(self):
        super().__init__()
        self.setScene(QGraphicsScene(self))
        self.item = QGraphicsPixmapItem()
        self.item.setTransformationMode(Qt.SmoothTransformation)
        self.scene().addItem(self.item)
        self.setDragMode(QGraphicsView.ScrollHandDrag)
        self.setTransformationAnchor(QGraphicsView.AnchorUnderMouse)
        self.setRenderHint(QPainter.SmoothPixmapTransform)
        self.setFrameShape(QGraphicsView.NoFrame)
        self.setBackgroundBrush(self.palette().window())
        self.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)  # the picture moves freely, even out of view
        self.setVerticalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        self.fitting = True
        self.picking = False  # clicks pick a subject instead of double-click zooming
        self.pressed_at = QPointF()
        self.brush = 0  # subject brush side in picture pixels, 0 = none
        pen = QPen(Qt.white, 1.5, Qt.DashLine)
        pen.setCosmetic(True)  # same thickness at any zoom
        self.brush_square = self.scene().addRect(QRectF(), pen)
        self.brush_square.setZValue(10)
        self.brush_square.hide()
        self.setMouseTracking(True)

    def show_brush(self, position=None):
        """The brush square under the mouse while picking a subject."""
        if position is None:
            position = self.mapFromGlobal(QCursor.pos())
        center = self.mapToScene(position)
        size = self.brush
        self.brush_square.setRect(center.x() - size / 2, center.y() - size / 2, size, size)
        self.brush_square.setVisible(bool(self.picking and size and self.viewport().rect().contains(position)))

    def mouseMoveEvent(self, event):
        super().mouseMoveEvent(event)
        if self.picking:
            self.show_brush(event.position().toPoint())

    def leaveEvent(self, event):
        super().leaveEvent(event)
        self.brush_square.hide()

    clicked = Signal(QPointF, object)  # scene point, keyboard modifiers: a click, not the end of a drag

    def mousePressEvent(self, event):
        self.pressed_at = event.position()
        super().mousePressEvent(event)

    def mouseReleaseEvent(self, event):
        super().mouseReleaseEvent(event)
        if event.button() == Qt.LeftButton and (event.position() - self.pressed_at).manhattanLength() < 5:
            self.clicked.emit(self.mapToScene(event.position().toPoint()), event.modifiers())

    def set_pixmap(self, pixmap):
        self.item.setPixmap(pixmap)
        room = 100 * max(pixmap.width(), pixmap.height(), 1)  # empty space all around: drag anywhere
        self.scene().setSceneRect(self.item.boundingRect().adjusted(-room, -room, room, room))
        if self.fitting:
            self.fit()

    def fit(self):
        self.fitting = True
        self.fitInView(self.item, Qt.KeepAspectRatio)
        self.zoomed.emit(self.transform().m11())

    def zoom(self, factor):
        self.fitting = False
        scale = min(max(self.transform().m11() * factor, 0.02), 50)
        self.setTransform(QTransform.fromScale(scale, scale))
        self.zoomed.emit(scale)

    def actual_size(self):
        self.fitting = False
        self.resetTransform()
        self.zoomed.emit(1.0)

    def wheelEvent(self, event):
        self.zoom(1.25 if event.angleDelta().y() > 0 else 1 / 1.25)

    def mouseDoubleClickEvent(self, event):
        if not self.picking:
            self.actual_size() if self.fitting else self.fit()

    def resizeEvent(self, event):
        super().resizeEvent(event)
        if self.fitting:
            self.fit()


def ocr_words(image, key=None):
    """Words tesseract finds in a QImage: [(text, (x, y, w, h), line id)]. Runs on a worker thread.
    Cached by file key, so each image is read once."""
    langs = [l for l in subprocess.run(["tesseract", "--list-langs"], capture_output=True, text=True)
             .stdout.splitlines()[1:] if l.strip() not in ("", "osd", "equ")] or ["eng"]  # line 1 is a header
    # The languages are part of the name: installing a new one re-reads images saved before it
    cached = CACHE / f"{key}.{'+'.join(langs)}.ocr2.json" if key else None  # ocr2: the photo-friendly settings
    if cached and cached.exists():
        cached.touch()
        return [(t, tuple(box), tuple(line)) for t, box, line in json.loads(cached.read_text())]
    (CACHE / "tmp").mkdir(parents=True, exist_ok=True)
    png = CACHE / "tmp" / f"ocr-{os.getpid()}-{time.monotonic_ns()}.png"
    # Tesseract wants letters ~20 px tall: enlarge small images (a phone photo of a card has 8 px text).
    # Boxes are divided back by `factor`, so they still match the shown image.
    longest = max(image.width(), image.height(), 1)
    factor = 1 if longest >= 1000 else min(3, 2200 / longest)
    gray = image.convertToFormat(QImage.Format_Grayscale8)
    if factor > 1:
        gray = gray.scaled(gray.size() * factor, Qt.KeepAspectRatio, Qt.SmoothTransformation)
    gray.save(str(png))
    try:
        # --psm 11 finds scattered words (photos, cards, signs) as well as plain pages;
        # thresholding_method=2 (Sauvola) cleans patterned and uneven backgrounds.
        # One thread: barely slower, and the rest of the PC stays responsive.
        result = subprocess.run(["tesseract", str(png), "stdout", "-l", "+".join(langs), "--psm", "11",
                                 "-c", "thresholding_method=2", "tsv"],
                                capture_output=True, text=True, timeout=120,
                                env={**os.environ, "OMP_THREAD_LIMIT": "1"})
    finally:
        png.unlink(missing_ok=True)
    words = []
    for row in result.stdout.splitlines()[1:]:
        cols = row.split("\t")
        # level 5 = a word; conf -1 = layout rows; low confidence is usually noise from pictures
        if len(cols) == 12 and cols[0] == "5" and cols[11].strip() and float(cols[10]) >= 40:
            box = tuple(round(int(v) / factor) for v in cols[6:10])
            words.append((cols[11], box, tuple(map(int, cols[2:5]))))
    if cached:
        cached.write_text(json.dumps(words))
    return words


def ocr_text(words):
    """Words back into lines and paragraphs."""
    lines, last = [], None
    for text, _, line in words:
        if line == last:
            lines[-1] += " " + text
        else:
            if last and line[0] != last[0]:  # new text block; tesseract's "paragraphs" split almost every line
                lines.append("")
            lines.append(text)
            last = line
    return "\n".join(lines)


class ImagePane(QWidget):
    """Zoomable image with a zoom toolbar, and text recognition (OCR) on request."""

    def __init__(self, source, path=None):
        super().__init__()
        self.path = Path(path) if path else None
        self.words = None  # None = not read yet
        self.boxes = []
        self.view = ZoomView()
        self.zoom_label = QLabel()
        self.view.zoomed.connect(lambda s: self.zoom_label.setText(f"{s * 100:.0f}%"))
        self.movie = None
        if isinstance(source, QImage):
            image = source
            self.view.set_pixmap(QPixmap.fromImage(image))
        elif QImageReader(str(source)).supportsAnimation():
            self.movie = QMovie(str(source), parent=self)
            self.movie.frameChanged.connect(lambda _: self.view.set_pixmap(self.movie.currentPixmap()))
            self.movie.start()
            image = self.movie.currentImage()
        else:
            image = decode_image(source)
            self.view.set_pixmap(QPixmap.fromImage(image))
        self.image = image

        self.ocr_status = QLabel()
        self.text_button = tool_button("edit-select-text", "Find the words in this image (T)", self.show_text,
                                       text="Text", checkable=True)
        self.text_button.setEnabled(bool(shutil.which("tesseract")))
        self.subject_button = tool_button("tools-wizard", "Pick a subject: click what to keep (S)", self.subject_mode,
                                          text="Subject", checkable=True)
        if not has_onnx():
            self.subject_button.setEnabled(False)
            self.subject_button.setToolTip("Install python-onnxruntime-cpu to pick subjects")
        self.subject_button.setEnabled(self.subject_button.isEnabled() and not self.movie)
        self.subject_status = QLabel(textFormat=Qt.RichText)
        self.subject_status.linkActivated.connect(show_in_folder)
        self.segmenter = None  # Segment Anything, once it has read the picture
        self.mask = None  # numpy 0..1: what's picked
        self.parts = []  # the picked objects, each with its own clicks (see Segmenter.part)
        self.logits = None  # all parts together, as log-odds
        self.layers = []  # dimming, glow and outline over the picture
        self.pending = []  # clicks waiting for the picture to be read, or for the previous click
        self.history = []  # earlier `parts`, for Ctrl+Z
        self.transition = None  # (animation, finish) while an add/remove/fade animation runs
        self.busy = False
        self.progress = Progress()
        self.ticker = QTimer(self, interval=250)
        self.ticker.timeout.connect(lambda: self.progress.text and self.subject_status.setText(self.progress.text))
        self.view.clicked.connect(self.picked)
        copy_subject = tool_button("edit-copy", "Copy the subject (Ctrl+C)", self.copy_subject, text="Copy")
        save_subject = tool_button("document-save", "Save the subject as a PNG next to the picture",
                                   self.save_subject, text="Save")
        save_subject.setEnabled(bool(self.path))
        self.reach = QSlider(Qt.Horizontal, minimum=-8, maximum=8, toolTip="Select less ⟷ more")
        self.reach.setFixedWidth(110)
        self.reach.valueChanged.connect(lambda _: self.refit.start())
        self.refit = QTimer(self, singleShot=True, interval=60)  # slider moves come fast: redo the mask once
        self.refit.timeout.connect(lambda: self.pending.append((0, 0, "reach", 0)) or self.next_pick())
        less, more = QLabel("Less"), QLabel("More")
        less.setEnabled(False)
        more.setEnabled(False)
        self.brush = QSlider(Qt.Horizontal, minimum=0, maximum=len(self.BRUSHES) - 1,
                             value=len(self.BRUSHES) - 1, toolTip="How much one click takes or cuts away")
        self.brush.setFixedWidth(90)
        self.brush_label = QLabel()
        self.brush_label.setMinimumWidth(QFontMetrics(self.font()).horizontalAdvance("Whole object"))
        self.brush.valueChanged.connect(self.show_brush)
        self.show_brush()
        self.subject_tray = Tray(less, self.reach, more)
        self.brush_tray = Tray(QLabel("Brush"), self.brush, self.brush_label)
        self.copy_tray = Tray(copy_subject, save_subject)
        self.text = QPlainTextEdit(readOnly=True)
        self.text.setProperty("peekFocus", True)  # selectable and copyable
        copy = QPushButton(QIcon.fromTheme("edit-copy"), "Copy text")
        copy.clicked.connect(lambda: QApplication.clipboard().setText(
            self.text.textCursor().selectedText().replace("\u2029", "\n") or self.text.toPlainText()))
        self.text_panel = vbox(self.text, copy)
        self.text_panel.setFixedWidth(300)
        self.search = SearchBar(self.view, self)
        self.search.changed.connect(lambda text: self.find(text, 0))
        self.search.step.connect(lambda delta: self.find(self.search.edit.text(), delta))
        self.search.edit.textEdited.connect(lambda _: self.words is None and self.read_text())
        self.matches, self.match = [], -1

        self.zoom_label.setMinimumWidth(44)
        self.zoom_label.setAlignment(Qt.AlignCenter)
        self.ocr_status.setEnabled(False)  # dimmed: it's a status, not a control
        statuses = QWidget()
        status_row = QHBoxLayout(statuses)
        status_row.setContentsMargins(0, 0, 0, 0)
        status_row.addWidget(self.ocr_status)
        status_row.addWidget(self.subject_status)
        status_row.addStretch()
        bar = toolbar(statuses, self.brush_tray, self.subject_tray, self.copy_tray,
                      (self.subject_button, self.text_button, self.search.button()),
                      (tool_button("zoom-out", "Zoom out (-)", lambda: self.view.zoom(1 / 1.25)),
                       self.zoom_label,
                       tool_button("zoom-in", "Zoom in (+)", lambda: self.view.zoom(1.25)),
                       tool_button("zoom-original", "Actual size (1)", self.view.actual_size),
                       tool_button("zoom-fit-best", "Fit window (0)", self.view.fit)))
        body = QHBoxLayout()
        body.addWidget(self.view, 1)
        body.addWidget(self.text_panel)
        layout = QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.addLayout(bar)
        layout.addLayout(body, 1)
        self.text_panel.hide()  # after parenting, or it pops up as its own window
        self.subject_tray.hide()
        self.copy_tray.hide()
        self.brush_tray.hide()
        QShortcut(QKeySequence.Undo, self, lambda: self.subject_button.isChecked()
                  and (self.pending.append((0, 0, "undo", 0)) or self.next_pick()))
        for keys, slot in ((("+", "="), lambda: self.view.zoom(1.25)), (("-",), lambda: self.view.zoom(1 / 1.25)),
                           (("1",), self.view.actual_size), (("0",), self.view.fit),
                           (("T",), self.text_button.toggle),
                           (("S",), lambda: self.subject_button.isEnabled() and self.subject_button.toggle())):
            for key in keys:
                QShortcut(key, self, slot)

    # ---- picking a subject: click what to keep, Shift+click adds, Ctrl+click removes

    HINT = "Click what to keep · Shift+click adds · Ctrl+click cuts away · Ctrl+Z undoes"

    BRUSHES = (0.03, 0.06, 0.12, 0.25, 0.5, None)  # brush side, share of the picture's longest side

    def brush_size(self):
        """Brush side in picture pixels, 0 = whole object."""
        share = self.BRUSHES[self.brush.value()]
        return round(share * max(self.image.width(), self.image.height())) if share else 0

    def show_brush(self):
        size = self.brush_size()
        self.brush_label.setText(f"{size} px" if size else "Whole object")
        self.view.brush = size

    def subject_mode(self, on):
        self.view.picking = on
        self.brush_tray.setVisible(on)
        self.view.show_brush()
        if not on:
            self.clear_subject()
            self.pending.clear()
            self.history = []
            self.subject_status.setText("")
            return
        if self.segmenter:
            self.subject_status.setText(self.HINT)
        elif not self.ticker.isActive():  # read the picture once
            self.subject_status.setText("Starting…")
            self.ticker.start()
            image, progress = self.image, self.progress
            in_background(lambda: run_catching(lambda: Segmenter(image, progress)), self, self.got_segmenter,
                          pool=_convert_pool)

    def got_segmenter(self, result):
        self.ticker.stop()
        if isinstance(result, Exception):
            self.subject_status.setText(f"Subject picking failed: {str(result)[:100]}")
            return
        self.segmenter = result
        if self.subject_button.isChecked():
            self.subject_status.setText(self.HINT)
            self.next_pick()
        # Meanwhile the edge model studies the picture; the outline sharpens when it's done
        image, progress = self.image, self.progress
        progress.text = "Clicks work now · sharpening the edges…"
        self.ticker.start()
        in_background(lambda: run_catching(lambda: upscale_array(matte(image, progress), image.height(),
                                                                 image.width())),
                      self, self.got_edges, pool=_convert_pool)

    def got_edges(self, result):
        self.ticker.stop()
        self.progress.text = ""
        if isinstance(result, Exception):
            self.subject_status.setText(f"Edge sharpening failed, outlines stay rougher: {str(result)[:80]}")
            return
        self.segmenter.edges = result
        if self.subject_button.isChecked():
            self.subject_status.setText(self.HINT)
            if self.parts:
                self.pending.append((0, 0, "reach", 0))  # redraw the current pick with sharp edges
                self.next_pick()

    def picked(self, point, modifiers):
        if not self.subject_button.isChecked() or not self.segmenter:
            return  # still reading the picture: clicks made now are ignored, not saved for later
        mode = "add" if modifiers & Qt.ShiftModifier else "remove" if modifiers & Qt.ControlModifier else "new"
        if not QRectF(self.image.rect()).contains(point):
            if mode == "new" and self.parts:
                self.history.append(self.parts)
                self.clear_subject()  # a click beside the picture lets go
            return
        self.pending = [p for p in self.pending if p[2] != "reach"]  # the click redoes the mask anyway
        self.pending.append((int(point.x()), int(point.y()), mode, self.brush_size()))
        self.next_pick()

    def next_pick(self):
        """Handle the oldest waiting click (or slider move), one at a time, off the GUI thread."""
        if self.busy or not self.segmenter or not self.pending:
            return
        x, y, mode, brush = self.pending.pop(0)
        if mode == "reach" and any(p[2] == "reach" for p in self.pending):
            return self.next_pick()  # a newer slider position is already waiting
        if mode == "undo":
            if not self.history:
                return self.next_pick()
            self.parts = self.history.pop()
            if not self.parts:
                self.clear_subject()
                return self.next_pick()
        segmenter, parts, bias = self.segmenter, list(self.parts), self.reach.value()
        box = None
        if brush:
            box = (max(0, x - brush // 2), max(0, y - brush // 2),
                   min(segmenter.width, x + brush // 2 + 1), min(segmenter.height, y + brush // 2 + 1))
        self.busy = True

        def work():
            nonlocal parts
            if mode == "new":
                parts = [segmenter.part([(x, y)], [1], box=box)]
            elif mode == "add":
                parts.append(segmenter.part([(x, y)], [1], box=box))
            elif mode == "remove" and box:
                # Brush: cut what the model sees under the click, only inside the brush
                local = segmenter.part([(x, y)], [1], box=box)["logits"]
                import numpy as np
                parts = [dict(part, logits=np.minimum(part["logits"], -local)) for part in parts]
                parts = [part for part in parts if (part["logits"] > 0).any()]
            elif mode == "remove":
                # "not here" for each picked object under the click: the model reshapes that object
                for i, part in enumerate(parts):
                    if part["logits"][y, x] > 0:
                        part = segmenter.part(part["points"] + [(x, y)], part["labels"] + [0], part,
                                              box=part.get("box"))
                        if part["logits"][y, x] > 0:
                            # Still there: the click contradicts the part's own clicks (e.g. Ctrl+click on
                            # something added with Shift+click). Cut the object under the click out of it.
                            import numpy as np
                            part["logits"] = np.minimum(part["logits"], -segmenter.part([(x, y)], [1])["logits"])
                        parts[i] = part
                parts = [part for part in parts if (part["logits"] > 0).any()]
            if not parts:
                return parts, None, None, None, mode, (x, y)
            logits = merge_parts(parts)
            height, width = logits.shape
            seeds = [pt for part in parts for pt, label in zip(part["points"], part["labels"]) if label == 1]
            reach = max(4, round(max(width, height) / 128))  # about 2 of the model's 256 px pixels
            inside = tidy(snap_to_edges(logits + bias > 0, segmenter.guide, reach), seeds)
            mask = feather(inside, max(1, round(max(width, height) / 1200)))
            if segmenter.edges is not None:
                mask = sharpen(inside, mask, segmenter.edges, reach, bias)
            return parts, logits, mask, subject_layers(mask), mode, (x, y)

        in_background(lambda: run_catching(work), self, self.show_subject, pool=_convert_pool)

    def show_subject(self, result):
        self.busy = False
        if isinstance(result, Exception):
            self.subject_status.setText(f"Subject picking failed: {str(result)[:100]}")
        elif self.subject_button.isChecked():
            parts, logits, mask, layers, mode, (x, y) = result
            if mode in ("new", "add", "remove"):
                self.history = (self.history + [self.parts])[-30:]  # Ctrl+Z goes back here
            if mask is None or mask.max() < 0.5:
                self.clear_subject()
            else:
                self.parts, self.logits, self.mask = parts, logits, mask
                self.draw_subject(*layers, mode, QPointF(x, y))
        self.next_pick()

    def draw_subject(self, dim, edge, mode, point):
        """Dim everything else and outline the subject with a glow.
        A new pick 'lifts': the dimming fades in while the glow tightens around the subject.
        Shift/Ctrl+click: the change spreads out from the click, so only what changed moves.
        Slider and Ctrl+Z: the old look fades into the new one."""
        scene = self.view.scene()
        self.end_transition()
        old = self.layers
        shade = QGraphicsPixmapItem(QPixmap.fromImage(dim))
        glow = QGraphicsPixmapItem(QPixmap.fromImage(edge))
        blur = QGraphicsBlurEffect(blurRadius=12)
        glow.setGraphicsEffect(blur)
        line = QGraphicsPixmapItem(QPixmap.fromImage(edge))
        line.setOpacity(0.85)
        new = [shade, glow, line]
        for z, item in enumerate(new, 1):
            item.setZValue(z)
            item.setTransformationMode(Qt.SmoothTransformation)
            scene.addItem(item)
        self.layers = new
        self.subject_tray.show()
        self.copy_tray.show()
        self.subject_status.setText(self.HINT)

        if mode == "new" or not old:
            for item in old:
                scene.removeItem(item)

            def step(t):
                shade.setOpacity(t)
                glow.setOpacity(t)
                line.setOpacity(0.85 * t)
                blur.setBlurRadius(12 + 36 * (1 - t))
            finish, duration = None, 650
        elif mode in ("add", "remove"):
            # Two clipping shapes: the new layers show inside a circle growing from the click, the old ones
            # everywhere else. Each spot shows exactly one of them, so nothing is dimmed twice.
            inside, outside = QGraphicsPathItem(), QGraphicsPathItem()
            for clip, items in ((inside, new), (outside, old)):
                clip.setPen(Qt.NoPen)
                clip.setFlag(QGraphicsItem.ItemClipsChildrenToShape)
                scene.addItem(clip)
                for item in items:
                    item.setParentItem(clip)
            area = QRectF(self.image.rect())
            reach = max(QLineF(point, corner).length() for corner in
                        (area.topLeft(), area.topRight(), area.bottomLeft(), area.bottomRight()))

            def step(t):
                circle = QPainterPath()
                circle.addEllipse(point, t * reach, t * reach)
                inside.setPath(circle)
                rest = QPainterPath()  # odd-even fill: the area minus the circle
                rest.addRect(area.adjusted(-100, -100, 100, 100))
                rest.addEllipse(point, t * reach, t * reach)
                outside.setPath(rest)

            def finish():
                for item in new:
                    item.setParentItem(None)
                scene.removeItem(inside)
                scene.removeItem(outside)  # with the old layers in it
            duration = 550
        else:
            def step(t):
                for item, opacity in zip(new, (1, 1, 0.85)):
                    item.setOpacity(opacity * t)
                for item, opacity in zip(old, (1, 1, 0.85)):
                    item.setOpacity(opacity * (1 - t))

            def finish():
                for item in old:
                    scene.removeItem(item)
            duration = 220
        animation = QVariantAnimation(self, startValue=0.0, endValue=1.0, duration=duration,
                                      easingCurve=QEasingCurve.OutCubic)
        animation.valueChanged.connect(step)
        step(0)
        if finish:
            def done():
                if self.transition and self.transition[0] is animation:
                    self.transition = None
                    finish()
            animation.finished.connect(done)
            self.transition = (animation, finish)
        animation.start(QAbstractAnimation.DeleteWhenStopped)

    def end_transition(self):
        """Jump a running add/remove/fade animation to its end, before drawing a newer pick."""
        if self.transition:
            animation, finish = self.transition
            self.transition = None
            animation.stop()
            finish()

    def clear_subject(self):
        self.end_transition()
        for item in self.layers:
            self.view.scene().removeItem(item)
        self.layers, self.mask, self.logits, self.parts = [], None, None, []
        self.subject_tray.hide()
        self.copy_tray.hide()
        if self.subject_button.isChecked() and self.segmenter:
            self.subject_status.setText(self.HINT)

    def leave_subject(self):
        """Esc: first lets go of the pick, then leaves picking."""
        if self.mask is not None:
            self.clear_subject()
        else:
            self.subject_button.setChecked(False)

    def cutout(self):
        return cut_out(self.image, self.mask) if self.mask is not None else None

    def copy_subject(self):
        if cut := self.cutout():
            QApplication.clipboard().setImage(cut)
            self.subject_status.setText("Subject copied")

    def save_subject(self):
        if not (cut := self.cutout()) or not self.path:
            return
        dst = free_path(self.path.parent, self.path.stem + " (subject)", ".png")
        if cut.save(str(dst)):
            self.subject_status.setText(f'Saved <a href="{QUrl.fromLocalFile(str(dst)).toLocalFile()}">'
                                        f'{escape(dst.name)}</a>')
        else:
            self.subject_status.setText("Couldn't write the PNG")

    def show_text(self, on):
        self.text_panel.setVisible(on)
        for box in self.boxes:
            box.setVisible(on)
        if on and self.words is None:
            self.read_text()

    def read_text(self):
        if self.words is not None or self.ocr_status.text() or not self.text_button.isEnabled():
            return  # already read, or reading
        self.ocr_status.setText("Reading text…")
        key = cache_key(self.path) if self.path and self.path.is_file() else None
        in_background(lambda: ocr_words(self.image, key), self, self.got_words)

    def got_words(self, words):
        self.words = words
        self.ocr_status.setText(f"{len(words)} words" if words else "No text found")
        self.text.setPlainText(ocr_text(words))
        accent = self.palette().highlight().color()
        fill = QColor(accent)
        fill.setAlpha(45)
        pen = QPen(accent, 0)  # 0 = one screen pixel at any zoom
        for _, (x, y, w, h), _ in words:
            box = self.view.scene().addRect(x - 1, y - 1, w + 2, h + 2, pen, fill)
            box.setVisible(self.text_button.isChecked())
            self.boxes.append(box)
        if self.search.edit.text():
            self.find(self.search.edit.text(), 0)

    def find(self, text, delta):
        """Highlight words containing `text`; Enter moves to the next one."""
        accent = self.palette().highlight().color()
        if delta == 0:
            self.matches = [i for i, (word, _, _) in enumerate(self.words or [])
                            if text and text.lower() in word.lower()]
            self.match = -1
            for i, box in enumerate(self.boxes):
                strong = i in self.matches
                fill = QColor(accent)
                fill.setAlpha(140 if strong else 45)
                box.setBrush(fill)
                box.setVisible(strong or self.text_button.isChecked())
        if self.words is None:
            self.search.label.setText("reading…" if text else "")
            return
        if not self.matches:
            self.search.label.setText("no matches" if text else "")
            return
        self.match = (self.match + (delta or 1)) % len(self.matches)
        self.view.centerOn(self.boxes[self.matches[self.match]])
        self.search.label.setText(f"{self.match + 1} of {len(self.matches)}")


# ---------------------------------------------------------------- documents

class DocView(QWidget):
    """PDF pages with a thumbnail sidebar, page counter, zoom and search."""
    editing_changed = Signal(bool)

    def __init__(self, path, source=None):
        super().__init__()
        self.path = Path(path)
        self.source = Path(source or path)  # the file you opened, when `path` is its converted PDF
        self.editor = None
        # Creation order = deletion order. Search model and doc go before the view that uses them;
        # a doc deleted after its view segfaults.
        self.search_model = QPdfSearchModel(self)
        self.doc = QPdfDocument(self)
        self.doc.load(str(display_pdf(path)))
        self.search_model.setDocument(self.doc)
        self.view = QPdfView()
        self.view.setDocument(self.doc)
        self.view.setSearchModel(self.search_model)
        self.view.setPageMode(QPdfView.PageMode.MultiPage)
        self.view.setZoomMode(QPdfView.ZoomMode.FitToWidth)
        self.view.viewport().installEventFilter(self)  # Ctrl+wheel zoom
        self.nav = self.view.pageNavigator()
        self.nav.currentPageChanged.connect(self.page_changed)

        self.thumbs = QListWidget()
        self.thumbs.setFixedWidth(180)
        self.thumbs.setIconSize(QSize(140, 200))
        self.thumbs.setViewMode(QListWidget.IconMode)
        self.thumbs.setMovement(QListWidget.Static)
        self.thumbs.setFlow(QListWidget.TopToBottom)
        self.thumbs.setWrapping(False)
        self.thumbs.setGridSize(QSize(160, 230))  # without a fixed cell, icon mode overlaps the pages
        self.thumbs.currentRowChanged.connect(self.go_to)
        for i in range(self.doc.pageCount()):
            QListWidgetItem(str(i + 1), self.thumbs).setTextAlignment(Qt.AlignHCenter)
        self.next_thumb = 0
        self.thumb_timer = QTimer(self, interval=0, timeout=self.render_thumb)  # one page per tick: never blocks
        self.thumb_timer.start()

        reader = QWidget()  # the reading page; its shortcuts sleep while the editor page shows
        self.search = SearchBar(self.view, reader)
        self.search.changed.connect(self.search_model.setSearchString)
        self.search.step.connect(self.next_match)
        self.search_model.modelReset.connect(self.update_matches)
        self.search_model.rowsInserted.connect(self.update_matches)

        self.page_label = QLabel()
        edit = tool_button("document-edit", "Draw shapes, highlight and write text on the pages (E)",
                           lambda: self.set_editing(True), text="Edit")
        if not pymupdf:
            edit.setEnabled(False)
            edit.setToolTip("Install python-pymupdf to draw and write on PDFs")
        bar = toolbar(self.page_label, (self.search.button(), edit),
                      (tool_button("zoom-out", "Zoom out", lambda: self.zoom(1 / 1.25)),
                       tool_button("zoom-in", "Zoom in", lambda: self.zoom(1.25)),
                       tool_button("zoom-fit-width", "Fit width", lambda: self.view.setZoomMode(QPdfView.ZoomMode.FitToWidth)),
                       tool_button("zoom-fit-best", "Fit page", lambda: self.view.setZoomMode(QPdfView.ZoomMode.FitInView))))
        if pymupdf:
            QShortcut("E", reader, lambda: self.set_editing(True))
        body = QHBoxLayout()
        body.addWidget(self.thumbs)
        body.addWidget(self.view, 1)
        reader_layout = QVBoxLayout(reader)
        reader_layout.setContentsMargins(0, 0, 0, 0)
        reader_layout.addLayout(bar)
        reader_layout.addLayout(body, 1)
        self.stack = QStackedWidget()
        self.stack.addWidget(reader)
        layout = QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.addWidget(self.stack)

        QShortcut(QKeySequence.MoveToPreviousPage, reader, lambda: self.go_to(self.nav.currentPage() - 1))
        QShortcut(QKeySequence.MoveToNextPage, reader, lambda: self.go_to(self.nav.currentPage() + 1))
        QShortcut(QKeySequence.ZoomIn, reader, lambda: self.zoom(1.25))
        QShortcut(QKeySequence.ZoomOut, reader, lambda: self.zoom(1 / 1.25))
        self.page_changed(0)

    @property
    def editing(self):
        return self.stack.currentIndex() == 1

    def set_editing(self, on):
        """Swap the page reader for the drawing canvas. False if you chose to stay because of unsaved drawings."""
        if on == self.editing:
            return True
        if on:
            self.editor = PdfEditor(self.doc, self.path, self.source)
            self.editor.done.connect(lambda: self.set_editing(False))
            self.stack.addWidget(self.editor)
            self.stack.setCurrentIndex(1)
            self.editor.canvas.setFocus()
            page = self.nav.currentPage()
            QTimer.singleShot(0, lambda: self.editor and self.editor.canvas.go_to(page))  # after the first layout
        else:
            if not self.editor.confirm_leave():
                return False
            page = self.editor.canvas.current_page()
            self.stack.setCurrentIndex(0)
            self.editor.deleteLater()  # the next Edit starts clean, from the file
            self.editor = None
            self.go_to(page)
            self.window().setFocus()
        self.editing_changed.emit(on)
        return True

    def render_thumb(self):
        i = self.next_thumb
        if i >= self.doc.pageCount():
            self.thumb_timer.stop()
            return
        size = self.doc.pagePointSize(i).toSize().scaled(QSize(140, 200), Qt.KeepAspectRatio)
        page = QImage(size, QImage.Format_RGB32)  # pages render transparent: put them on paper white
        page.fill(Qt.white)
        painter = QPainter(page)
        painter.drawImage(0, 0, self.doc.render(i, size))
        painter.end()
        self.thumbs.item(i).setIcon(QIcon(QPixmap.fromImage(page)))
        self.next_thumb += 1

    def go_to(self, page):
        if 0 <= page < self.doc.pageCount() and page != self.nav.currentPage():
            self.nav.jump(page, QPointF(), self.nav.currentZoom())

    def page_changed(self, page):
        self.page_label.setText(f"Page {page + 1} of {self.doc.pageCount()}")
        self.thumbs.blockSignals(True)  # highlighting the thumbnail must not jump back
        self.thumbs.setCurrentRow(page)
        self.thumbs.blockSignals(False)

    def update_matches(self):
        count = self.search_model.rowCount(QModelIndex())
        self.search.label.setText(f"{count} matches" if self.search.edit.text() else "")

    def next_match(self, delta):
        count = self.search_model.rowCount(QModelIndex())
        if not count:
            return
        index = (self.view.currentSearchResultIndex() + delta) % count
        self.view.setCurrentSearchResultIndex(index)
        self.nav.jump(self.search_model.resultAtIndex(index))
        self.search.label.setText(f"{index + 1} of {count}")

    def zoom(self, factor):
        self.view.setZoomMode(QPdfView.ZoomMode.Custom)
        self.view.setZoomFactor(min(max(self.view.zoomFactor() * factor, 0.1), 10))

    def eventFilter(self, obj, event):
        if event.type() == QEvent.Wheel and event.modifiers() & Qt.ControlModifier:
            self.zoom(1.25 if event.angleDelta().y() > 0 else 1 / 1.25)
            return True
        return False


# ---------------------------------------------------------------- PDF editor

EDIT_SCALE = 2  # scene units per PDF point
PAGE_GAP = 24  # scene units around and between pages
EDIT_COLORS = ["#e53935", "#1e88e5", "#43a047", "#fdd835", "#212121"]
EDIT_WIDTHS = {"Thin": 1.5, "Medium": 3, "Thick": 6}  # points
TEXT_SIZES = {"Small": 10, "Normal": 14, "Large": 20, "Huge": 28}  # points
EDIT_TOOLS = (("select", "edit-select", "Select and move (S) — double-click text to change it"),
              ("rect", "draw-rectangle", "Rectangle (R) — Shift for a square"),
              ("ellipse", "draw-ellipse", "Ellipse (O) — Shift for a circle"),
              ("arrow", "draw-arrow", "Arrow (A)"),
              ("pen", "draw-freehand", "Pen (P)"),
              ("highlight", "draw-highlight", "Highlight text (H) — drag across the words"),
              ("text", "draw-text", "Text (T) — click to write, click a text to change it"))
KIND, PAGE, SHAPE = 0, 1, 2  # QGraphicsItem.data() keys of a drawing
RENDER_ANNOTATIONS = QPdfDocumentRenderOptions()
RENDER_ANNOTATIONS.setRenderFlags(QPdfDocumentRenderOptions.RenderFlag.Annotations)


def display_pdf(path):
    """`path`, or a cached copy with its annotations painted into the pages:
    Qt's PDF view leaves annotations (comments, highlights, filled forms) out."""
    if not pymupdf:
        return path
    try:
        with open(path, "rb") as f, mmap.mmap(f.fileno(), 0, access=mmap.ACCESS_READ) as data:
            if data.find(b"/Annot") < 0:  # most PDFs: no annotations, nothing to do
                return path
        cached = CACHE / f"{cache_key(Path(path))}.baked.pdf"
        if cached.exists():
            cached.touch()
            return cached
        with pymupdf.open(path) as doc:
            if doc.needs_pass or not any(page.first_annot or page.first_widget for page in doc):
                return path
            doc.bake()
            CACHE.mkdir(parents=True, exist_ok=True)
            doc.save(cached, garbage=1, deflate=True)
        return cached
    except Exception:  # damaged or unusual PDF: show it the way Qt can
        return path


def page_words(page):
    """Word boxes of a PyMuPDF page in reading order, in displayed (rotated) points: [(QRectF, line id)]."""
    words = []
    for x0, y0, x1, y1, _, block, line, _ in page.get_text("words", sort=True):
        box = pymupdf.Rect(x0, y0, x1, y1) * page.rotation_matrix
        words.append((QRectF(box.x0, box.y0, box.width, box.height), (block, line)))
    return words


def words_between(words, start, end):
    """The line boxes covering the words from the one nearest `start` to the one nearest `end`, like a text selection."""
    def nearest(point):
        def distance(box):
            dx = max(box.left() - point.x(), 0, point.x() - box.right())
            dy = max(box.top() - point.y(), 0, point.y() - box.bottom())
            return dy * 4 + dx  # lines matter more than columns
        return min(range(len(words)), key=lambda i: distance(words[i][0]))
    first, last = sorted((nearest(start), nearest(end)))
    lines = {}
    for box, line in words[first:last + 1]:
        lines[line] = lines[line].united(box) if line in lines else box
    return list(lines.values())


def arrow_path(start, end, width):
    """A line with an open arrow head at `end`."""
    path = QPainterPath(start)
    path.lineTo(end)
    line = QLineF(end, start)
    head = width * 3 + 10
    for turn in (28, -28):
        side = QLineF(line)
        side.setLength(head)
        side.setAngle(line.angle() + turn)
        path.moveTo(end)
        path.lineTo(side.p2())
    return path


class TextBox(QGraphicsTextItem):
    """Text on a page. Typing only while editing; otherwise it moves like any drawing."""

    def __init__(self, size, color):
        super().__init__()
        self.document().setDocumentMargin(2 * EDIT_SCALE)
        self.set_style(size, color)

    def set_style(self, size, color):
        self.size = size
        font = QFont("Helvetica")  # same shapes as the PDF's built-in Helvetica the saved text uses
        font.setStyleHint(QFont.SansSerif)
        font.setPixelSize(round(size * EDIT_SCALE))
        self.setFont(font)
        self.setDefaultTextColor(color)

    def start_editing(self):
        self.setTextInteractionFlags(Qt.TextEditorInteraction)
        self.setFocus()

    def focusOutEvent(self, event):
        super().focusOutEvent(event)
        if event.reason() == Qt.PopupFocusReason:  # the right-click menu: still typing
            return
        self.setTextInteractionFlags(Qt.NoTextInteraction)
        cursor = self.textCursor()
        cursor.clearSelection()
        self.setTextCursor(cursor)
        if not self.toPlainText().strip():  # an empty box is a click that wrote nothing
            QTimer.singleShot(0, lambda: self.scene() and self.scene().views()[0].discard(self))

    def mouseDoubleClickEvent(self, event):
        if not self.textInteractionFlags() & Qt.TextEditorInteraction:
            self.start_editing()
        super().mouseDoubleClickEvent(event)

    def paint(self, painter, option, widget=None):
        super().paint(painter, option, widget)
        if self.hasFocus():  # a frame while typing shows where the box is
            painter.setPen(QPen(option.palette.highlight().color(), 0, Qt.DashLine))
            painter.drawRect(self.boundingRect())


class PdfCanvas(QGraphicsView):
    """Every page of a PDF in one scrolling column, with drawing tools on top.
    Pages are rendered only near the screen, sharp at the current zoom."""
    selection_changed = Signal()
    page_changed = Signal(int)

    def __init__(self, doc, words):
        super().__init__()
        self.doc, self.words = doc, words
        self.setScene(QGraphicsScene(self))
        self.setRenderHint(QPainter.Antialiasing)
        self.setBackgroundBrush(self.palette().window())
        self.setFrameShape(QGraphicsView.NoFrame)
        self.setTransformationAnchor(QGraphicsView.AnchorUnderMouse)
        self.pages = []
        widest, y = 0, 0
        edge = QColor(self.palette().shadow().color())
        edge.setAlpha(70)
        for i in range(doc.pageCount()):
            size = doc.pagePointSize(i) * EDIT_SCALE
            rect = QRectF(-size.width() / 2, y, size.width(), size.height())
            self.scene().addRect(rect, QPen(edge, 0), QColor(Qt.white)).setZValue(-2)  # paper until rendered
            picture = self.scene().addPixmap(QPixmap())
            picture.setPos(rect.topLeft())
            picture.setZValue(-1)
            picture.setTransformationMode(Qt.SmoothTransformation)
            self.pages.append(types.SimpleNamespace(rect=rect, picture=picture, scale=0))
            widest = max(widest, size.width())
            y += size.height() + PAGE_GAP
        self.scene().setSceneRect(QRectF(-widest / 2 - PAGE_GAP, -PAGE_GAP, widest + 2 * PAGE_GAP, y + PAGE_GAP))
        self.history = []  # ("add", [items]) / ("delete", [items]), for Undo
        self.tool = "rect"
        self.color = QColor(EDIT_COLORS[0])
        self.width_pt = EDIT_WIDTHS["Medium"]
        self.text_size = TEXT_SIZES["Normal"]
        self.drawing = None
        self.fitting = True
        self.render_timer = QTimer(self, singleShot=True, interval=15, timeout=self.render_visible)
        self.verticalScrollBar().valueChanged.connect(self.render_timer.start)
        self.scene().selectionChanged.connect(self.selection_changed)

    # ---- pages

    def render_visible(self):
        view = self.mapToScene(self.viewport().rect()).boundingRect()
        near = view.adjusted(0, -view.height(), 0, view.height())  # a screen above and below, for smooth scrolling
        zoom = self.transform().m11() * self.devicePixelRatioF()
        for i, page in enumerate(self.pages):
            if not page.rect.intersects(near):
                if page.scale:  # far away: give the memory back
                    page.picture.setPixmap(QPixmap())
                    page.scale = 0
                continue
            # pixels per scene unit; capped at 16 Mpx a page, so deep zoom can't eat gigabytes
            want = min(zoom, 3, (16e6 / (page.rect.width() * page.rect.height())) ** 0.5)
            if page.scale and abs(page.scale - want) < want * 0.1:
                continue
            pixels = (page.rect.size() * want).toSize()
            image = QImage(pixels, QImage.Format_RGB32)  # pages render transparent: paper white first
            image.fill(Qt.white)
            painter = QPainter(image)
            painter.drawImage(0, 0, self.doc.render(i, pixels, RENDER_ANNOTATIONS))
            painter.end()
            page.picture.setPixmap(QPixmap.fromImage(image))
            page.picture.setScale(1 / want)
            page.scale = want
        self.page_changed.emit(self.current_page())

    def page_at(self, point):
        return next((i for i, page in enumerate(self.pages) if page.rect.contains(point)), None)

    def current_page(self):
        middle = self.mapToScene(self.viewport().rect().center()).y()
        return min(range(len(self.pages)),
                   key=lambda i: 0 if self.pages[i].rect.top() <= middle <= self.pages[i].rect.bottom()
                   else min(abs(self.pages[i].rect.top() - middle), abs(self.pages[i].rect.bottom() - middle)))

    def go_to(self, page):
        if 0 <= page < len(self.pages):
            top = self.mapFromScene(QPointF(0, self.pages[page].rect.top() - PAGE_GAP / 2)).y()
            self.verticalScrollBar().setValue(self.verticalScrollBar().value() + top)

    def zoom(self, factor, under_mouse=False):
        self.fitting = False
        factor = min(max(self.transform().m11() * factor, 0.1), 8) / self.transform().m11()
        if not under_mouse:
            self.setTransformationAnchor(QGraphicsView.AnchorViewCenter)
        self.scale(factor, factor)
        self.setTransformationAnchor(QGraphicsView.AnchorUnderMouse)
        self.render_timer.start()

    def fit(self):
        """Page width fills the view."""
        self.fitting = True
        page = self.current_page()
        scale = self.viewport().width() / self.scene().sceneRect().width()
        self.setTransform(QTransform.fromScale(scale, scale))
        self.go_to(page)
        self.render_timer.start()

    def wheelEvent(self, event):
        if event.modifiers() & Qt.ControlModifier:
            self.zoom(1.25 if event.angleDelta().y() > 0 else 1 / 1.25, under_mouse=True)
        else:
            super().wheelEvent(event)

    def resizeEvent(self, event):
        super().resizeEvent(event)
        if self.fitting:
            self.fit()
        self.render_timer.start()

    # ---- drawing

    def set_tool(self, tool):
        self.tool = tool
        self.finish_typing()
        self.setDragMode(QGraphicsView.RubberBandDrag if tool == "select" else QGraphicsView.NoDrag)
        self.viewport().setCursor(Qt.ArrowCursor if tool == "select" else
                                  Qt.IBeamCursor if tool in ("text", "highlight") else Qt.CrossCursor)

    def finish_typing(self):
        """End the text box being typed in. False if there was none."""
        item = self.scene().focusItem()
        if not isinstance(item, TextBox):
            return False
        item.clearFocus()
        return True

    def pen(self, kind, color=None, width=None):
        color = QColor(color or self.color)
        if kind == "marker":  # highlighter on a page without text: a wide see-through stroke
            color.setAlpha(100)
            pen = QPen(color, 12 * EDIT_SCALE)
            pen.setCapStyle(Qt.FlatCap)
        else:
            pen = QPen(color, (width or self.width_pt) * EDIT_SCALE)
            pen.setCapStyle(Qt.RoundCap)
        pen.setJoinStyle(Qt.RoundJoin)
        return pen

    def clamp(self, point):
        """`point` kept on the page being drawn on: outside it would be cut away."""
        rect = self.pages[self.drawing.data(PAGE)].rect
        return QPointF(min(max(point.x(), rect.left()), rect.right()), min(max(point.y(), rect.top()), rect.bottom()))

    def page_points(self, page, point):
        return (point - self.pages[page].rect.topLeft()) / EDIT_SCALE

    def mousePressEvent(self, event):
        if event.button() != Qt.LeftButton or self.tool == "select":
            return super().mousePressEvent(event)
        pos = self.mapToScene(event.position().toPoint())
        typing = self.scene().focusItem()
        if isinstance(typing, TextBox) and not typing.sceneBoundingRect().contains(pos):
            typing.clearFocus()  # a click elsewhere ends typing; the next click starts a new text
            return
        page = self.page_at(pos)
        if page is None:
            return
        if self.tool == "text":
            hit = next((i for i in self.items(event.position().toPoint()) if isinstance(i, TextBox)), None)
            if hit:  # change an existing text: the click places the cursor
                hit.start_editing()
                return super().mousePressEvent(event)
            box = TextBox(self.text_size, self.color)
            box.setPos(pos - QPointF(2 * EDIT_SCALE, self.text_size * EDIT_SCALE * 0.75))  # first line on the click
            self.add(box, "text", page)
            box.start_editing()
            return
        self.start = pos
        if self.tool in ("rect", "ellipse"):
            item = (QGraphicsRectItem if self.tool == "rect" else QGraphicsEllipseItem)(QRectF(pos, pos))
            kind = self.tool
        elif self.tool == "highlight" and self.words(page):
            item = QGraphicsPathItem()
            item.setPen(Qt.NoPen)
            kind = "highlight"
        else:  # arrow, pen, highlight on a scanned page
            # our own copy of the path: the item drops a path that is only a starting point
            self.path = QPainterPath(pos)
            item = QGraphicsPathItem(self.path)
            kind = "marker" if self.tool == "highlight" else self.tool
        if kind == "highlight":
            fill = QColor(self.color)
            fill.setAlpha(90)
            item.setBrush(fill)
        else:
            item.setPen(self.pen(kind))
        self.add(item, kind, page, record=False)
        self.drawing = item

    def mouseMoveEvent(self, event):
        if not self.drawing:
            return super().mouseMoveEvent(event)
        pos = self.clamp(self.mapToScene(event.position().toPoint()))
        kind, page = self.drawing.data(KIND), self.drawing.data(PAGE)
        if kind in ("rect", "ellipse"):
            rect = QRectF(self.start, pos).normalized()
            if event.modifiers() & Qt.ShiftModifier:  # Shift: square / circle
                side = max(rect.width(), rect.height())
                rect.setSize(QSizeF(side, side))
            self.drawing.setRect(rect)
        elif kind == "arrow":
            self.drawing.setPath(arrow_path(self.start, pos, self.width_pt * EDIT_SCALE))
            self.drawing.setData(SHAPE, (self.start, pos))
        elif kind == "highlight":
            boxes = words_between(self.words(page), self.page_points(page, self.start), self.page_points(page, pos))
            path = QPainterPath()
            origin = self.pages[page].rect.topLeft()
            for box in boxes:
                path.addRect(QRectF(origin + box.topLeft() * EDIT_SCALE, box.size() * EDIT_SCALE))
            self.drawing.setPath(path)
            self.drawing.setData(SHAPE, boxes)
        else:
            self.path.lineTo(pos)
            self.drawing.setPath(self.path)

    def mouseReleaseEvent(self, event):
        item, self.drawing = self.drawing, None
        if item:
            size = item.boundingRect().size()
            if item.data(KIND) != "highlight" and size.width() + size.height() < 6 + item.pen().widthF() * 2:
                self.scene().removeItem(item)  # a click without dragging draws nothing
            else:
                self.history.append(("add", [item]))
        super().mouseReleaseEvent(event)

    def add(self, item, kind, page, record=True):
        # add to the flags, don't replace them: a text box's own "focusable" flag is what lets you type in it
        item.setFlags(item.flags() | QGraphicsItem.ItemIsSelectable | QGraphicsItem.ItemIsMovable)
        item.setData(KIND, kind)
        item.setData(PAGE, page)
        self.scene().addItem(item)
        if record:
            self.history.append(("add", [item]))

    def drawings(self):
        return [item for item in self.scene().items() if item.data(KIND)]

    def discard(self, item):
        """Remove an empty text box without a trace in Undo."""
        if item.scene():
            self.scene().removeItem(item)
        self.history = [(action, items) for action, items in self.history if item not in items]

    def undo(self):
        if not self.history:
            return
        action, items = self.history.pop()
        for item in items:
            if action == "add":
                self.scene().removeItem(item)
            else:
                self.scene().addItem(item)

    def delete_selected(self):
        items = self.scene().selectedItems()
        for item in items:
            self.scene().removeItem(item)
        if items:
            self.history.append(("delete", items))

    def restyle(self, color=None, width=None, size=None):
        """Apply a new color, line width or text size to the selected drawings."""
        for item in self.scene().selectedItems():
            kind = item.data(KIND)
            if kind == "text":
                item.set_style(size or item.size, color or item.defaultTextColor())
            elif kind == "highlight" and color:
                fill = QColor(color)
                fill.setAlpha(90)
                item.setBrush(fill)
            elif kind != "highlight" and (color or width and kind != "marker"):
                item.setPen(self.pen(kind, color or item.pen().color(), width or item.pen().widthF() / EDIT_SCALE))
                if color and kind == "marker":
                    item.setPen(self.pen(kind, color))

    def annotations(self):
        """Every drawing as plain data for the save thread: page, kind, geometry in PDF points, style."""
        out = []
        for item in reversed(self.drawings()):  # oldest first, so later drawings stay on top
            kind, page = item.data(KIND), item.data(PAGE)
            here = lambda point: self.page_points(page, item.mapToScene(point))
            points = lambda p: (p.x(), p.y())
            if kind == "text":
                if not item.toPlainText().strip():
                    continue
                color = item.defaultTextColor()
                rect = QRectF(here(item.boundingRect().topLeft()), here(item.boundingRect().bottomRight()))
                geometry = (rect.left(), rect.top(), rect.right() + 2, rect.bottom() + 2)  # PDF text needs a hair more room
                out.append((page, kind, geometry, color.getRgbF()[:3], 0, item.toPlainText(), item.size))
                continue
            color = item.brush().color() if kind == "highlight" else item.pen().color()
            width = item.pen().widthF() / EDIT_SCALE
            if kind in ("rect", "ellipse"):
                rect = QRectF(here(item.rect().topLeft()), here(item.rect().bottomRight()))
                rect = rect.adjusted(-width / 2, -width / 2, width / 2, width / 2)  # PDF borders sit inside the box
                geometry = (rect.left(), rect.top(), rect.right(), rect.bottom())
            elif kind == "arrow":
                geometry = [points(here(p)) for p in item.data(SHAPE)]
            elif kind == "highlight":
                shift = (item.scenePos()) / EDIT_SCALE
                geometry = [(b.left() + shift.x(), b.top() + shift.y(), b.right() + shift.x(), b.bottom() + shift.y())
                            for b in item.data(SHAPE)]
            else:  # pen, marker
                path = item.path()
                geometry = [points(here(QPointF(path.elementAt(i).x, path.elementAt(i).y)))
                            for i in range(path.elementCount())]
            out.append((page, kind, geometry, color.getRgbF()[:3], width, "", 0))
        return out


def write_annotations(src, annotations, dst):
    """Copy PDF `src` to `dst` with the drawings added as real annotations,
    so any PDF app can still move or delete them. Worker thread."""
    doc = pymupdf.open(src)
    if doc.needs_pass and not doc.authenticate(""):  # many "protected" PDFs only restrict editing
        raise RuntimeError("this PDF needs a password")
    for number, kind, geometry, rgb, width, text, size in annotations:
        page = doc[number]
        flip = page.derotation_matrix  # we drew on the page as shown; annotations use its unrotated corners
        if kind in ("rect", "ellipse"):
            add = page.add_rect_annot if kind == "rect" else page.add_circle_annot
            annot = add(pymupdf.Rect(geometry) * flip)
        elif kind == "arrow":
            annot = page.add_line_annot(*(pymupdf.Point(p) * flip for p in geometry))
            annot.set_line_ends(pymupdf.PDF_ANNOT_LE_NONE, pymupdf.PDF_ANNOT_LE_OPEN_ARROW)
        elif kind in ("pen", "marker"):
            annot = page.add_ink_annot([[tuple(pymupdf.Point(p) * flip) for p in geometry]])
        elif kind == "highlight":
            annot = page.add_highlight_annot(quads=[pymupdf.Rect(box) * flip for box in geometry])
        else:
            annot = page.add_freetext_annot(pymupdf.Rect(geometry) * flip, text, fontsize=size, fontname="helv",
                                            text_color=rgb, rotate=page.rotation)
            annot.update()
            continue
        annot.set_colors(stroke=rgb)
        if kind in ("marker", "highlight"):
            annot.set_opacity(0.4)  # see-through in every viewer, also the ones that ignore highlight blending
            width = 12
        if kind != "highlight":
            annot.set_border(width=width)
        annot.update()
    doc.save(dst, garbage=1, deflate=True)
    return dst


class PdfEditor(QWidget):
    """Draw shapes, arrows, freehand lines, highlights and text on a PDF; save them as a new file."""
    done = Signal()

    def __init__(self, doc, pdf, source):
        super().__init__()
        self.pdf, self.source = Path(pdf), Path(source)
        self.fz = pymupdf.open(pdf)  # word boxes for the highlighter; saving reopens the file
        words = {}
        self.canvas = PdfCanvas(doc, lambda page: words.setdefault(page, page_words(self.fz[page])))
        self.canvas.setProperty("peekFocus", True)  # typing text needs the keyboard
        self.shortcuts = []  # off while typing, so letters, Ctrl+Z and Del reach the text
        self.saved = set()  # drawings as they were at the last save

        tools = []
        group = QButtonGroup(self)
        for tool, icon, tip in EDIT_TOOLS:
            button = tool_button(icon, tip, lambda on, t=tool: on and self.set_tool(t), checkable=True)
            group.addButton(button)
            tools.append(button)
            self.shortcuts.append(QShortcut(tip.split("(")[1][0], self, button.click))
        self.tool_buttons = dict(zip((t[0] for t in EDIT_TOOLS), tools))
        swatches = []
        colors = QButtonGroup(self)
        for color in EDIT_COLORS:
            swatch = QToolButton(checkable=True, autoRaise=True, toolTip="Color")
            swatch.setIcon(color_icon(color))
            swatch.toggled.connect(lambda on, c=color: on and self.set_color(QColor(c)))
            colors.addButton(swatch)
            swatches.append(swatch)
        swatches[0].setChecked(True)
        self.width_box = QComboBox(toolTip="Line width")
        self.width_box.addItems(EDIT_WIDTHS)
        self.width_box.setCurrentText("Medium")
        self.width_box.textActivated.connect(self.set_width)
        self.size_box = QComboBox(toolTip="Text size")
        self.size_box.addItems(TEXT_SIZES)
        self.size_box.setCurrentText("Normal")
        self.size_box.textActivated.connect(self.set_size)
        self.page_label = QLabel()
        self.page_label.setEnabled(False)  # dimmed: information, not a control
        self.status = QLabel()
        save = QPushButton(QIcon.fromTheme("document-save-as"), "Save copy")
        save.setToolTip("Write a new PDF next to the original (Ctrl+S). The original is never changed.")
        save.clicked.connect(self.save)

        bar = QHBoxLayout()
        bar.setSpacing(8)
        for group_widgets in (tools, swatches, (self.width_box, self.size_box),
                              (tool_button("edit-undo", "Undo (Ctrl+Z)", self.canvas.undo),
                               tool_button("edit-delete", "Delete selected (Del)", self.canvas.delete_selected))):
            bar.addWidget(Tray(*group_widgets))
        bar.addStretch()
        bar.addWidget(self.page_label)
        bar.addWidget(Tray(tool_button("zoom-out", "Zoom out (Ctrl+-)", lambda: self.canvas.zoom(1 / 1.25)),
                           tool_button("zoom-in", "Zoom in (Ctrl++)", lambda: self.canvas.zoom(1.25)),
                           tool_button("zoom-fit-width", "Fit width", self.canvas.fit)))
        bar.addWidget(self.status)
        bar.addWidget(Tray(tool_button("dialog-ok", "Back to reading (Esc)", self.done.emit, text="Done")))
        bar.addWidget(save)
        layout = QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.addLayout(bar)
        layout.addWidget(self.canvas, 1)

        self.shortcuts += [QShortcut(QKeySequence.Undo, self, self.canvas.undo),
                           QShortcut(QKeySequence.Delete, self, self.canvas.delete_selected)]
        QShortcut(QKeySequence.Save, self, self.save)
        QShortcut(QKeySequence.MoveToPreviousPage, self, lambda: self.canvas.go_to(self.canvas.current_page() - 1))
        QShortcut(QKeySequence.MoveToNextPage, self, lambda: self.canvas.go_to(self.canvas.current_page() + 1))
        QShortcut(QKeySequence.ZoomIn, self, lambda: self.canvas.zoom(1.25))
        QShortcut(QKeySequence.ZoomOut, self, lambda: self.canvas.zoom(1 / 1.25))
        self.canvas.scene().focusItemChanged.connect(
            lambda item, *_: [k.setEnabled(not isinstance(item, TextBox)) for k in self.shortcuts])
        self.canvas.selection_changed.connect(self.show_style_controls)
        self.canvas.page_changed.connect(lambda page: self.page_label.setText(f"Page {page + 1} of {len(self.canvas.pages)}"))
        # Every child except the canvas leaves the keyboard alone, so typed text always lands there
        for w in self.findChildren(QWidget):
            if w is not self.canvas and w is not self.canvas.viewport():
                w.setFocusPolicy(Qt.NoFocus)
        self.tool_buttons["rect"].setChecked(True)

    def set_tool(self, tool):
        self.canvas.set_tool(tool)
        self.show_style_controls()

    def show_style_controls(self):
        """Line width for shapes, text size for text, neither for the highlighter."""
        kinds = {item.data(KIND) for item in self.canvas.scene().selectedItems()}
        if not kinds:
            kinds = {self.canvas.tool}
        self.width_box.setVisible(bool(kinds & {"rect", "ellipse", "arrow", "pen"}))
        self.size_box.setVisible("text" in kinds)

    def set_color(self, color):
        self.canvas.color = color
        self.canvas.restyle(color=color)

    def set_width(self, name):
        self.canvas.width_pt = EDIT_WIDTHS[name]
        self.canvas.restyle(width=EDIT_WIDTHS[name])

    def set_size(self, name):
        self.canvas.text_size = TEXT_SIZES[name]
        self.canvas.restyle(size=TEXT_SIZES[name])

    def finish_typing(self):
        return self.canvas.finish_typing()

    def snapshot(self):
        return {repr(a) for a in self.canvas.annotations()}

    def save(self, wait=False):
        """Write "name (edited).pdf" next to the original. `wait`: on this thread, True if it worked."""
        self.finish_typing()
        annotations = self.canvas.annotations()
        if not annotations:
            self.status.setText("Nothing drawn yet")
            return False
        snapshot = {repr(a) for a in annotations}
        dst = free_path(self.source.parent, self.source.stem + " (edited)", ".pdf")
        self.status.setText("Saving…")

        def finished(result):
            if isinstance(result, Exception):
                self.status.setText(f"Save failed: {result}")
                return False
            self.saved = snapshot
            self.status.setText(f"Saved {result.name}")
            self.status.setToolTip(str(result))
            return True

        if wait:
            return finished(run_catching(lambda: write_annotations(self.pdf, annotations, dst)))
        in_background(lambda: run_catching(lambda: write_annotations(self.pdf, annotations, dst)), self, finished,
                      pool=_convert_pool)
        return True

    def confirm_leave(self):
        """True when it's fine to leave: nothing new since the last save, or you chose Save or Discard."""
        self.finish_typing()
        if self.snapshot() <= self.saved:
            return True
        answer = QMessageBox.question(self, "Unsaved drawings", "Save your drawings as a new PDF before leaving?",
                                      QMessageBox.Save | QMessageBox.Discard | QMessageBox.Cancel,
                                      QMessageBox.Save)
        if answer == QMessageBox.Save:
            return self.save(wait=True)
        return answer == QMessageBox.Discard


def run_catching(fn):
    try:
        return fn()
    except Exception as error:  # shown in the window instead of lost on a worker thread
        return error


def color_icon(color):
    pixmap = QPixmap(16, 16)
    pixmap.fill(Qt.transparent)
    painter = QPainter(pixmap)
    painter.setRenderHint(QPainter.Antialiasing)
    painter.setBrush(QColor(color))
    painter.setPen(QPen(QColor(128, 128, 128, 170), 1))  # mid gray: the dark swatch stays visible on dark themes
    painter.drawEllipse(QRectF(1.5, 1.5, 13, 13))
    painter.end()
    return QIcon(pixmap)


class TextPane(QWidget):
    """Plain text, code or rendered Markdown/HTML, with search."""

    def __init__(self, editor):
        super().__init__()
        self.editor = editor
        self.search = SearchBar(editor, self)
        self.search.changed.connect(lambda text: self.find(text, 0))
        self.search.step.connect(lambda delta: self.find(self.search.edit.text(), delta))
        layout = QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.addWidget(editor, 1)

    def find(self, text, delta):
        if not text:
            self.search.label.setText("")
            cursor = self.editor.textCursor()
            cursor.clearSelection()  # drop the last match's highlight
            self.editor.setTextCursor(cursor)
            return
        flags = QTextDocument.FindFlag(0) if delta >= 0 else QTextDocument.FindFlag.FindBackward
        if delta == 0:  # new search text: start from the top
            cursor = self.editor.textCursor()
            cursor.movePosition(cursor.MoveOperation.Start)
            self.editor.setTextCursor(cursor)
        found = self.editor.find(text, flags)
        if not found:  # wrap around
            cursor = self.editor.textCursor()
            cursor.movePosition(cursor.MoveOperation.Start if delta >= 0 else cursor.MoveOperation.End)
            self.editor.setTextCursor(cursor)
            found = self.editor.find(text, flags)
        count = self.editor.toPlainText().lower().count(text.lower())
        self.search.label.setText(f"{count} matches" if found else "no matches")


def plain_text_pane(text):
    edit = QPlainTextEdit(text, readOnly=True)
    edit.setFont(QFontDatabase.systemFont(QFontDatabase.FixedFont))
    return TextPane(edit)


def rich_text_pane(path, kind):
    """Markdown or HTML rendered by Qt. Relative images load from the file's folder; nothing from the web."""
    browser = QTextBrowser(openExternalLinks=True)
    browser.document().setBaseUrl(QUrl.fromLocalFile(str(Path(path).parent) + "/"))
    text = Path(path).read_bytes()[:TEXT_LIMIT].decode("utf-8", errors="replace")
    browser.setMarkdown(text) if kind == "markdown" else browser.setHtml(text)
    return TextPane(browser)


def code_pane(text, path):
    result = code_html(text, Path(path).name)
    if not result:
        return plain_text_pane(text)
    html, background = result
    browser = QTextBrowser()
    browser.setFont(QFontDatabase.systemFont(QFontDatabase.FixedFont))
    browser.setLineWrapMode(QTextBrowser.LineWrapMode.NoWrap)
    palette = browser.palette()
    palette.setColor(QPalette.Base, QColor(background))  # the theme's own background, so colors stay readable
    browser.setPalette(palette)
    browser.setHtml(html)
    return TextPane(browser)


# ---------------------------------------------------------------- spreadsheets

def excel_color(color):
    """openpyxl Color -> QColor, or None. Theme/indexed colors are skipped."""
    # ponytail: only explicit RGB colors; resolve theme colors via the workbook theme if they matter
    if color is not None and color.type == "rgb" and isinstance(color.rgb, str) and len(color.rgb) == 8:
        return QColor("#" + color.rgb[2:])
    return None


def format_cell(value, number_format):
    """Show a cell value roughly like Excel does, in the system locale."""
    locale = QLocale()
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime.datetime):
        return locale.toString(value, QLocale.ShortFormat)
    if isinstance(value, datetime.date):
        return locale.toString(QDate(value.year, value.month, value.day), QLocale.ShortFormat)
    if not isinstance(value, (int, float)):
        return str(value)
    if number_format == "General":
        return locale.toString(float(value), "g", 10)
    # ponytail: covers decimals, thousands, %, currency and literal text; not conditional or fraction formats
    section = number_format.split(";")[1 if value < 0 and ";" in number_format else 0]
    section = re.sub(r"\[\$([^-\]]*)[^\]]*\]", r"\1", section)  # [$€-410] -> €
    section = re.sub(r"\[[^\]]*\]|_.|\*.|\\|\"", "", section)  # colors, padding, fill, escapes, quotes
    match = re.search(r"[#0?][#0?,.]*", section)
    if not match:
        return section
    digits = match.group()
    decimals = len(digits.split(".")[1]) if "." in digits else 0
    number = abs(value) * 100 if "%" in section else abs(value)
    # Group in Python, not QLocale: Qt skips the separator on 4-digit numbers in some locales
    # (Italian "2490,00"), Excel doesn't ("2.490,00")
    text = f"{number:{',' if ',' in digits else ''}.{decimals}f}"
    text = text.replace(",", "\0").replace(".", locale.decimalPoint()).replace("\0", locale.groupSeparator())
    sign = "-" if value < 0 and ";" not in number_format else ""
    return sign + section[:match.start()] + text + section[match.end():]


class SheetView(QWidget):
    """Spreadsheet grid: sheet tabs, column letters, cell styles, a formula bar for the clicked cell, search."""

    def __init__(self, path):
        super().__init__()
        formulas = openpyxl.load_workbook(path)  # cell.value is "=SUM(...)" for formula cells
        values = openpyxl.load_workbook(path, data_only=True)  # cached results Excel saved

        self.ref = QLabel("A1")
        self.ref.setMinimumWidth(60)
        self.formula = QLineEdit(readOnly=True, placeholderText="Click a cell to inspect it")
        show_hidden = tool_button("view-visible", "Show rows and columns the file hides or collapses",
                                  self.show_hidden, text="Show hidden", checkable=True)
        self.tabs = QTabWidget(tabPosition=QTabWidget.South, documentMode=True)
        self.search = SearchBar(self.tabs, self)
        self.search.changed.connect(lambda text: self.find(text, 0))
        self.search.step.connect(lambda delta: self.find(self.search.edit.text(), delta))
        self.matches, self.match = [], -1

        bar = QHBoxLayout()
        bar.addWidget(self.ref)
        bar.addWidget(self.formula, 1)
        bar.addWidget(self.search.button())
        bar.addWidget(show_hidden)

        for ws in formulas.worksheets:
            self.tabs.addTab(self.build_table(ws, values[ws.title]), ws.title)
        self.tabs.currentChanged.connect(lambda _: self.find(self.search.edit.text(), 0))
        layout = QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.addLayout(bar)
        layout.addWidget(self.tabs, 1)

    def build_table(self, ws_formulas, ws):
        rows, cols = min(ws.max_row, MAX_ROWS), min(ws.max_column, MAX_COLS)
        table = QTableWidget(rows, cols)
        table.setHorizontalHeaderLabels([get_column_letter(c) for c in range(1, cols + 1)])
        table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        table.setShowGrid(True)  # always: seeing the cells is the point of this view
        table.hidden_rows, table.hidden_cols = [], []  # what the file hides, for the "Show hidden" toggle

        default_width = ws.sheet_format.defaultColWidth or 8.43
        for c in range(cols):
            table.setColumnWidth(c, int(default_width * 7 + 5))  # Excel width unit ~ 7 px at 96 dpi
        for dim in ws.column_dimensions.values():
            for c in range(dim.min, min(dim.max, cols) + 1):
                if dim.width:
                    table.setColumnWidth(c - 1, int(dim.width * 7 + 5))
                if dim.hidden:
                    table.setColumnHidden(c - 1, True)
                    table.hidden_cols.append(c - 1)
        for r, dim in ws.row_dimensions.items():
            if r <= rows:
                if dim.ht:
                    table.setRowHeight(r - 1, int(dim.ht * 96 / 72))  # points -> px
                if dim.hidden:
                    table.setRowHidden(r - 1, True)
                    table.hidden_rows.append(r - 1)

        for row in ws.iter_rows(max_row=rows, max_col=cols):
            for cell in row:
                raw = ws_formulas.cell(cell.row, cell.column).value
                if cell.value is None and raw is None and not cell.has_style:
                    continue
                item = QTableWidgetItem(format_cell(cell.value, cell.number_format))
                item.setData(Qt.UserRole, raw)  # formula or raw value, for the formula bar
                font = item.font()
                font.setBold(bool(cell.font.b))
                font.setItalic(bool(cell.font.i))
                font.setUnderline(bool(cell.font.u))
                if cell.font.sz:
                    font.setPointSizeF(float(cell.font.sz))
                item.setFont(font)
                if color := excel_color(cell.font.color):
                    item.setForeground(color)
                if cell.fill.fill_type == "solid" and (color := excel_color(cell.fill.fgColor)):
                    item.setBackground(color)
                align = {"center": Qt.AlignHCenter, "right": Qt.AlignRight, "left": Qt.AlignLeft}.get(
                    cell.alignment.horizontal,
                    Qt.AlignRight if isinstance(cell.value, (int, float)) else Qt.AlignLeft)  # Excel's default
                item.setTextAlignment(align | Qt.AlignVCenter)
                table.setItem(cell.row - 1, cell.column - 1, item)

        for merged in ws_formulas.merged_cells.ranges:
            if merged.min_row <= rows and merged.min_col <= cols:
                table.setSpan(merged.min_row - 1, merged.min_col - 1,
                              min(merged.max_row, rows) - merged.min_row + 1,
                              min(merged.max_col, cols) - merged.min_col + 1)
        table.currentCellChanged.connect(lambda r, c, *_: self.inspect(table, r, c))
        return table

    def show_hidden(self, on):
        for i in range(self.tabs.count()):
            table = self.tabs.widget(i)
            for r in table.hidden_rows:
                table.setRowHidden(r, not on)
            for c in table.hidden_cols:
                table.setColumnHidden(c, not on)

    def inspect(self, table, row, col):
        self.ref.setText(f"{get_column_letter(col + 1)}{row + 1}")
        item = table.item(row, col)
        raw = item.data(Qt.UserRole) if item else None
        self.formula.setText("" if raw is None else str(raw))

    def find(self, text, delta):
        table = self.tabs.currentWidget()
        if delta == 0:
            found = table.findItems(text, Qt.MatchContains) if text else []
            self.matches = sorted(found, key=lambda item: (item.row(), item.column()))
            self.match = -1
        if not self.matches:
            self.search.label.setText("no matches" if text else "")
            return
        self.match = (self.match + (delta or 1)) % len(self.matches)
        item = self.matches[self.match]
        table.setCurrentItem(item)
        table.scrollToItem(item)
        self.search.label.setText(f"{self.match + 1} of {len(self.matches)}")


# ---------------------------------------------------------------- 3D, fonts, archives, folders, audio

class ModelView(QWidget):
    """Rotatable 3D view with fly mode, wireframe, auto-spin, floor grid and model stats."""
    flying_changed = Signal(bool)
    toggles = {"wireframe": False, "spin": False, "grid": True}  # stay as you left them across files

    def __init__(self, path):
        super().__init__()
        CACHE.mkdir(parents=True, exist_ok=True)
        qml = CACHE / "viewer.qml"
        if not qml.exists() or qml.read_text() != VIEWER_QML:
            qml.write_text(VIEWER_QML)
        self.quick = QQuickWidget()
        self.quick.setResizeMode(QQuickWidget.SizeRootObjectToView)
        palette = QApplication.palette()
        self.quick.setInitialProperties({"source": QUrl.fromLocalFile(str(path)), **ModelView.toggles,
                                         "background": palette.window().color(),
                                         "textColor": palette.text().color()})
        self.quick.setSource(QUrl.fromLocalFile(str(qml)))
        self.flying = False
        self.quick.rootObject().animationsReady.connect(self.show_animations)
        in_background(lambda: animated_scene(Path(path)), self, self.load_scene)

        self.stats = QLabel()
        in_background(lambda: model_stats(Path(path)), self.stats, lambda text: self.stats.setText(text or ""))
        self.animation = QComboBox(visible=False, toolTip="Animation")
        self.animation.currentIndexChanged.connect(lambda _: self.play(True))
        self.play_button = tool_button("media-playback-pause", "Pause the animation (P)",
                                       lambda on: self.play(not on), text="Pause", checkable=True)
        self.play_button.setVisible(False)
        QShortcut("P", self, self.play_button.toggle)
        self.fly_button = tool_button("transform-move", "Walk inside the model", self.set_flying,
                                      text="Fly", checkable=True)
        QShortcut("Home", self, self.recenter)
        self.stats.setEnabled(False)  # dimmed: information, not a control
        self.animation_tray = Tray(self.animation, self.play_button)
        self.animation_tray.hide()  # shown once balsam finds animations
        bar = toolbar(self.stats, self.animation_tray,
                      (tool_button("zoom-fit-best", "Back to the starting view (Home)", self.recenter,
                                   text="Recenter"), self.fly_button),
                      (tool_button("draw-polyline", "Show the triangles", lambda on: self.toggle("wireframe", on),
                                   text="Wireframe", checkable=True, checked=ModelView.toggles["wireframe"]),
                       tool_button("object-rotate-right", "Turn the model slowly", lambda on: self.toggle("spin", on),
                                   text="Spin", checkable=True, checked=ModelView.toggles["spin"]),
                       tool_button("view-grid", "Floor grid", lambda on: self.toggle("grid", on),
                                   text="Grid", checkable=True, checked=ModelView.toggles["grid"])))
        layout = QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.addLayout(bar)
        layout.addWidget(self.quick, 1)

    def load_scene(self, qml):
        if qml:
            self.quick.rootObject().setProperty("scene", QUrl.fromLocalFile(str(qml)))

    def show_animations(self, names):
        self.animation.blockSignals(True)  # filling the list must not restart playback
        self.animation.clear()
        self.animation.addItems(list(names))
        self.animation.blockSignals(False)
        self.animation.setVisible(len(names) > 1)
        self.play_button.setVisible(bool(names))
        self.animation_tray.setVisible(bool(names))

    def play(self, playing):
        self.play_button.blockSignals(True)
        self.play_button.setChecked(not playing)
        self.play_button.blockSignals(False)
        self.play_button.setText("Pause" if playing else "Play")
        self.play_button.setIcon(QIcon.fromTheme("media-playback-pause" if playing else "media-playback-start"))
        QMetaObject.invokeMethod(self.quick.rootObject(), "playAnimation",
                                 Q_ARG("QVariant", max(self.animation.currentIndex(), 0)),
                                 Q_ARG("QVariant", playing))

    def toggle(self, name, on):
        ModelView.toggles[name] = on
        self.quick.rootObject().setProperty(name, on)

    def recenter(self):
        """Land if flying, then frame the whole model again from the starting angle."""
        self.set_flying(False)
        QMetaObject.invokeMethod(self.quick.rootObject(), "fit")

    def set_flying(self, on):
        self.flying = on
        self.fly_button.setChecked(on)
        self.quick.rootObject().setProperty("flying", on)
        self.quick.setFocusPolicy(Qt.StrongFocus if on else Qt.NoFocus)
        if on:
            self.quick.setFocus()
        self.flying_changed.emit(on)


class FontView(QWidget):
    """Font sample in several sizes. The font is loaded for this preview only."""

    def __init__(self, path):
        super().__init__()
        font_id = QFontDatabase.addApplicationFont(str(path))
        if font_id < 0:
            raise ValueError("Qt can't load this font")
        self.destroyed.connect(lambda: QFontDatabase.removeApplicationFont(font_id))
        family = QFontDatabase.applicationFontFamilies(font_id)[0]
        styles = ", ".join(QFontDatabase.styles(family))
        sample = "The quick brown fox jumps over the lazy dog"
        html = [f"<h2 style=\"font-family:'{escape(family)}'\">{escape(family)}</h2>",
                f"<p style='opacity:.7'>{escape(styles)}</p>"]
        for size in (12, 18, 24, 36, 48, 72):
            html.append(f"<p style=\"font-family:'{escape(family)}'; font-size:{size}pt\">{sample}</p>")
        html.append(f"<p style=\"font-family:'{escape(family)}'; font-size:20pt\">"
                    "ABCDEFGHIJKLMNOPQRSTUVWXYZ<br>abcdefghijklmnopqrstuvwxyz<br>0123456789 !?&amp;@€$%()</p>")
        browser = QTextBrowser()
        browser.setHtml("".join(html))
        layout = QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.addWidget(browser)


class ArchiveView(QWidget):
    """File tree of an archive, read on a worker thread so big archives don't freeze the window."""

    def __init__(self, path):
        super().__init__()
        self.archive_size = Path(path).stat().st_size
        self.summary = QLabel("Reading archive…")
        self.tree = QTreeWidget()
        self.tree.setHeaderLabels(["Name", "Size", "Modified"])
        layout = QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.addWidget(self.summary)
        layout.addWidget(self.tree, 1)
        in_background(lambda: list_archive(path), self, self.fill)

    def fill(self, entries):
        if entries is None:
            self.summary.setText("Can't read this archive (is 7z installed?)")
            return
        fmt = QLocale().formattedDataSize
        nodes = {}

        def node(name):  # the tree item for a folder path, creating its parents as needed
            if name not in nodes:
                parent, _, leaf = name.rpartition("/")
                item = QTreeWidgetItem([leaf])
                item.setIcon(0, QIcon.fromTheme("folder"))
                (node(parent).addChild if parent else self.tree.addTopLevelItem)(item)
                nodes[name] = item
            return nodes[name]

        files = total = 0
        mimes = QMimeDatabase()
        for name, size, modified, is_dir in sorted(entries, key=lambda e: e[0].lower()):
            name = name.strip("/")
            if not name:
                continue
            if is_dir:
                node(name)
                continue
            files += 1
            total += size or 0
            parent, _, leaf = name.rpartition("/")
            item = QTreeWidgetItem([leaf, fmt(size) if size is not None else "",
                                    modified.strftime("%Y-%m-%d %H:%M") if modified else ""])
            item.setIcon(0, QIcon.fromTheme(mimes.mimeTypeForFile(leaf, QMimeDatabase.MatchExtension).iconName()))
            item.setTextAlignment(1, Qt.AlignRight | Qt.AlignVCenter)
            (node(parent).addChild if parent else self.tree.addTopLevelItem)(item)
        if self.tree.topLevelItemCount() <= 3:
            self.tree.expandToDepth(0)
        self.tree.resizeColumnToContents(0)
        cut = f" (first {MAX_ENTRIES})" if len(entries) >= MAX_ENTRIES else ""
        self.summary.setText(f"{files} files{cut}, {len(nodes)} folders  ·  {fmt(total)} unpacked, "
                             f"{fmt(self.archive_size)} packed")


class FolderView(QWidget):
    """Folder contents as icons/thumbnails, with the total size counted on a worker thread."""

    def __init__(self, path, on_open):
        super().__init__()
        self.summary = QLabel("Counting…")
        grid = QListWidget(viewMode=QListWidget.IconMode, movement=QListWidget.Static,
                           resizeMode=QListWidget.Adjust, wordWrap=True)
        grid.setIconSize(QSize(96, 96))
        grid.setGridSize(QSize(130, 140))
        grid.itemActivated.connect(lambda item: on_open(item.data(Qt.UserRole)))
        mimes = QMimeDatabase()
        children = sorted_like_dolphin(path)
        for child in children[:FOLDER_ITEMS]:
            thumb = cached_thumbnail(child, ("large", "normal"))
            icon = (QIcon(QPixmap.fromImage(thumb)) if thumb else
                    QIcon.fromTheme(mimes.mimeTypeForFile(str(child)).iconName()))
            item = QListWidgetItem(icon, child.name, grid)
            item.setData(Qt.UserRole, child)
            item.setToolTip(child.name)
        more = QLabel(f"… and {len(children) - FOLDER_ITEMS} more" if len(children) > FOLDER_ITEMS else "")
        layout = QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.addWidget(self.summary)
        layout.addWidget(grid, 1)
        layout.addWidget(more)
        self.direct = len(children)
        in_background(lambda: folder_summary(path), self, self.show_summary)

    def show_summary(self, result):
        files, folders, size, complete = result
        more = "" if complete else " (counting stopped after 3 s)"
        self.summary.setText(f"{self.direct} items here  ·  {files} files in {folders} subfolders  ·  "
                             f"{'' if complete else 'at least '}{QLocale().formattedDataSize(size)}{more}")


def clock(ms):
    """1:05 or 1:02:05"""
    seconds = max(ms, 0) // 1000
    hours, rest = divmod(seconds, 3600)
    return f"{hours}:{rest // 60:02d}:{rest % 60:02d}" if hours else f"{rest // 60}:{rest % 60:02d}"


class SeekSlider(QSlider):
    """A slider that jumps to where you click, instead of stepping toward it."""

    def mousePressEvent(self, event):
        if event.button() == Qt.LeftButton:
            option = QStyleOptionSlider()
            self.initStyleOption(option)
            groove = self.style().subControlRect(QStyle.CC_Slider, option, QStyle.SC_SliderGroove, self)
            handle = self.style().subControlRect(QStyle.CC_Slider, option, QStyle.SC_SliderHandle, self)
            x = event.position().x() - groove.x() - handle.width() / 2
            self.setValue(QStyle.sliderValueFromPosition(self.minimum(), self.maximum(), round(x),
                                                         groove.width() - handle.width()))
            self.sliderMoved.emit(self.value())
        super().mousePressEvent(event)  # then the handle under the mouse can be dragged on


class MediaPane(QWidget):
    """Video or audio with a player bar: play/pause, timeline, time and volume.
    Keys: K play/pause, J / L back / forward 5 s, M mute."""

    def __init__(self, content, player):
        super().__init__()
        self.player = player
        audio = player.audioOutput()
        self.play = tool_button("media-playback-start", "Play / pause (K)", self.toggle)
        self.elapsed = QLabel()
        self.total = QLabel()
        numbers = self.elapsed.font()
        numbers.setFeature(QFont.Tag("tnum"), 1)  # equal-width digits: the time doesn't wobble while playing
        for label in (self.elapsed, self.total):
            label.setFont(numbers)
            label.setEnabled(False)  # dimmed: secondary to the timeline
        self.seek = SeekSlider(Qt.Horizontal)
        self.seek.sliderMoved.connect(player.setPosition)
        self.mute = tool_button("audio-volume-high", "Mute (M)", lambda: audio.setMuted(not audio.isMuted()))
        self.volume = QSlider(Qt.Horizontal, maximum=100, toolTip="Volume")
        self.volume.setFixedWidth(90)
        self.volume.valueChanged.connect(lambda v: (audio.setVolume(QtAudio.convertVolume(
            v / 100, QtAudio.LogarithmicVolumeScale, QtAudio.LinearVolumeScale)), audio.setMuted(False)))

        bar = QHBoxLayout()
        bar.setSpacing(8)
        bar.addWidget(Tray(self.play))
        bar.addWidget(self.elapsed)
        bar.addWidget(self.seek, 1)
        bar.addWidget(self.total)
        sound = Tray(self.mute, self.volume)
        sound.layout().setContentsMargins(3, 3, 9, 3)  # room for the handle at 100%
        bar.addWidget(sound)
        layout = QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.addWidget(content, 1)
        layout.addLayout(bar)
        for w in self.findChildren(QWidget):
            if w is not content:
                w.setFocusPolicy(Qt.NoFocus)  # Space and the arrows keep flipping files

        player.playbackStateChanged.connect(self.show_state)
        player.durationChanged.connect(lambda ms: (self.seek.setMaximum(ms), self.total.setText(clock(ms))))
        player.positionChanged.connect(self.show_position)
        audio.mutedChanged.connect(self.show_volume)
        audio.volumeChanged.connect(self.show_volume)
        QShortcut("K", self, self.toggle)
        QShortcut("M", self, self.mute.click)
        QShortcut("J", self, lambda: player.setPosition(max(player.position() - 5000, 0)))
        QShortcut("L", self, lambda: player.setPosition(min(player.position() + 5000, player.duration())))
        self.seek.setMaximum(player.duration())
        self.total.setText(clock(player.duration()))
        self.show_position(player.position())
        self.show_state()
        self.show_volume()

    def toggle(self):
        self.player.pause() if self.player.isPlaying() else self.player.play()

    def show_state(self, *_):
        playing = self.player.isPlaying()
        self.play.setIcon(QIcon.fromTheme("media-playback-pause" if playing else "media-playback-start"))

    def show_position(self, ms):
        if not self.seek.isSliderDown():  # don't yank the handle out of your hand while dragging
            self.seek.setValue(ms)
        self.elapsed.setText(clock(ms))

    def show_volume(self, *_):
        audio = self.player.audioOutput()
        self.volume.blockSignals(True)  # following the volume must not set it again
        # sliders feel linear to the ear only on a logarithmic scale
        self.volume.setValue(round(100 * QtAudio.convertVolume(audio.volume(), QtAudio.LinearVolumeScale,
                                                                QtAudio.LogarithmicVolumeScale)))
        self.volume.blockSignals(False)
        level = 0 if audio.isMuted() else self.volume.value()
        self.mute.setIcon(QIcon.fromTheme("audio-volume-muted" if level == 0 else "audio-volume-low" if level < 34
                                          else "audio-volume-medium" if level < 67 else "audio-volume-high"))
        self.mute.setToolTip("Unmute (M)" if audio.isMuted() else "Mute (M)")


class AudioView(QWidget):
    """Cover art and tags of the playing file. Click to pause."""

    def __init__(self, player, mime, path):
        super().__init__()
        self.player, self.mime, self.path = player, mime, path
        self.cover = QLabel(alignment=Qt.AlignCenter)
        self.title = QLabel(alignment=Qt.AlignCenter, wordWrap=True)
        font = self.title.font()
        font.setPointSizeF(font.pointSizeF() * 1.6)
        font.setBold(True)
        self.title.setFont(font)
        self.details = QLabel(alignment=Qt.AlignCenter, wordWrap=True)
        layout = QVBoxLayout(self)
        layout.addStretch()
        for w in (self.cover, self.title, self.details):
            layout.addWidget(w)
        layout.addStretch()
        player.metaDataChanged.connect(self.refresh)
        player.durationChanged.connect(self.refresh)
        self.refresh()

    def refresh(self):
        meta = self.player.metaData()
        cover = meta.value(QMediaMetaData.CoverArtImage) or meta.value(QMediaMetaData.ThumbnailImage)
        if isinstance(cover, QImage) and not cover.isNull():
            self.cover.setPixmap(QPixmap.fromImage(cover).scaled(300, 300, Qt.KeepAspectRatio,
                                                                 Qt.SmoothTransformation))
        else:
            self.cover.setPixmap(QIcon.fromTheme(self.mime.iconName()).pixmap(128))
        self.title.setText(meta.stringValue(QMediaMetaData.Title) or Path(self.path).stem)
        artist = (meta.stringValue(QMediaMetaData.ContributingArtist)
                  or meta.stringValue(QMediaMetaData.AlbumArtist))
        seconds = self.player.duration() // 1000
        parts = [artist, meta.stringValue(QMediaMetaData.AlbumTitle),
                 f"{seconds // 60}:{seconds % 60:02d}" if seconds else ""]
        self.details.setText("  ·  ".join(p for p in parts if p))

    def mousePressEvent(self, event):
        self.player.pause() if self.player.isPlaying() else self.player.play()


# ---------------------------------------------------------------- details panel data

def image_rows(path):
    """Size and format from the file header (no full decode), plus camera details from EXIF."""
    reader = QImageReader(str(path))
    size = reader.size()
    if not size.isValid():
        return []
    rows = [("Dimensions", f"{size.width()} × {size.height()} px"),
            ("Megapixels", f"{size.width() * size.height() / 1e6:.1f}"),
            ("Format", bytes(reader.format()).decode().upper() or "?")]
    if not Image:
        return rows
    try:
        with Image.open(path) as im:
            exif = im.getexif()
            tag = {ExifTags.TAGS.get(k, k): v for k, v in
                   [*exif.items(), *exif.get_ifd(ExifTags.IFD.Exif).items()]}
            gps = {ExifTags.GPSTAGS.get(k, k): v for k, v in exif.get_ifd(ExifTags.IFD.GPSInfo).items()}
            if profile := im.info.get("icc_profile"):
                # the profile's description sits after a 'desc' tag; good enough for "sRGB" / "Display P3"
                name = re.search(rb"desc.{8}([\x20-\x7e]{3,60})", profile, re.S)
                rows.append(("Color profile", name.group(1).decode().strip() if name else "embedded"))
    except Exception:  # Pillow can't open every format Qt can (kra, xcf...): skip camera details
        return rows

    def text(key):
        return str(tag[key]).strip("\x00 ")
    if camera := " ".join(text(k) for k in ("Make", "Model") if k in tag):
        rows.append(("Camera", camera))
    if "LensModel" in tag:
        rows.append(("Lens", text("LensModel")))
    if "DateTimeOriginal" in tag or "DateTime" in tag:
        rows.append(("Taken", text("DateTimeOriginal" if "DateTimeOriginal" in tag else "DateTime")))
    if tag.get("ExposureTime"):
        t = float(tag["ExposureTime"])
        rows.append(("Exposure", f"1/{round(1 / t)} s" if t < 1 else f"{t:g} s"))
    if "FNumber" in tag:
        rows.append(("Aperture", f"f/{float(tag['FNumber']):g}"))
    if "ISOSpeedRatings" in tag:
        rows.append(("ISO", text("ISOSpeedRatings")))
    if "FocalLength" in tag:
        rows.append(("Focal length", f"{float(tag['FocalLength']):g} mm"))
    if "Software" in tag:
        rows.append(("Software", text("Software")))
    if "GPSLatitude" in gps and "GPSLongitude" in gps:
        def degrees(dms, ref):
            value = float(dms[0]) + float(dms[1]) / 60 + float(dms[2]) / 3600
            return -value if ref in ("S", "W") else value
        lat = degrees(gps["GPSLatitude"], gps.get("GPSLatitudeRef"))
        lon = degrees(gps["GPSLongitude"], gps.get("GPSLongitudeRef"))
        rows.append(("Location", (f"{lat:.5f}, {lon:.5f}",
                                  f"https://www.openstreetmap.org/?mlat={lat}&mlon={lon}#map=15/{lat}/{lon}")))
    return rows


def media_rows(path):
    """Codec, resolution, bitrate and tags from ffprobe."""
    if not shutil.which("ffprobe"):
        return []
    out = subprocess.run(["ffprobe", "-v", "error", "-print_format", "json", "-show_format", "-show_streams",
                          str(path)], capture_output=True, text=True, timeout=15).stdout
    info = json.loads(out or "{}")
    fmt = info.get("format", {})
    rows = []
    if seconds := float(fmt.get("duration", 0) or 0):
        rows.append(("Duration", f"{int(seconds // 3600)}:{int(seconds % 3600 // 60):02d}:{seconds % 60:04.1f}"
                     if seconds >= 3600 else f"{int(seconds // 60)}:{seconds % 60:04.1f}"))
    if fmt.get("format_long_name"):
        rows.append(("Container", fmt["format_long_name"]))
    if fmt.get("bit_rate"):
        rows.append(("Bitrate", f"{int(fmt['bit_rate']) // 1000} kb/s"))
    for s in info.get("streams", []):
        codec = s.get("codec_long_name") or s.get("codec_name", "?")
        if s.get("codec_type") == "video" and not s.get("disposition", {}).get("attached_pic"):
            num, _, den = s.get("avg_frame_rate", "0/1").partition("/")
            fps = float(num) / float(den) if float(den or 0) else 0
            rows.append(("Video", f"{codec}, {s.get('width')} × {s.get('height')}" + (f", {fps:.3g} fps" if fps else "")))
        elif s.get("codec_type") == "audio":
            layout = s.get("channel_layout") or f"{s.get('channels')} channels"
            rows.append(("Audio", f"{codec}, {int(s.get('sample_rate', 0)) / 1000:g} kHz, {layout}"))
    tags = {k.lower(): v for k, v in fmt.get("tags", {}).items()}
    for key in ("title", "artist", "album", "date", "genre", "encoder"):
        if tags.get(key):
            rows.append((key.capitalize(), tags[key]))
    return rows


def pdf_rows(path):
    if not shutil.which("pdfinfo"):
        return []
    out = subprocess.run(["pdfinfo", str(path)], capture_output=True, text=True, timeout=15).stdout
    fields = dict(line.split(":", 1) for line in out.splitlines() if ":" in line)
    wanted = [("Title", "Title"), ("Author", "Author"), ("Creator", "Made with"), ("Producer", "Producer"),
              ("CreationDate", "Created"), ("Pages", "Pages"), ("Page size", "Page size"),
              ("PDF version", "PDF version"), ("Encrypted", "Encrypted")]
    return [(label, fields[key].strip()) for key, label in wanted if fields.get(key, "").strip()]


def office_rows(path):
    """Author, dates and counts that Office (docProps) and OpenDocument (meta.xml) files carry."""
    z = zipfile.ZipFile(path)
    names = z.namelist()
    found = {}
    for part in ("docProps/core.xml", "docProps/app.xml", "meta.xml"):
        if part in names:
            for el in ElementTree.fromstring(z.read(part)).iter():
                tag = el.tag.rsplit("}", 1)[-1]
                if el.text and el.text.strip():
                    found.setdefault(tag, el.text.strip())
                if tag == "document-statistic":  # OpenDocument keeps counts as attributes
                    for k, v in el.attrib.items():
                        found.setdefault(k.rsplit("}", 1)[-1], v)
    wanted = [("title", "Title"), ("creator", "Last author"), ("initial-creator", "Author"),
              ("lastModifiedBy", "Last edited by"), ("created", "Created"), ("creation-date", "Created"),
              ("modified", "Modified"), ("date", "Modified"), ("Pages", "Pages"), ("page-count", "Pages"),
              ("Words", "Words"), ("word-count", "Words"), ("Slides", "Slides"), ("Application", "Made with"),
              ("generator", "Made with")]
    rows, seen = [], set()
    for key, label in wanted:
        if key in found and label not in seen:
            seen.add(label)
            rows.append((label, found[key].replace("T", " ").rstrip("Z") if label in ("Created", "Modified")
                         else found[key]))
    return rows


def archive_rows(path):
    entries = list_archive(path)
    if entries is None:
        return []
    files = [e for e in entries if not e[3]]
    unpacked = sum(e[1] or 0 for e in files)
    rows = [("Files", f"{len(files)}"), ("Unpacked", QLocale().formattedDataSize(unpacked))]
    if unpacked:
        rows.append(("Compression", f"{100 - Path(path).stat().st_size * 100 / unpacked:.0f}% smaller"))
    if zipfile.is_zipfile(path):
        locked = any(info.flag_bits & 1 for info in zipfile.ZipFile(path).infolist())  # bit 0 = encrypted
        rows.append(("Password", "yes" if locked else "no"))
    return rows


def text_rows(path):
    data = Path(path).read_bytes()[:TEXT_LIMIT]
    try:
        data.decode("utf-8")
        encoding = "UTF-8 with BOM" if data.startswith(b"\xef\xbb\xbf") else "UTF-8"
    except UnicodeDecodeError:
        encoding = "not UTF-8 (probably Windows-1252 / Latin-1)"
    crlf = data.count(b"\r\n")
    lf = data.count(b"\n") - crlf
    endings = "Windows (CRLF)" if crlf and not lf else "mixed" if crlf else "Linux (LF)"
    cut = "+" if len(data) == TEXT_LIMIT else ""
    return [("Lines", f"{data.count(b'\n') + (not data.endswith(b'\n') and bool(data))}{cut}"),
            ("Encoding", encoding), ("Line endings", endings)]


def font_rows(path):
    """Version, designer and license from the font's 'name' table (TrueType / OpenType only)."""
    data = Path(path).read_bytes()
    if data[:4] not in (b"\x00\x01\x00\x00", b"OTTO", b"true"):
        return []
    tables = {data[12 + 16 * i:16 + 16 * i]: struct.unpack_from(">II", data, 20 + 16 * i)
              for i in range(struct.unpack_from(">H", data, 4)[0])}  # tag -> (offset, length)
    rows = []
    if b"maxp" in tables:
        rows.append(("Glyphs", str(struct.unpack_from(">H", data, tables[b"maxp"][0] + 4)[0])))
    if b"name" not in tables:
        return rows
    base = tables[b"name"][0]
    count, strings = struct.unpack_from(">HH", data, base + 2)
    names = {}
    for i in range(count):
        platform, _, _, name_id, length, offset = struct.unpack_from(">HHHHHH", data, base + 6 + 12 * i)
        raw = data[base + strings + offset:base + strings + offset + length]
        if platform in (0, 3):  # Unicode / Windows: UTF-16 big endian
            names.setdefault(name_id, raw.decode("utf-16-be", errors="replace"))
        elif platform == 1:  # old Macintosh names
            names.setdefault(name_id, raw.decode("latin-1"))
    for name_id, label in ((5, "Version"), (9, "Designer"), (8, "Foundry"), (0, "Copyright"),
                           (13, "License"), (14, "License URL")):
        if names.get(name_id, "").strip():
            value = names[name_id].strip()
            rows.append((label, value if len(value) < 300 else value[:300] + "…"))
    return rows


def file_details(path, mime=None):
    """[(section, [(label, value)])] for the Details panel. Runs on a worker thread.
    Pass `mime` from the GUI thread: a QMimeDatabase's first use on a new thread reloads the
    whole type database (~0.3 s)."""
    path = Path(path)
    st = path.stat()
    mime = mime or QMimeDatabase().mimeTypeForFile(str(path))
    when = lambda t: datetime.datetime.fromtimestamp(t).strftime("%Y-%m-%d %H:%M")
    try:
        owner = f"{pwd.getpwuid(st.st_uid).pw_name} : {grp.getgrgid(st.st_gid).gr_name}"
    except KeyError:
        owner = f"{st.st_uid} : {st.st_gid}"
    general = [("Folder", str(path.parent)),  # the type is in the panel's header
               ("Modified", when(st.st_mtime))]
    if path.is_file():
        general.insert(0, ("Size", f"{QLocale().formattedDataSize(st.st_size)} ({QLocale().toString(st.st_size)} bytes)"))
    if birth := getattr(st, "st_birthtime", None):
        general.append(("Created", when(birth)))
    general += [("Opened", when(st.st_atime)), ("Owner", owner), ("Permissions", stat.filemode(st.st_mode))]
    if path.is_symlink():
        general.append(("Links to", os.readlink(path)))
    sections = [("General", general)]

    suffix, name = path.suffix.lower(), mime.name()
    readers = []
    if path.is_dir():
        files, folders, size, complete = folder_summary(path)
        readers.append(("Contents", lambda: [("Files", str(files)), ("Subfolders", str(folders)),
                                             ("Total size", ("" if complete else "at least ")
                                              + QLocale().formattedDataSize(size))]))
    elif name.startswith(("audio/", "video/")):
        readers.append(("Media", lambda: media_rows(path)))
    elif name == "application/pdf":
        readers.append(("Document", lambda: pdf_rows(path)))
    elif suffix in OFFICE and zipfile.is_zipfile(path):
        readers.append(("Document", lambda: office_rows(path)))
    elif suffix in MODELS:
        readers.append(("Model", lambda: [tuple(p.split(" ", 1)[::-1]) for p in
                                          (model_stats(path) or "").split("  ·  ") if " " in p]))
    elif suffix in FONTS or name.startswith("font/"):
        readers.append(("Font", lambda: font_rows(path)))
    elif name in ARCHIVES:
        readers.append(("Archive", lambda: archive_rows(path)))
    elif QImageReader(str(path)).canRead():
        readers.append(("Image", lambda: image_rows(path)))
    elif mime.inherits("text/plain"):
        readers.append(("Text", lambda: text_rows(path)))
    for title, read in readers:
        try:
            if rows := read():
                sections.append((title, rows))
        except Exception:  # a damaged file must not break the panel; General still shows
            pass
    return sections


def breakable(text):
    """Qt's rich text can't break long paths or codec names: add invisible break points."""
    text = re.sub(r"([/_.,-])", "\\1\u200b", text)
    return re.sub(r"(\S{18})(?=\S)", "\\1\u200b", text)


def details_html(sections):
    """All sections in one two-column table, so the columns line up. A (text, url) value is a link."""
    html = ["<table cellspacing=0 cellpadding=2 width=100%>"]
    for i, (title, rows) in enumerate(sections):
        top = 4 if i == 0 else 14
        html.append(f"<tr><td colspan=2 style='padding-top:{top}px'><b>{escape(title)}</b></td></tr>")
        for label, value in rows:
            shown = (f"<a href='{escape(value[1])}'>{escape(value[0])}</a>" if isinstance(value, tuple)
                     else escape(breakable(str(value))))
            html.append(f"<tr><td width=36% style='opacity:.65'>{escape(label)}</td><td>{shown}</td></tr>")
    html.append("</table>")
    return "".join(html)


# ---------------------------------------------------------------- local AI (runs on this PC, nothing is uploaded)

MODELS_DIR = Path.home() / ".local/share/peek/models"  # not in the cache: cleanup must never delete these
AI_MODELS = {  # name: (file, url, size in bytes) -- downloaded the first time a tool is used
    # Segment Anything 2.1 (Hiera Base+), zipped: the encoder reads a picture once, the decoder turns each
    # click into a mask. Unpacked by sam2_files.
    "sam2": ("sam2.1_hiera_base_plus_20260221.zip",
             "https://huggingface.co/vietanhdev/segment-anything-2.1-onnx-models/resolve/main/"
             "sam2.1_hiera_base_plus_20260221.zip",
             271998747),
    # BiRefNet retrained on anime and illustrations (ToonOut): outlines of drawn subjects
    "toonout": ("birefnet-toonout.onnx",
                "https://huggingface.co/sprited/birefnet-toonout-onnx/resolve/main/birefnet-toonout.onnx",
                974344828),
    "birefnet": ("BiRefNet-general-bb_swin_v1_tiny-epoch_232.onnx",
                 "https://github.com/danielgatis/rembg/releases/download/v0.0.0/"
                 "BiRefNet-general-bb_swin_v1_tiny-epoch_232.onnx",
                 224005088),
}
_model_locks = {name: threading.Lock() for name in AI_MODELS}  # two tools asking at once download once


class Progress:
    """Status text a worker thread writes and a GUI timer shows."""
    text = ""


def ensure_model(name, progress):
    """Path to the model file, downloading it first if needed. Worker thread."""
    file, url, size = AI_MODELS[name]
    path = MODELS_DIR / file
    with _model_locks[name]:
        if path.exists() and path.stat().st_size == size:
            return path
        MODELS_DIR.mkdir(parents=True, exist_ok=True)
        part = path.with_name(path.name + ".part")  # renamed when complete: a cut download is never used
        done = 0
        with urllib.request.urlopen(url, timeout=30) as response, open(part, "wb") as f:
            while chunk := response.read(1 << 20):
                f.write(chunk)
                done += len(chunk)
                progress.text = f"Downloading the AI model, once: {done * 100 // size}% of {size // 2 ** 20} MB"
        if part.stat().st_size != size:
            part.unlink()
            raise RuntimeError("the model download was cut off; try again")
        part.rename(path)
        return path


SAM2_FILES = ("sam2.1_hiera_base_plus.encoder.onnx", "sam2.1_hiera_base_plus.decoder.onnx")
_unpack_lock = threading.Lock()
_sessions = {}  # model path: its ONNX Runtime session, kept for the next picture


def sam2_files(progress):
    """Paths of the Segment Anything 2.1 encoder and decoder, downloaded and unpacked once. Worker thread."""
    paths = [MODELS_DIR / name for name in SAM2_FILES]
    with _unpack_lock:
        if not all(path.exists() for path in paths):
            archive = ensure_model("sam2", progress)
            progress.text = "Unpacking the AI model, once…"
            with zipfile.ZipFile(archive) as zipped:
                for path in paths:
                    part = path.with_name(path.name + ".part")  # renamed when complete, like downloads
                    with zipped.open(path.name) as src, open(part, "wb") as dst:
                        shutil.copyfileobj(src, dst)
                    part.rename(path)
            archive.unlink()  # unpacked: the archive is not needed any more
    return paths


def session(path, keep=True):
    """An ONNX Runtime session for a model file. `keep`: cached for the next picture."""
    import onnxruntime
    if path in _sessions:
        return _sessions[path]
    options = onnxruntime.SessionOptions()
    options.enable_cpu_mem_arena = False  # give memory back after each run: peek stays open between previews
    options.add_session_config_entry("session.intra_op.allow_spinning", "0")  # no busy-waiting when idle
    options.intra_op_num_threads = max(1, (os.cpu_count() or 2) // 2)  # half the cores: the PC stays responsive
    model = onnxruntime.InferenceSession(str(path), options, providers=["CPUExecutionProvider"])
    if keep:
        _sessions[path] = model
    return model


def free_ai_memory():
    """Unload the AI models and hand freed memory back to the system: peek stays running hidden."""
    import ctypes
    import gc
    _sessions.clear()
    gc.collect()
    try:
        ctypes.CDLL("libc.so.6").malloc_trim(0)  # glibc keeps freed big blocks for itself otherwise
    except (OSError, AttributeError):
        pass


def model_pixels(image, width, height):
    """`image` squeezed to width×height, as the 1×3×H×W float array these models read (ImageNet colours)."""
    import numpy as np
    small = image.convertToFormat(QImage.Format_RGB888).scaled(width, height, Qt.IgnoreAspectRatio,
                                                                Qt.SmoothTransformation)
    rows = np.frombuffer(small.constBits(), np.uint8).reshape(height, small.bytesPerLine())
    pixels = rows[:, :width * 3].reshape(height, width, 3).astype(np.float32) / 255
    pixels = (pixels - np.array([0.485, 0.456, 0.406], np.float32)) / np.array([0.229, 0.224, 0.225], np.float32)
    return pixels.transpose(2, 0, 1)[None]


def looks_drawn(image):
    """True for drawings and illustrations: large areas of flat or smoothly shaded colour. Photos have grain,
    so neighbouring pixels almost always differ."""
    # ponytail: one brightness statistic; a very clean photo (studio, heavy denoise) passes as a drawing.
    # Upgrade path: a small classifier model if that ever picks the wrong edge model.
    import numpy as np
    gray = image.convertToFormat(QImage.Format_Grayscale8).scaled(512, 512, Qt.KeepAspectRatio,
                                                                  Qt.FastTransformation)  # no smoothing: keeps grain
    a = np.frombuffer(gray.constBits(), np.uint8).reshape(gray.height(), gray.bytesPerLine())[:, :gray.width()]
    a = a.astype(np.int16)
    dx, dy = np.abs(np.diff(a, axis=1))[:-1], np.abs(np.diff(a, axis=0))[:, :-1]
    return ((dx <= 1) & (dy <= 1)).mean() > 0.45


def matte(image, progress):
    """How sure the edge model is that each pixel belongs to the main subject, 0..1, at 1024×1024.
    ToonOut for drawings, BiRefNet for photos. Worker thread."""
    import numpy as np
    name = "toonout" if looks_drawn(image) else "birefnet"
    model = session(ensure_model(name, progress), keep=False)  # up to 1 GB: loaded per picture, then freed
    progress.text = "Finding the edges…"
    out = model.run(None, {model.get_inputs()[0].name: model_pixels(image, 1024, 1024)})[0][0, 0]
    if out.min() < 0 or out.max() > 1:
        out = 1 / (1 + np.exp(-out))  # log-odds (BiRefNet) to 0..1; ToonOut answers in 0..1 already
    return (out - out.min()) / (out.max() - out.min() + 1e-8)  # 0 = background, 1 = subject


def upscale_array(a, height, width):
    """Bilinear resize of a 2-D numpy array."""
    import numpy as np
    ys = np.clip((np.arange(height) + 0.5) * a.shape[0] / height - 0.5, 0, a.shape[0] - 1)
    xs = np.clip((np.arange(width) + 0.5) * a.shape[1] / width - 0.5, 0, a.shape[1] - 1)
    y0, x0 = ys.astype(int), xs.astype(int)
    y1, x1 = np.minimum(y0 + 1, a.shape[0] - 1), np.minimum(x0 + 1, a.shape[1] - 1)
    fy, fx = (ys - y0)[:, None], (xs - x0)[None, :]
    top = a[y0][:, x0] * (1 - fx) + a[y0][:, x1] * fx
    bottom = a[y1][:, x0] * (1 - fx) + a[y1][:, x1] * fx
    return (top * (1 - fy) + bottom * fy).astype(np.float32)


class Segmenter:
    """Segment Anything 2.1: reads a picture once (a few seconds), then turns a click into the mask
    of the object under it (a few hundredths of a second). Worker threads."""

    def __init__(self, image, progress):
        import numpy as np
        encoder, decoder = (session(path) for path in sam2_files(progress))
        try:
            import scipy.ndimage  # noqa: F401  slow to load: here, not on the first click (see tidy)
        except ImportError:
            pass
        progress.text = "Looking at the picture…"
        self.decoder = decoder
        self.height, self.width = image.height(), image.width()
        gray = image.convertToFormat(QImage.Format_Grayscale8)
        rows = np.frombuffer(gray.constBits(), np.uint8).reshape(self.height, gray.bytesPerLine())
        self.guide = rows[:, :self.width].astype(np.float32) / 255  # full-size brightness, to snap outlines to
        self.edges = None  # the edge model's 0..1 map at full size, once ImagePane has it (see sharpen)
        # The model sees the picture squeezed into 1024×1024; its masks cover that same square at 256 px
        names = [output.name for output in encoder.get_outputs()]
        self.features = dict(zip(names, encoder.run(None, {"image": model_pixels(image, 1024, 1024)})))

    def part(self, points, labels, previous=None, box=None):
        """One picked object: the clicks on it (label 1 = keep here, 0 = not here) and how sure the model is
        that each pixel belongs to it, as log-odds (above 0 = inside). `previous`: the same object before
        this click, so a "not here" click reshapes it instead of starting over. `box` (x0, y0, x1, y1):
        the brush, only what's inside it."""
        import numpy as np
        corners, kinds = (list(zip(box[::2], box[1::2])), [2, 3]) if box else ([], [])  # 2, 3: box corners
        sx, sy = 1024 / self.width, 1024 / self.height
        masks, scores = self.decoder.run(None, {
            **self.features,
            "point_coords": np.array([[(x * sx, y * sy) for x, y in [*points, *corners]]], np.float32),
            "point_labels": np.array([[*labels, *kinds]], np.float32),
            "mask_input": previous["low"] if previous else np.zeros((1, 1, 256, 256), np.float32),
            "has_mask_input": np.ones(1, np.float32) if previous else np.zeros(1, np.float32)})
        scores = scores[0]
        if len(points) == 1 and not previous and not box:
            # One plain click: the model offers a small part, a part and the whole object. Take the biggest
            # one it is still fairly sure of: a click on a body usually means the whole figure.
            good = [i for i in range(len(scores)) if scores[i] >= scores.max() - 0.2]
            best = max(good, key=lambda i: (masks[0, i] > 0).sum())
        else:
            best = int(scores.argmax())
        low = masks[:, best:best + 1]
        logits = upscale_array(low[0, 0], self.height, self.width)
        if box:
            x0, y0, x1, y1 = box
            inside = logits[y0:y1, x0:x1].copy()
            logits[:] = -20  # surely outside
            logits[y0:y1, x0:x1] = inside
        return {"points": points, "labels": labels, "low": low, "logits": logits, "box": box}


def sharpen(inside, mask, edges, radius, bias):
    """`mask` with its outline redrawn by the edge model where the two agree. The edge model draws at
    1024 px with soft hair and fur; Segment Anything at 256 px. Where the edge model disagrees (it only
    knows "main subject", so e.g. a picked pond is background to it) the selection's own outline stays.
    `bias` is the Less/More slider (-8..8): it moves how sure the edge model must be."""
    import numpy as np
    near = box(inside.astype(np.float32), radius)
    band = (near > 0) & (near < 1)  # within `radius` of the outline
    agree = box(((edges > 0.5) == inside).astype(np.float32), 2 * radius) > 0.8
    middle = 0.5 - bias / 8 * 0.4  # More: also keep what the edge model is less sure of
    edges = np.clip((edges - middle) / 0.9 + 0.5, 0, 1)  # its faintest guesses count as fully out: no haze
    return np.where(band & agree, edges, mask)


def merge_parts(parts):
    """All picked objects as one log-odds map. Where two touch, both are only half sure at the shared
    edge: that seam is filled."""
    import numpy as np
    logits = np.maximum.reduce([p["logits"] for p in parts])
    if len(parts) > 1:
        inside = logits > 0
        logits[grow(grow(inside, 4), 4, shrink=True) & ~inside] = 4
    return logits


def box(a, radius):
    """Each pixel replaced by the average of the (2·radius+1)² square around it (running sums: fast)."""
    import numpy as np
    size = 2 * radius + 1
    for axis in (0, 1):  # one direction at a time
        pad = [(0, 0), (0, 0)]
        pad[axis] = (radius + 1, radius)
        sums = np.moveaxis(np.cumsum(np.pad(a, pad, mode="edge"), axis=axis), axis, 0)
        a = np.moveaxis((sums[size:] - sums[:-size]) / size, 0, axis)
    return a


def feather(inside, radius):
    """A hard in/out mask with only its edge softened over `radius` pixels: no haze where the model
    was unsure, and no jagged outline."""
    import numpy as np
    return box(inside.astype(np.float32), radius)


def snap_to_edges(inside, guide, radius):
    """The mask's outline moved onto the picture's own edges nearby (guided filter, He et al. 2010).
    The model draws at 256 px, so on a big picture its outline is blocky and a few pixels off."""
    import numpy as np
    p = inside.astype(np.float32)
    mean_i, mean_p = box(guide, radius), box(p, radius)
    a = (box(guide * p, radius) - mean_i * mean_p) / (box(guide * guide, radius) - mean_i ** 2 + 1e-3)
    b = mean_p - a * mean_i
    return box(a, radius) * guide + box(b, radius) > 0.5


def tidy(inside, seeds):
    """Drops the specks not connected to any click in `seeds` and fills small holes.
    Without scipy the mask is returned as it is."""
    try:
        from scipy import ndimage
    except ImportError:
        return inside
    import numpy as np
    labels, _ = ndimage.label(inside)
    keep = {labels[y, x] for x, y in seeds} - {0}
    if keep:
        inside = np.isin(labels, list(keep))
    holes, count = ndimage.label(~inside)
    if count:
        sizes = np.bincount(holes.ravel())
        border = set(np.unique(np.concatenate([holes[0], holes[-1], holes[:, 0], holes[:, -1]])))
        small = [i for i in range(1, count + 1) if sizes[i] < inside.size / 1000 and i not in border]
        inside = inside | np.isin(holes, small)
    return inside


def grow(inside, steps, shrink=False):
    """A boolean mask grown (or shrunk) by `steps` pixels."""
    out = inside.copy()
    for _ in range(steps):
        op = (lambda a, b: a & b) if shrink else (lambda a, b: a | b)
        out[1:] = op(out[1:], out[:-1].copy())
        out[:-1] = op(out[:-1], out[1:].copy())
        out[:, 1:] = op(out[:, 1:], out[:, :-1].copy())
        out[:, :-1] = op(out[:, :-1], out[:, 1:].copy())
    return out


def subject_layers(mask):
    """From a 0..1 mask: the dimming over everything else, and the outline for the glow, as QImages."""
    import numpy as np
    height, width = mask.shape
    dim = QImage(width, height, QImage.Format_ARGB32_Premultiplied)
    dim.fill(Qt.black)
    dim.setAlphaChannel(gray_image(((1 - mask) * 150).astype(np.uint8)))
    inside = mask > 0.5
    shrunk = grow(inside, max(2, round(max(width, height) / 450)), shrink=True)  # outline grows with the picture
    edge = QImage(width, height, QImage.Format_ARGB32_Premultiplied)
    edge.fill(Qt.white)
    edge.setAlphaChannel(gray_image(((inside & ~shrunk) * 255).astype(np.uint8)))
    return dim, edge


def gray_image(array):
    """A uint8 numpy array as an 8-bit gray QImage (copied, so the array can go)."""
    height, width = array.shape
    return QImage(array.tobytes(), width, height, width, QImage.Format_Grayscale8).copy()


def cut_out(image, mask):
    """`image` with only the masked part visible, cropped to it. None if the mask is empty."""
    import numpy as np
    ys, xs = np.nonzero(mask > 0.5)
    if not len(xs):
        return None
    result = image.convertToFormat(QImage.Format_ARGB32)
    result.setAlphaChannel(gray_image((mask * 255).astype(np.uint8)))
    return result.copy(int(xs.min()), int(ys.min()), int(xs.max() - xs.min() + 1), int(ys.max() - ys.min() + 1))


def waifu2x_models():
    """The photo model folder of the installed waifu2x-ncnn-vulkan, or None."""
    for root in (Path("/usr/share/waifu2x-ncnn-vulkan"), Path("/usr/lib/waifu2x-ncnn-vulkan")):
        for name in ("models-upconv_7_photo", "models-cunet"):
            if (root / name).is_dir():
                return root / name
    return None


def upscale(path, scale, dst, progress):
    """`scale`× bigger image by waifu2x on the graphics card (Vulkan). Worker thread."""
    progress.text = f"Enlarging {scale}×…"
    (CACHE / "tmp").mkdir(parents=True, exist_ok=True)
    src = CACHE / "tmp" / f"upscale-{os.getpid()}-{time.monotonic_ns()}.png"
    decode_image(path).save(str(src))  # waifu2x reads few formats; Qt reads them all
    try:
        run(["waifu2x-ncnn-vulkan", "-i", str(src), "-o", str(dst), "-s", str(scale), "-n", "1",
             "-m", str(waifu2x_models())])
    finally:
        src.unlink(missing_ok=True)


AI_TARGETS = {  # Tools panel entries: code -> (format shown, label, file name suffix)
    "up2": ("PNG", "2× bigger (AI)", " (2x)"),
    "up4": ("PNG", "4× bigger (AI)", " (4x)"),
}


def has_onnx():
    """The local AI runtime is installed (python-onnxruntime-cpu, which brings numpy)."""
    return bool(importlib.util.find_spec("onnxruntime") and importlib.util.find_spec("numpy"))


def ai_image_targets():
    if shutil.which("waifu2x-ncnn-vulkan") and waifu2x_models():
        return [("up2", AI_TARGETS["up2"][1]), ("up4", AI_TARGETS["up4"][1])]
    return []


# ---------------------------------------------------------------- convert panel data

X2T_CODES = {".docx": 65, ".odt": 67, ".rtf": 68, ".txt": 69, ".html": 70, ".epub": 72, ".fb2": 73,
             ".pptx": 129, ".odp": 131, ".xlsx": 257, ".ods": 259, ".csv": 260, ".pdf": 513}
WORD = {".docx", ".docm", ".doc", ".odt", ".rtf", ".epub", ".fb2"}
SHEET = {".xlsx", ".xlsm", ".xls", ".ods", ".csv"}
SLIDES = {".pptx", ".ppt", ".odp"}
RASTER = [".png", ".jpg", ".webp", ".avif", ".tiff", ".bmp"]
AUDIO_OUT = [".mp3", ".ogg", ".opus", ".flac", ".wav", ".m4a"]
VIDEO_OUT = [".mp4", ".webm", ".mkv", ".gif", ".mp3"]
MODEL_OUT = [".glb", ".gltf", ".obj", ".stl", ".fbx", ".ply", ".dae"]


def convert_targets(path):
    """[(extension, label)] this file can be turned into. The label says what the result is."""
    path = Path(path)
    if path.is_dir():
        return []
    suffix, name = path.suffix.lower(), QMimeDatabase().mimeTypeForFile(str(path)).name()
    out = []
    if name.startswith("video/"):
        out = [(e, "audio only" if e == ".mp3" else "animation" if e == ".gif" else "") for e in VIDEO_OUT]
    elif name.startswith("audio/"):
        out = [(e, "") for e in AUDIO_OUT]
    elif name == "application/pdf":
        out = [(".pdf", "smaller file"), (".png", "one image per page"), (".docx", "editable, layout may shift")]
    elif suffix in WORD and X2T.exists():
        out = [(e, "") for e in (".docx", ".odt", ".rtf", ".pdf", ".txt", ".html", ".epub")]
    elif suffix in SHEET and X2T.exists():
        out = [(e, "first sheet" if e == ".csv" else "") for e in (".xlsx", ".ods", ".csv", ".pdf")]
    elif suffix in SLIDES and X2T.exists():
        out = [(e, "") for e in (".pptx", ".odp", ".pdf")]
    elif suffix in {".djvu", ".xps", ".oxps"} and X2T.exists():
        out = [(".pdf", "")]
    elif suffix in MODELS and shutil.which("assimp"):
        out = [(e, "") for e in MODEL_OUT]
    elif suffix == ".blend" and blender_command() and shutil.which("assimp"):
        out = [(e, "") for e in MODEL_OUT]
    elif name in ARCHIVES and shutil.which("bsdtar"):
        out = [("/", "extract into a folder")]
    elif QImageReader(str(path)).canRead():
        writable = {"." + bytes(f).decode() for f in QImageWriter.supportedImageFormats()} | {".jpg"}
        out = ai_image_targets() + [(e, "") for e in RASTER if e in writable] + [(".pdf", "")]
        if QImageReader(str(path)).supportsAnimation():
            out += [(".mp4", "video"), (".webm", "video")]
    elif name in {"text/markdown", "text/html"} or mime_inherits(name, "text/plain"):
        out = [(".pdf", "")] + ([(".html", "")] if name == "text/markdown" else [])
    return [(e, label) for e, label in out if e != suffix or label]  # keep "PDF → smaller PDF"


def mime_inherits(name, parent):
    return QMimeDatabase().mimeTypeForName(name).inherits(parent)


def free_path(folder, stem, ext):
    """`stem.ext` in `folder`, or `stem (2).ext`… when taken: converting never overwrites a file."""
    candidate = folder / f"{stem}{ext}"
    n = 2
    while candidate.exists():
        candidate = folder / f"{stem} ({n}){ext}"
        n += 1
    return candidate


def run(cmd, timeout=1800):
    result = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout, stdin=subprocess.DEVNULL)
    if result.returncode != 0:
        raise RuntimeError((result.stderr or result.stdout).strip().splitlines()[-1:] or ["failed"])


def x2t_convert(src, dst):
    task = CACHE / f"convert-{os.getpid()}-{time.monotonic_ns()}.xml"
    (CACHE / "tmp").mkdir(parents=True, exist_ok=True)
    task.write_text(f"""<?xml version="1.0" encoding="utf-8"?>
<TaskQueueDataConvert>
  <m_sFileFrom>{escape(str(src))}</m_sFileFrom>
  <m_sFileTo>{escape(str(dst))}</m_sFileTo>
  <m_nFormatTo>{X2T_CODES[dst.suffix]}</m_nFormatTo>
  <m_sTempDir>{escape(str(CACHE / "tmp"))}</m_sTempDir>
  <m_sFontDir>{escape(str(X2T_FONTS))}</m_sFontDir>
  <m_sAllFontsPath>{escape(str(X2T_FONTS / "AllFonts.js"))}</m_sAllFontsPath>
  <m_bIsNoBase64>true</m_bIsNoBase64>
</TaskQueueDataConvert>""")
    try:
        run([str(X2T), str(task)])
    finally:
        task.unlink(missing_ok=True)


def convert_file(path, ext, name=None, progress=None):
    """Write `path` converted to `ext` next to it and return the new path. Runs on a worker thread;
    pass the mime `name` from the GUI thread (see file_details)."""
    path = Path(path)
    suffix = path.suffix.lower()
    name = name or QMimeDatabase().mimeTypeForFile(str(path)).name()
    progress = progress or Progress()
    stem = path.name[:-len(".tar" + path.suffix)] if path.name.lower().endswith(".tar" + suffix) else path.stem
    if ext in AI_TARGETS:
        dst = free_path(path.parent, stem + AI_TARGETS[ext][2], ".png")
        upscale(path, int(ext[2:]), dst, progress)
        return dst
    if ext == "/":  # archive: extract into a new folder
        dst = free_path(path.parent, stem, "")
        dst.mkdir()
        run(["bsdtar", "-xf", str(path), "-C", str(dst)])
        return dst
    if name == "application/pdf" and ext == ".pdf":
        dst = free_path(path.parent, f"{stem} (compressed)", ".pdf")
        # /ebook = 150 dpi images: much smaller, still sharp on screen
        run(["gs", "-q", "-dBATCH", "-dNOPAUSE", "-dSAFER", "-sDEVICE=pdfwrite", "-dPDFSETTINGS=/ebook",
             f"-sOutputFile={dst}", str(path)])
        return dst
    if name == "application/pdf" and ext == ".png":
        dst = free_path(path.parent, f"{stem} pages", "")
        dst.mkdir()
        run(["pdftoppm", "-r", "150", "-png", str(path), str(dst / "page")])
        return dst
    dst = free_path(path.parent, stem, ext)
    if name.startswith(("audio/", "video/")) or (ext in {".mp4", ".webm"} and not suffix in MODELS):
        args = {".mp4": ["-c:v", "libx264", "-preset", "veryfast", "-crf", "22", "-pix_fmt", "yuv420p",
                         "-vf", "scale=trunc(iw/2)*2:trunc(ih/2)*2", "-c:a", "aac", "-movflags", "+faststart"],
                ".webm": ["-c:v", "libvpx-vp9", "-deadline", "realtime", "-cpu-used", "8", "-b:v", "0",
                          "-crf", "32", "-c:a", "libopus"],
                ".mkv": ["-c", "copy"],  # same streams in a new box: instant, no quality loss
                ".gif": ["-vf", "fps=12,scale=480:-1:flags=lanczos,split[a][b];[a]palettegen[p];[b][p]paletteuse",
                         "-loop", "0"],
                ".mp3": ["-vn", "-q:a", "2"], ".ogg": ["-vn"], ".opus": ["-vn"], ".wav": ["-vn"],
                ".flac": [], ".m4a": ["-vn"]}[ext]
        run(["ffmpeg", "-y", "-loglevel", "error", "-i", str(path), *args, "-map_metadata", "0", str(dst)])
        return dst
    if ext in X2T_CODES and suffix in WORD | SHEET | SLIDES | {".djvu", ".xps", ".oxps", ".pdf"}:
        x2t_convert(path, dst)
        return dst
    if suffix in MODELS or suffix == ".blend":
        source = path
        if suffix == ".blend":  # Blender exports glTF once (cached for the 3D view too); assimp does the rest
            cmd, part, out, _ = conversion(path)
            if not out.exists():
                CACHE.mkdir(parents=True, exist_ok=True)
                run(cmd)
                part.rename(out)
            source = out
        run(["assimp", "export", str(source), str(dst)])
        return dst
    if ext in RASTER or (ext == ".pdf" and QImageReader(str(path)).canRead()):
        image = decode_image(path)
        if ext == ".pdf":
            writer = QPdfWriter(str(dst))
            # 1 px = 1 pt, exact: the default fuzzy match would snap the page to the nearest paper size
            writer.setPageSize(QPageSize(QSizeF(image.size()), QPageSize.Unit.Point, "",
                                         QPageSize.SizeMatchPolicy.ExactMatch))
            writer.setPageMargins(QMarginsF(0, 0, 0, 0))
            writer.setResolution(72)
            painter = QPainter(writer)
            painter.drawImage(0, 0, image)
            painter.end()
        else:
            if ext in {".jpg", ".bmp"} and image.hasAlphaChannel():  # no transparency there: flatten on white
                flat = QImage(image.size(), QImage.Format_RGB32)
                flat.fill(Qt.white)
                painter = QPainter(flat)
                painter.drawImage(0, 0, image)
                painter.end()
                image = flat
            if not image.save(str(dst), quality=92):
                raise RuntimeError("Qt could not write this format")
        return dst
    raise RuntimeError("no converter for this pair")


def text_to_document(path, ext):
    """Markdown / HTML / text to PDF or HTML with Qt's text engine. GUI thread only (uses fonts)."""
    path = Path(path)
    name = QMimeDatabase().mimeTypeForFile(str(path)).name()
    text = path.read_bytes()[:TEXT_LIMIT].decode("utf-8", errors="replace")
    doc = QTextDocument()
    doc.setBaseUrl(QUrl.fromLocalFile(str(path.parent) + "/"))
    if name == "text/markdown":
        doc.setMarkdown(text)
    elif name == "text/html":
        doc.setHtml(text)
    else:
        doc.setPlainText(text)
        doc.setDefaultFont(QFontDatabase.systemFont(QFontDatabase.FixedFont))
    dst = free_path(path.parent, path.stem, ext)
    if ext == ".html":
        dst.write_text(doc.toHtml())
    else:
        writer = QPdfWriter(str(dst))
        writer.setPageSize(QPageSize(QPageSize.A4))
        doc.print_(writer)
    return dst


# ---------------------------------------------------------------- side panel

class SidePanel(QWidget):
    """The side panel: the file's name and type, everything peek knows about it (read on a worker thread),
    then its tools: one button per action and the latest result. Results land next to the original;
    nothing is overwritten."""
    open_file = Signal(object)

    def __init__(self):
        super().__init__()
        self.path = None
        self.thumb = QLabel(alignment=Qt.AlignCenter)
        self.thumb.setFixedSize(56, 56)
        self.name = QLabel(wordWrap=True, textFormat=Qt.RichText)
        header = QHBoxLayout()
        header.setSpacing(10)
        header.addWidget(self.thumb)
        header.addWidget(self.name, 1)
        self.details = QLabel(textFormat=Qt.RichText, alignment=Qt.AlignTop, wordWrap=True,
                              openExternalLinks=True, textInteractionFlags=Qt.TextBrowserInteraction)
        tools = QLabel("<b>Tools</b>", textFormat=Qt.RichText)
        tools.setContentsMargins(0, 10, 0, 0)
        self.buttons = QVBoxLayout()
        self.buttons.setSpacing(4)
        self.result = QHBoxLayout()  # the latest job: progress, then a link to what it made
        self.result.setSpacing(4)
        note = QLabel("Saved next to the original. Existing files are never replaced.", wordWrap=True)
        note.setEnabled(False)  # the theme's dimmed text color
        inner = QWidget()
        column = QVBoxLayout(inner)
        column.setContentsMargins(10, 8, 10, 10)
        column.setSpacing(8)
        column.addLayout(header)
        column.addWidget(self.details)
        column.addWidget(tools)
        column.addLayout(self.result)
        column.addLayout(self.buttons)
        column.addWidget(note)
        column.addStretch()
        scroll = QScrollArea(widgetResizable=True, frameShape=QFrame.NoFrame)
        scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)  # long values wrap instead
        scroll.setWidget(inner)
        layout = QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.addWidget(scroll)

    def clear_result(self):
        while item := self.result.takeAt(0):
            item.widget().deleteLater()

    def show_file(self, path):
        self.path = path
        self.clear_result()  # a result belongs to the file it came from
        mime = QMimeDatabase().mimeTypeForFile(str(path))
        reader = QImageReader(str(path))
        if mime.name().startswith("image/") and reader.size().isValid():
            reader.setScaledSize(reader.size().scaled(56, 56, Qt.KeepAspectRatio))  # decodes only a small copy
            self.thumb.setPixmap(QPixmap.fromImage(reader.read()))
        else:
            self.thumb.setPixmap(QIcon.fromTheme(mime.iconName(), QIcon.fromTheme("unknown")).pixmap(48))
        self.name.setText(f"<b>{escape(Path(path).name)}</b><br><small>{escape(mime.comment())}</small>")
        self.details.setText("<p style='opacity:.6'>Reading…</p>")
        in_background(lambda: file_details(path, mime), self,
                      lambda sections: path == self.path and self.details.setText(details_html(sections)))
        while item := self.buttons.takeAt(0):
            item.widget().deleteLater()
        targets = convert_targets(path)
        if not targets:
            label = QLabel("No tools for this type of file.", wordWrap=True)
            label.setEnabled(False)
            self.buttons.addWidget(label)
        mimes = QMimeDatabase()
        for ext, label in targets:
            fmt = "Folder" if ext == "/" else AI_TARGETS[ext][0] if ext in AI_TARGETS else ext[1:].upper()
            icon = QIcon.fromTheme("folder" if ext == "/" else "tools-wizard" if ext in AI_TARGETS else
                                   mimes.mimeTypeForFile("x" + ext, QMimeDatabase.MatchExtension).iconName())
            button = QPushButton(icon, f"{fmt}  —  {label}" if label else fmt)
            button.setStyleSheet("text-align: left; padding: 5px 8px")
            button.clicked.connect(lambda _=False, e=ext, f=fmt: self.start(e, f))
            self.buttons.addWidget(button)

    def start(self, ext, fmt):
        path = self.path
        self.clear_result()
        status = QLabel(f"{fmt}…", wordWrap=True, textFormat=Qt.RichText)
        status.linkActivated.connect(lambda link: self.open_file.emit(Path(link)))
        busy = QProgressBar(maximum=0, textVisible=False, maximumWidth=60, maximumHeight=8)  # max 0 = spinning
        self.result.addWidget(status, 1)
        self.result.addWidget(busy)
        progress = Progress()
        ticker = QTimer(status, interval=250)  # the worker writes progress.text; this shows it
        ticker.timeout.connect(lambda: progress.text and status.setText(f"{fmt}<br><small>{escape(progress.text)}</small>"))
        ticker.start()

        def done(result):
            ticker.stop()
            busy.hide()
            if isinstance(result, Exception):
                status.setText(f"<b>{fmt} failed</b><br><span style='opacity:.7'>{escape(str(result))[:300]}</span>")
                return
            # the name is a link: click it to preview the new file here
            status.setText(f'✓ <a href="{escape(str(result))}">{escape(result.name)}</a>')
            status.setToolTip("Preview it here")
            self.result.addWidget(tool_button("document-open-folder", "Show in folder",
                                              lambda: show_in_folder(result)))

        def work():
            try:
                return convert_file(path, ext, name, progress)
            except Exception as error:  # shown in the result line; the panel stays usable
                return error

        text_formats = {"text/markdown", "text/html"}
        name = QMimeDatabase().mimeTypeForFile(str(path)).name()
        if ext in {".pdf", ".html"} and (name in text_formats or mime_inherits(name, "text/plain")) \
                and Path(path).suffix.lower() not in OFFICE | EBOOKS:
            try:  # Qt's text engine wants the GUI thread; it's fast enough there
                done(text_to_document(path, ext))
            except Exception as error:
                done(error)
        else:
            in_background(work, status, done, pool=_convert_pool)


def show_in_folder(path):
    """Open the folder in the file manager with the file selected."""
    if shutil.which("dolphin"):
        QProcess.startDetached("dolphin", ["--select", str(path)])
    else:
        QDesktopServices.openUrl(QUrl.fromLocalFile(str(Path(path).parent)))


def apps_for(mime):
    """Apps that open files of type `mime` (default first), from the desktop's own lists:
    [(name, icon, .desktop path)]."""
    names = []
    defaults = configparser.ConfigParser(interpolation=None, strict=False)
    defaults.read(Path(QStandardPaths.writableLocation(QStandardPaths.GenericConfigLocation)) / "mimeapps.list")
    types = [mime.name(), *mime.allAncestors()]  # a .py file also opens in anything that takes text/plain
    names += [n for t in types for n in defaults.get("Default Applications", t, fallback="").split(";")]
    for folder in QStandardPaths.standardLocations(QStandardPaths.ApplicationsLocation):
        cache = configparser.ConfigParser(interpolation=None, strict=False)
        cache.read(Path(folder) / "mimeinfo.cache")
        names += [n for t in types for n in cache.get("MIME Cache", t, fallback="").split(";")]
    apps = []
    for name in dict.fromkeys(n for n in names if n):  # in order, once each
        file = QStandardPaths.locate(QStandardPaths.ApplicationsLocation, name)
        entry = configparser.ConfigParser(interpolation=None, strict=False)
        try:
            entry.read(file, encoding="utf-8")
        except (configparser.Error, UnicodeDecodeError):
            continue
        if not file or "Desktop Entry" not in entry or entry["Desktop Entry"].get("NoDisplay") == "true" \
                or entry["Desktop Entry"].get("Hidden") == "true":
            continue
        apps.append((entry["Desktop Entry"].get("Name", name), entry["Desktop Entry"].get("Icon", ""), file))
    return apps


def move_to_trash(path):
    """True if `path` went to the Trash (restorable from Dolphin)."""
    result = QFile.moveToTrash(str(path))
    return result[0] if isinstance(result, tuple) else result


# ---------------------------------------------------------------- choosing a preview

def preview_widget(path, mime, player, image=None, on_open=None):
    name = mime.name()
    suffix = Path(path).suffix.lower()
    if Path(path).is_dir():
        return FolderView(Path(path), on_open)
    if name.startswith(("video/", "audio/")):
        player.setSource(QUrl.fromLocalFile(path))
        player.play()
        if name.startswith("audio/"):
            return MediaPane(AudioView(player, mime, path), player)
        video = QVideoWidget()
        player.setVideoOutput(video)
        # Space closes the window like Sushi, so pausing goes on a click
        video.mousePressEvent = lambda e: player.pause() if player.isPlaying() else player.play()
        return MediaPane(video, player)
    if name == "application/pdf":
        return DocView(path)
    if suffix in SHEETS and openpyxl:
        try:
            return SheetView(path)
        except Exception:  # openpyxl rejects some producers' xlsx (e.g. OnlyOffice's): use the text preview
            pass
    if suffix in MODELS:
        return ModelView(path)
    if suffix in FONTS or name.startswith("font/"):
        try:
            return FontView(path)
        except ValueError:
            pass
    # Any format a Qt image plugin decodes, whatever its mime: kimageformats adds kra, psd, xcf, raw...
    if image is not None and not image.isNull():
        return ImagePane(image, path)
    if QImageReader(path).canRead():
        return ImagePane(path, path)
    try:
        extracted = EXTRACTORS[suffix](path) if suffix in EXTRACTORS else None
        text = office_text(path) if suffix in OFFICE else None
    except (OSError, ValueError, KeyError, struct.error, sqlite3.Error, zipfile.BadZipFile,
            ElementTree.ParseError, zstd.ZstdError):
        extracted = text = None  # damaged or unexpected file: fall through to the generic previews
    if extracted and not extracted.isNull():
        return ImagePane(extracted)
    if text:
        return plain_text_pane(text)
    if name == "text/markdown":
        return rich_text_pane(path, "markdown")
    if name in {"text/html", "application/xhtml+xml"}:
        return rich_text_pane(path, "html")
    if name in ARCHIVES:
        return ArchiveView(path)
    if mime.inherits("text/plain"):
        with open(path, "rb") as f:
            text = f.read(TEXT_LIMIT).decode("utf-8", errors="replace")
        return code_pane(text, path)
    thumb = cached_thumbnail(path)
    return ImagePane(thumb) if thumb else icon_label(mime)


EXTRACTORS = {".blend": blend_thumbnail, ".clip": clip_preview}


# ---------------------------------------------------------------- the window

class Peek(QWidget):
    def __init__(self, path):
        super().__init__()
        self.preview = QWidget()
        self.files, self.index = [], 0
        self.preloaded = {}  # path -> Future[QImage] for the neighbours of the shown file
        self.player = QMediaPlayer(self)
        self.player.setAudioOutput(QAudioOutput(self.player))
        # Without GStreamer, Qt 6.12's FFmpeg backend freezes when a very short file reaches its end.
        # Looping never reaches the end; pausing when the position jumps back plays it once.
        self.player.setLoops(QMediaPlayer.Loops.Infinite)
        self.last_pos = 0
        self.player.positionChanged.connect(self.pause_on_wrap)
        self.converter = QProcess(self)  # one at a time: flipping fast must not pile up Blender runs
        self.converter.finished.connect(self.conversion_done)
        self.job = None

        # Bottom bar: draggable file icon, file info, "Open" button
        self.drag_icon = DragIcon()
        self.info = QLabel(textFormat=Qt.RichText)
        open_btn = QPushButton(QIcon.fromTheme("document-open"), "Open")
        open_btn.clicked.connect(self.open_external)
        self.panel = SidePanel()
        self.panel.open_file.connect(self.open)
        line = QFrame(frameShape=QFrame.VLine, frameShadow=QFrame.Plain)
        line.setEnabled(False)  # the theme's faint line color
        self.side_box = QWidget()  # the panel plus the thin line separating it from the preview
        side_layout = QHBoxLayout(self.side_box)
        side_layout.setContentsMargins(6, 0, 0, 0)
        self.panel_top = side_layout  # its top margin keeps the panel under the floating button: resizeEvent
        side_layout.addWidget(line)
        side_layout.addWidget(self.panel, 1)
        self.side_box.setFixedWidth(330)
        self.side_box.hide()
        self.panel_button = tool_button("sidebar-expand-right", "Details and tools (I)", self.toggle_side,
                                        checkable=True)
        self.panel_button.toggled.disconnect()  # toggle_side decides the checked state
        self.panel_button.clicked.connect(self.toggle_side)
        self.panel_tray = Tray(self.panel_button)
        self.panel_tray.setParent(self)  # floats in the top right corner: see resizeEvent
        bar = QHBoxLayout()
        bar.setSpacing(8)
        bar.addWidget(self.drag_icon)
        bar.addWidget(self.info, 1)
        bar.addWidget(open_btn)
        body = QHBoxLayout()
        body.addWidget(self.preview, 1)
        body.addWidget(self.side_box)
        self.layout_ = QVBoxLayout(self)
        self.layout_.addLayout(body, 1)
        self.layout_.addLayout(bar)
        for w in [self.panel, *self.panel.findChildren(QWidget), self.panel_button, open_btn]:
            w.setFocusPolicy(Qt.NoFocus)  # clicking them must not steal Space and the arrow keys

        # Switched off while flying or typing in a search field
        self.shortcuts = [QShortcut(key, self, self.close) for key in ("Space", "Q")]
        self.shortcuts.append(QShortcut("Return", self, self.open_external))
        self.shortcuts += [QShortcut(key, self, lambda: self.step(-1)) for key in ("Left", "Up")]
        self.shortcuts += [QShortcut(key, self, lambda: self.step(1)) for key in ("Right", "Down")]
        self.shortcuts += [QShortcut("I", self, self.toggle_side),
                           QShortcut("C", self, self.toggle_side),
                           QShortcut(QKeySequence.Copy, self, self.copy_file),
                           QShortcut("Ctrl+Shift+C", self, self.copy_path),
                           QShortcut("F2", self, self.rename),
                           QShortcut(QKeySequence.Delete, self, self.trash)]
        QApplication.instance().installEventFilter(self)  # right-clicks anywhere in the window: see eventFilter
        QShortcut("Escape", self, self.escape)
        QShortcut("F11", self, lambda: self.showNormal() if self.isFullScreen() else self.showFullScreen())
        QApplication.instance().focusChanged.connect(self.update_shortcuts)

        self.settings = QSettings("peek", "peek")
        self.player.audioOutput().setVolume(float(self.settings.value("volume", 1.0)))
        if geometry := self.settings.value("geometry"):
            self.restoreGeometry(geometry)
        else:
            screen = QApplication.primaryScreen().availableSize()
            self.resize(int(screen.width() * 0.6), int(screen.height() * 0.7))
        self.open(path)

    def toggle_side(self):
        self.side_box.setVisible(not self.side_open())
        self.refresh_side()

    def resizeEvent(self, event):
        super().resizeEvent(event)
        margins = self.layout_.contentsMargins()
        self.panel_tray.adjustSize()
        self.panel_tray.move(self.width() - margins.right() - self.panel_tray.width(), margins.top())
        self.panel_tray.raise_()
        self.panel_top.setContentsMargins(6, self.panel_tray.height() + 6, 0, 0)

    def side_open(self):
        return not self.side_box.isHidden()  # isVisible() would say False while the window is hidden

    def refresh_side(self, *_):
        """Fill only the visible panel page: browsing with the panel closed costs nothing."""
        self.panel_button.setChecked(self.side_open())
        self.panel_button.setIcon(QIcon.fromTheme("sidebar-collapse-right" if self.side_open()
                                                  else "sidebar-expand-right"))
        if not self.side_open() or not self.files:
            return
        path = self.files[self.index]
        if self.panel.path != path:
            self.panel.show_file(path)
            for w in self.panel.findChildren(QWidget):
                w.setFocusPolicy(Qt.NoFocus)

    def open(self, path):
        path = Path(path)
        self.files = siblings(path)
        self.index = self.files.index(path)
        self.show_file()

    def receive(self, socket):
        """Another `peek FILE` launch: show its file here, or close if it's the file already shown."""
        if not socket.bytesAvailable():
            socket.waitForReadyRead(1000)
        token, _, path = bytes(socket.readAll()).decode().partition("\n")
        socket.deleteLater()
        if path == "--quit":
            QApplication.quit()
            return
        path = Path(path)
        if not path.exists():
            return
        if self.isVisible() and self.files and path == self.files[self.index]:
            self.close()  # Space on the same file again closes, like Sushi
            return
        self.side_box.hide()  # a new launch starts with the panel closed; the arrow keys keep it open
        self.refresh_side()
        self.open(path)
        if token:  # Wayland only lets a window come forward with the launcher's activation token
            os.environ["XDG_ACTIVATION_TOKEN"] = token
        self.show()
        self.raise_()
        self.activateWindow()

    def closeEvent(self, event):
        if isinstance(self.preview, DocView) and not self.preview.set_editing(False):
            event.ignore()  # you chose to keep editing
            return
        # Stay running hidden: the next `peek FILE` shows up instantly instead of starting Python again
        if not self.isFullScreen():
            self.settings.setValue("geometry", self.saveGeometry())
        self.settings.setValue("volume", self.player.audioOutput().volume())
        event.ignore()
        self.hide()
        if isinstance(self.preview, ImagePane) and self.preview.segmenter:
            self.preview.subject_button.setChecked(False)  # lets go of the pick and its history
            self.preview.segmenter = None
            free_ai_memory()
        self.player.stop()
        self.player.setSource(QUrl())
        if self.converter.state() != QProcess.NotRunning:
            self.converter.kill()
            self.converter.waitForFinished()
        self.job = None
        self.preloaded.clear()
        self.set_preview(QWidget())  # free the old preview's memory while hidden

    def escape(self):
        if isinstance(self.preview, ModelView) and self.preview.flying:
            self.preview.set_flying(False)
        elif isinstance(self.preview, DocView) and self.preview.editing:
            if not self.preview.editor.finish_typing():  # first Esc ends the text box, the next one leaves
                self.preview.set_editing(False)
        elif isinstance(self.preview, ImagePane) and self.preview.subject_button.isChecked():
            self.preview.leave_subject()
        elif (bar := self.preview.findChild(SearchBar)) and bar.isVisible():
            bar.close_bar()
        elif self.isFullScreen():
            self.showNormal()
        else:
            self.close()

    def update_shortcuts(self, *_):
        flying = (isinstance(self.preview, ModelView) and self.preview.flying
                  or isinstance(self.preview, DocView) and self.preview.editing)
        typing = isinstance(self.focusWidget(), QLineEdit)
        for shortcut in self.shortcuts:
            shortcut.setEnabled(not (flying or typing))

    def pause_on_wrap(self, pos):
        if pos < self.last_pos and self.player.duration() - self.last_pos < 1500:  # wrapped at the end, not a seek
            self.player.pause()
        self.last_pos = pos

    def step(self, delta):
        self.index = (self.index + delta) % len(self.files)
        self.show_file()

    def show_file(self):
        path = self.files[self.index]
        mime = QMimeDatabase().mimeTypeForFile(str(path))
        self.player.stop()
        self.player.setVideoOutput(None)  # detach before the old video widget is deleted
        future = self.preloaded.get(path)
        image = future.result() if future and future.done() and not future.exception() else None
        self.set_preview(preview_widget(str(path), mime, self.player, image, self.open))

        size = QLocale().formattedDataSize(path.stat().st_size) if path.is_file() else ""
        self.details = f"<b>{escape(path.name)}</b><br>{mime.comment()}  {size}  ·  {self.index + 1}/{len(self.files)}"
        self.info.setText(self.details)
        self.refresh_side()
        self.drag_icon.set_file(path, QIcon.fromTheme(mime.iconName(), QIcon.fromTheme(mime.genericIconName())))
        self.setWindowTitle(path.name)
        self.setWindowIcon(QIcon.fromTheme(mime.iconName()))

        # Show the quick preview first, then swap in the 3D model / real pages once converted
        if self.converter.state() != QProcess.NotRunning:
            self.converter.kill()
            self.converter.waitForFinished()
        self.job = conversion(path) if path.is_file() else None
        if self.job:
            cmd, part, out, task = self.job
            if out.exists():
                out.touch()  # recently used: the cache cleanup keeps it longer
                self.show_converted(out)
            else:
                (CACHE / "tmp").mkdir(parents=True, exist_ok=True)
                if task:
                    out.with_suffix(".xml").write_text(task)
                self.info.setText(self.details + "  ·  <i>loading full preview…</i>")
                self.converter.start(cmd[0], cmd[1:])
        self.preload_neighbours()

    def preload_neighbours(self):
        """Decode the previous and next image on worker threads, so flipping to them is instant."""
        keep = {self.files[self.index]}
        for delta in (-1, 1):
            path = self.files[(self.index + delta) % len(self.files)]
            keep.add(path)
            if path in self.preloaded or not path.is_file():
                continue
            mime = QMimeDatabase().mimeTypeForFile(str(path), QMimeDatabase.MatchExtension).name()
            reader = QImageReader(str(path))
            if ((mime.startswith("image/") or path.suffix.lower() in {".kra", ".ora", ".xcf", ".psd"})
                    and reader.canRead() and not reader.supportsAnimation() and mime != "image/svg+xml"):
                self.preloaded[path] = _preload_pool.submit(decode_image, path)
        for path in list(self.preloaded):
            if path not in keep:
                del self.preloaded[path]

    def conversion_done(self, code, status):
        if not self.job or status != QProcess.NormalExit:
            return  # killed because the user flipped to another file
        _, part, out, _ = self.job
        out.with_suffix(".xml").unlink(missing_ok=True)  # x2t's task file, if this was an Office job
        self.info.setText(self.details)
        if code == 0 and part.exists():
            part.rename(out)
            self.show_converted(out)
            prune_cache()

    def show_converted(self, out):
        self.set_preview(ModelView(out) if out.suffix == ".glb" else DocView(out, self.files[self.index]))

    def set_preview(self, new):
        # Keep keyboard focus on the window, or a text/PDF view would eat the arrow keys.
        # Search fields are the exception: they need focus to type in.
        for w in [new, *new.findChildren(QWidget)]:
            if not w.property("peekFocus"):
                w.setFocusPolicy(Qt.NoFocus)
        if isinstance(new, ModelView):
            new.flying_changed.connect(self.update_shortcuts)
        if isinstance(new, DocView):
            new.editing_changed.connect(self.update_shortcuts)
        self.layout_.replaceWidget(self.preview, new)
        self.preview.deleteLater()
        self.preview = new
        self.panel_tray.raise_()  # the new preview was stacked on top of it
        self.setFocus()
        self.update_shortcuts()

    def open_external(self):
        QDesktopServices.openUrl(QUrl.fromLocalFile(str(self.files[self.index])))
        self.close()

    def eventFilter(self, obj, event):
        """Right-click anywhere in the preview: the file menu. Text views keep their Copy / Select All on top."""
        if (event.type() != QEvent.ContextMenu or not isinstance(obj, QWidget) or obj.window() is not self
                or isinstance(obj, QLineEdit) or self.panel.isAncestorOf(obj)
                or isinstance(self.preview, ModelView) and self.preview.flying):
            return False
        text = next((w for w in (obj, obj.parent()) if isinstance(w, (QPlainTextEdit, QTextEdit))), None)
        if text:
            menu = text.createStandardContextMenu(text.viewport().mapFromGlobal(event.globalPos()))
            menu.addSeparator()
        else:
            menu = QMenu(self)
        self.fill_menu(menu)
        menu.exec(event.globalPos())
        menu.deleteLater()
        return True

    def fill_menu(self, menu):
        path = self.files[self.index]
        mime = QMimeDatabase().mimeTypeForFile(str(path))

        def add(icon, text, slot, key=None, into=menu):
            action = into.addAction(QIcon.fromTheme(icon), text)
            action.triggered.connect(slot)
            if key:
                action.setShortcut(QKeySequence(key))  # shown as a hint; the window's shortcut does the work
            return action

        add("document-open", "Open", self.open_external, "Return")
        apps = apps_for(mime) if path.is_file() else []
        if apps:
            others = menu.addMenu(QIcon.fromTheme("system-run"), "Open With")
            for name, icon, desktop in apps:
                add(icon, name, lambda _=False, d=desktop: self.open_with(d), into=others)
        add("document-open-folder", "Show in Folder", lambda: show_in_folder(path))
        menu.addSeparator()
        add("edit-copy", "Copy", self.copy_file, QKeySequence.Copy)
        add("edit-copy-path", "Copy Path", self.copy_path, "Ctrl+Shift+C")
        menu.addSeparator()
        add("edit-rename", "Rename…", self.rename, "F2")
        add("user-trash", "Move to Trash", self.trash, QKeySequence.Delete)
        if isinstance(self.preview, ImagePane) and shutil.which("plasma-apply-wallpaperimage"):
            add("viewimage", "Set as Wallpaper",
                lambda: QProcess.startDetached("plasma-apply-wallpaperimage", [str(path)]))
        menu.addSeparator()
        panel = add("sidebar-expand-right", "Details and Tools", self.toggle_side, "I")
        panel.setCheckable(True)
        panel.setChecked(self.side_open())
        full = add("view-fullscreen", "Full Screen", lambda: self.showNormal() if self.isFullScreen()
                   else self.showFullScreen(), "F11")
        full.setCheckable(True)
        full.setChecked(self.isFullScreen())
        menu.addSeparator()
        add("window-close", "Close", self.close, "Space")

    def note(self, text):
        """A short message in the file line, gone after two seconds."""
        self.info.setText(f"{self.details}  ·  <i>{escape(text)}</i>")
        QTimer.singleShot(2000, lambda: self.info.setText(self.details))

    def open_with(self, desktop):
        # gio fills in the app's own command line from its .desktop file
        QProcess.startDetached("gio", ["launch", desktop, str(self.files[self.index])])
        self.close()

    def selected_text(self):
        for w in [*self.preview.findChildren(QPlainTextEdit), *self.preview.findChildren(QTextEdit)]:
            if w.isVisible() and w.textCursor().hasSelection():
                return w.textCursor().selectedText().replace("\u2029", "\n")  # Qt marks line breaks with U+2029
        return ""

    def copy_file(self):
        """Ctrl+C: selected text if there is some, else the file itself: Dolphin pastes the file,
        chat apps and image editors paste the picture, text fields paste the path."""
        if text := self.selected_text():
            QApplication.clipboard().setText(text)
            return
        if isinstance(self.preview, ImagePane) and self.preview.mask is not None:
            self.preview.copy_subject()
            return
        path = self.files[self.index]
        url = QUrl.fromLocalFile(str(path))
        data = QMimeData()
        data.setUrls([url])
        data.setText(str(path))
        data.setData("x-special/gnome-copied-files", f"copy\n{url.toString()}".encode())  # GNOME apps' paste
        if isinstance(self.preview, ImagePane) and not self.preview.image.isNull():
            data.setImageData(self.preview.image)
        QApplication.clipboard().setMimeData(data)
        self.note("copied")

    def copy_path(self):
        QApplication.clipboard().setText(str(self.files[self.index]))
        self.note("path copied")

    def rename(self):
        path = self.files[self.index]
        dialog = QInputDialog(self, windowTitle="Rename", labelText="New name:", textValue=path.name)
        field = dialog.findChild(QLineEdit)
        stem = len(path.stem) if path.is_file() and path.suffix else len(path.name)
        QTimer.singleShot(0, lambda: field.setSelection(0, stem))  # like Dolphin: the extension stays unselected
        if not dialog.exec():
            return
        name = dialog.textValue().strip()
        new = path.with_name(name)
        if not name or name == path.name:
            return
        if "/" in name or new.exists():
            QMessageBox.warning(self, "Rename", f"Can't use “{name}”: " +
                                ("a name can't contain /." if "/" in name else "a file with that name exists."))
            return
        try:
            path.rename(new)
        except OSError as error:
            QMessageBox.warning(self, "Rename", f"Couldn't rename: {error.strerror}.")
            return
        self.files[self.index] = new
        self.show_file()

    def trash(self):
        path = self.files[self.index]
        if not move_to_trash(path):
            QMessageBox.warning(self, "Move to Trash", f"Couldn't move “{path.name}” to the Trash.")
            return
        del self.files[self.index]
        if not self.files:
            self.close()
            return
        self.index %= len(self.files)
        self.show_file()
        self.note(f"“{path.name}” moved to the Trash")


def to_local_path(arg):
    """Accept a path, a file:// URL, or the Plasma desktop's desktop:/ URL."""
    url = QUrl(arg)
    if url.scheme() == "file":
        return url.toLocalFile()
    if url.scheme() == "desktop":
        desktop = QStandardPaths.writableLocation(QStandardPaths.DesktopLocation)
        return desktop + url.path()
    return arg


def send_to_running(message):
    """Hand a path (or --quit) to the running peek. False if none is running."""
    socket = QLocalSocket()
    socket.connectToServer(SOCKET)
    if not socket.waitForConnected(300):
        return False
    socket.write(f"{os.environ.get('XDG_ACTIVATION_TOKEN', '')}\n{message}".encode())
    socket.waitForBytesWritten(1000)
    socket.disconnectFromServer()
    return True


def main():
    app = QApplication(sys.argv)
    if sys.argv[1:] == ["--quit"]:  # stop the hidden background peek, e.g. after updating peek
        print("stopped" if send_to_running("--quit") else "peek is not running")
        return 0
    path = Path(to_local_path(sys.argv[1])) if len(sys.argv) == 2 else None
    if not path or not path.exists():
        sys.exit("usage: peek FILE | peek --quit")
    if send_to_running(path.resolve()):
        return 0
    app.setQuitOnLastWindowClosed(False)  # closing hides the window; peek keeps running for next time
    prune_cache(clean_parts=True)
    window = Peek(path.resolve())
    server = QLocalServer(window)
    QLocalServer.removeServer(SOCKET)  # clears a socket left behind by a crashed peek
    server.listen(SOCKET)
    server.newConnection.connect(lambda: window.receive(server.nextPendingConnection()))
    window.show()
    # Return instead of sys.exit() here: the window must be freed when main() ends,
    # before Python tears down QApplication, or Qt deletes widgets twice and segfaults.
    return app.exec()


if __name__ == "__main__":
    sys.exit(main())
